from __future__ import annotations

from datetime import date
from decimal import ROUND_HALF_EVEN, Decimal, localcontext
import hashlib
import importlib.util
import json
from pathlib import Path
import re
import runpy
import shutil
import sys
from typing import Any
import xml.etree.ElementTree as ET
import zipfile

import pytest
from openpyxl import load_workbook  # type: ignore[import-untyped]
from openpyxl.utils.cell import range_boundaries  # type: ignore[import-untyped]

from quant_engine.core import (
    AdapterIdentity,
    CoverageClaim,
    CoverageStatus,
    GenerationManifest,
    SourceArtifact,
    VerificationStatus,
    publish_generation,
    read_current,
    stage_generation,
)
from quant_engine.core.contracts import canonical_json_bytes
from quant_engine.portfolio_ledger import decimal_text


REPOSITORY_ROOT = Path(__file__).resolve().parents[2]
MODULE_PATH = REPOSITORY_ROOT / "tools" / "portfolio_report.py"
SPEC = importlib.util.spec_from_file_location("portfolio_report", MODULE_PATH)
if SPEC is None or SPEC.loader is None:
    raise RuntimeError("unable to load portfolio report builder")
reporter = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = reporter
SPEC.loader.exec_module(reporter)
DEFAULT_PRODUCT_CONTRACT = reporter._PORTABLE_HWM_PRODUCT_CONTRACT
_SYNTHETIC_PRODUCT_CONTRACT_PATHS: set[Path] = set()


def _policy_unavailable_appendices(tmp_path: Path) -> tuple[Path, Path, dict[str, Any]]:
    proof = {
        "provider_id": "portfolio123",
        "data_class": "PROVIDER_LIVE_RETURN",
        "disposition": "EXCLUDE",
        "admitted_item_count": 0,
        "authority_ref": "synthetic-policy-receipt",
        "contract_raw_sha256": "a" * 64,
    }
    reconciliation_rows = []
    attribution_rows = []
    for strategy_id in ("grizzlybears-vtma", "p123-ai-factor", "tt-portfolio"):
        metrics: dict[str, Any] = {
            "status": "AVAILABLE",
            "observations": 3,
            "observed_start": "2026-01-01",
            "observed_end": "2026-01-03",
        }
        bridge: dict[str, Any] = {
            "model_return": 0.1,
            "account_return": 0.2,
            "gap_pp": 10.0,
            "categories": [{"effect_pp": 10.0}],
        }
        row = {
            "strategy_id": strategy_id,
            "verification": "VERIFIED",
            "return_reconciliation": metrics,
            "root_causes": [{"cause": "synthetic"}],
        }
        if strategy_id == "p123-ai-factor":
            metrics.clear()
            metrics.update(
                {
                    "status": "UNAVAILABLE_POLICY_NOT_ADMITTED",
                    "observations": 0,
                    "model_cumulative_return": None,
                    "account_minus_model_cumulative": None,
                    "policy_evidence": proof,
                }
            )
            row["account_live_return_summary"] = {
                "status": "AVAILABLE",
                "observations": 3,
                "observed_start": "2026-01-01",
                "observed_end": "2026-01-03",
                "cumulative_return": 0.2,
            }
            row["return_window_comparison"] = {
                "status": "UNAVAILABLE_POLICY_NOT_ADMITTED",
                "expanded_exact_common_gap": None,
                "expanded_minus_declared_gap": None,
                "prior_declared_overlap": None,
                "policy_evidence": proof,
            }
            bridge = {
                "status": "UNAVAILABLE_POLICY_NOT_ADMITTED",
                "model_return": None,
                "account_return": 0.2,
                "gap_pp": None,
                "categories": [],
                "policy_evidence": proof,
            }
        reconciliation_rows.append(row)
        attribution_rows.append(
            {
                "strategy_id": strategy_id,
                "verification": "PARTIAL_EXACT_ACTUAL_PNL_PARTIAL_COUNTERFACTUAL_GAP",
                "observed_start": "2026-01-01",
                "observed_end": "2026-01-03",
                "bridge": bridge,
            }
        )
        if strategy_id == "p123-ai-factor":
            for field in ("execution_shortfall", "effective_account_model_mix"):
                attribution_rows[-1][field] = {
                    "status": "UNAVAILABLE_POLICY_NOT_ADMITTED",
                    "policy_evidence": proof,
                }
            attribution_rows[-1]["actual_return_contribution_by_trade_provenance"] = {
                "daily_observations": 3,
                "contributions_pp": {"actual_activity": 20.0},
            }
    reconciliation_path = tmp_path / "reconciliation.json"
    attribution_path = tmp_path / "attribution.json"
    reconciliation_path.write_text(
        json.dumps(
            {
                "schema_version": 1,
                "builder": "portfolio-implementation-reconciliation",
                "verification": "VERIFIED",
                "as_of_date": "2026-01-03",
                "reconciliation_id": "synthetic",
                "strategies": reconciliation_rows,
            }
        ),
        encoding="utf-8",
    )
    attribution_path.write_text(
        json.dumps(
            {
                "schema_version": 1,
                "builder": "portfolio-implementation-return-attribution",
                "verification": "VERIFIED_WITH_DECLARED_PARTIAL_CAUSAL_BOUNDARIES",
                "as_of_date": "2026-01-03",
                "attribution_id": "synthetic",
                "strategies": attribution_rows,
            }
        ),
        encoding="utf-8",
    )
    return reconciliation_path, attribution_path, proof


@pytest.mark.parametrize("attribution", [False, True])
@pytest.mark.parametrize(
    "tamper", [None, "unbound", "source", "number", "status", "strategy"]
)
def test_policy_unavailable_comparison_requires_source_proof_without_invented_returns(
    tmp_path: Path,
    attribution: bool,
    tamper: str | None,
) -> None:
    reconciliation_path, attribution_path, proof = _policy_unavailable_appendices(
        tmp_path
    )
    path = attribution_path if attribution else reconciliation_path
    validator = (
        reporter._implementation_return_attribution_contract
        if attribution
        else reporter._implementation_reconciliation_contract
    )
    payload = json.loads(path.read_text(encoding="utf-8"))
    row = payload["strategies"][1]
    comparison = row["bridge" if attribution else "return_reconciliation"]
    if tamper == "source":
        comparison["policy_evidence"]["contract_raw_sha256"] = "b" * 64
    elif tamper == "number":
        comparison["model_return" if attribution else "model_cumulative_return"] = 0.0
    elif tamper == "status":
        comparison["status"] = "AVAILABLE"
    elif tamper == "strategy":
        row["strategy_id"] = "tt-portfolio"
    path.write_text(json.dumps(payload), encoding="utf-8")
    if tamper is not None:
        with pytest.raises(reporter.ReportValidationError):
            validator(
                path,
                portfolio_as_of="2026-01-03",
                expected_policy_evidence=None if tamper == "unbound" else proof,
            )
    else:
        rows, reference = validator(
            path, portfolio_as_of="2026-01-03", expected_policy_evidence=proof
        )
        assert len(rows) == reference["strategy_count"] == 3
        assert (
            rows["p123-ai-factor"]["bridge"]["account_return"] == 0.2
            if attribution
            else rows["p123-ai-factor"]["account_live_return_summary"][
                "cumulative_return"
            ]
            == 0.2
        )


@pytest.mark.parametrize("tamper", ["missing", "amount", "count"])
def test_policy_unavailable_retains_reconciled_actual_trade_attribution(
    tmp_path: Path, tamper: str
) -> None:
    _, path, proof = _policy_unavailable_appendices(tmp_path)
    payload = json.loads(path.read_text(encoding="utf-8"))
    row = payload["strategies"][1]
    if tamper == "missing":
        del row["actual_return_contribution_by_trade_provenance"]
    elif tamper == "amount":
        row["actual_return_contribution_by_trade_provenance"]["contributions_pp"][
            "actual_activity"
        ] = 21.0
    else:
        row["actual_return_contribution_by_trade_provenance"]["daily_observations"] = 0
    path.write_text(json.dumps(payload), encoding="utf-8")
    with pytest.raises(reporter.ReportValidationError):
        reporter._implementation_return_attribution_contract(
            path, portfolio_as_of="2026-01-03", expected_policy_evidence=proof
        )


def test_policy_unavailable_comparison_suppresses_historical_attribution_metrics(
    tmp_path: Path,
) -> None:
    reconciliation_path, _, proof = _policy_unavailable_appendices(tmp_path)
    rows, _ = reporter._implementation_reconciliation_contract(
        reconciliation_path,
        portfolio_as_of="2026-01-03",
        expected_policy_evidence=proof,
    )
    strategy_id = "p123-ai-factor"
    result = reporter._merge_attribution(
        [
            {
                "strategy_id": strategy_id,
                "implementation_reconciliation": rows[strategy_id],
            }
        ],
        {
            "benchmark_assignments": [{"strategy_id": strategy_id}],
            "comparisons": [
                {
                    "strategy_id": strategy_id,
                    "status": "DESCRIPTIVE_EXACT_DATE_VERSION_UNRESOLVED",
                    "metrics": {
                        "model_return": 0.1,
                        "cumulative_gap": -0.3,
                        "reference_cumulative_return": 0.5,
                        "observations": 200,
                    },
                }
            ],
            "costs": [{"strategy_id": strategy_id}],
        },
    )[0]
    assert result["provider_comparison_status"] == "UNAVAILABLE_POLICY_NOT_ADMITTED"
    assert result["provider_reference_cumulative_return"] is None
    assert result["provider_cumulative_gap"] is None
    assert result["provider_comparison_observations"] is None
    assert (
        result["implementation_reconciliation"]["account_live_return_summary"][
            "cumulative_return"
        ]
        == 0.2
    )


@pytest.mark.parametrize(
    ("change", "expected"),
    [
        ({}, 3),
        ({"strategy_version": "other-version"}, 0),
        ({"mode_admission": "EXCLUDE_COMPARISON_ONLY"}, 0),
        ({"evidence_layer": "PROVIDER_LIVE"}, 0),
        ({"observed_end": "2025-01-02"}, 0),
        ({"observed_end": "2025-01-03"}, 0),
    ],
)
def test_report_projects_only_accepted_exact_version_prehistory(
    change: dict[str, Any], expected: int
) -> None:
    history = {
        "segments": [
            {
                "evidence_layer": "ACCOUNT_LIVE",
                "mode_admission": "INCLUDE_SELECTED_LIVE",
                "observed_start": "2025-01-02",
                "strategy_version": "synthetic-v1",
            },
            {
                "evidence_layer": "BACKTEST",
                "mode_admission": "INCLUDE_VERSION_COMPATIBLE_PREHISTORY",
                "observed_end": "2024-12-31",
                "return_observations": 3,
                "strategy_version": "synthetic-v1",
                **change,
            },
        ]
    }
    assert reporter._accepted_pre_live_observations(history, "synthetic-v1") == expected
    assert (
        reporter._accepted_pre_live_observations({"segments": []}, "synthetic-v1") == 0
    )
    history["segments"] = history["segments"][1:]
    assert reporter._accepted_pre_live_observations(history, "synthetic-v1") == 0


def test_report_history_requires_actual_packaged_coverage_authority(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    namespace = runpy.run_path(
        str(
            REPOSITORY_ROOT
            / "tests/portfolio_analytics/test_portfolio_current_dependence.py"
        )
    )
    current = namespace["current"]
    selected = namespace["_unmocked_fixture"](tmp_path, monkeypatch)
    built = current.build_source(
        contract_path=selected["contract_path"],
        policy_path=selected["policy_path"],
        history_generation_root=selected["history_root"],
        allocation_generation_root=selected["allocation_root"],
        evaluated_on=date(2026, 1, 3),
    )
    source = tmp_path / built["receipt"]["bundle_id"]
    source.mkdir()
    for name in (
        "dependence-source.json",
        "history-coverage.json",
        "returns.csv",
        "segments.csv",
        "spec.json",
    ):
        (source / name).write_bytes(built[name])
    monkeypatch.setattr(
        reporter, "_PORTABLE_HWM_PRODUCT_CONTRACT", selected["contract_path"]
    )
    monkeypatch.setattr(reporter, "DEPENDENCE_POLICY", selected["policy_path"])
    helpers = runpy.run_path(
        str(REPOSITORY_ROOT / "tests/core/mvp2_consumers_fixture.py")
    )
    helpers["_install_dynamic_dependence_loader"](monkeypatch, reporter, current)
    names = set(json.loads(built["spec.json"])["series_ids"])
    rows, _identity = reporter._history_contract(
        source / "dependence-source.json",
        portfolio_as_of="2026-01-03",
        strategy_names=names | {"Excluded synthetic strategy"},
    )
    assert all(rows[name]["backtest_observations"] == 0 for name in names)
    assert all(
        rows[name]["history_status"] == "UNAVAILABLE_NO_MATCHED_BACKTEST_PREHISTORY"
        for name in names
    )
    # A new claim in the consumer-side coverage file cannot manufacture history.
    coverage = json.loads((source / "history-coverage.json").read_bytes())
    coverage["strategies"] = [{"series_id": next(iter(names)), "segments": []}]
    (source / "history-coverage.json").write_bytes(canonical_json_bytes(coverage))
    with pytest.raises(
        reporter.ReportValidationError, match="history coverage differs"
    ):
        reporter._history_contract(
            source / "dependence-source.json",
            portfolio_as_of="2026-01-03",
            strategy_names=names | {"Excluded synthetic strategy"},
        )


HWM_TEST_NAMESPACE = runpy.run_path(
    str(REPOSITORY_ROOT / "tests" / "portfolio_ledger" / "test_hwm.py")
)


def _portable_hwm_bundle(
    *, schema_version: int = 2
) -> tuple[dict[str, Any], bytes, bytes]:
    fixture_name = (
        "_as_held_direct_net_return_fixture"
        if schema_version == 3
        else "_as_held_hwm_fixture"
    )
    fixture = HWM_TEST_NAMESPACE[fixture_name]()
    return (
        HWM_TEST_NAMESPACE["_compute_as_held"](fixture),
        fixture["input_generation_bytes"],
        fixture["calendar_bytes"],
    )


def _portable_hwm_receipt(*, schema_version: int = 2) -> dict[str, Any]:
    return _portable_hwm_bundle(schema_version=schema_version)[0]


def test_public_report_value_redacts_portfolio123_entity_identifiers() -> None:
    value = {
        "provider": {"entity_id": "1853536", "provider_id": "portfolio123"},
        "version_ref": "p123-entity:1832531",
        "reason": "No exact return series reconciles entity 1832531.",
        "unchanged": "13 accepted holdings",
    }

    assert reporter._public_report_value(value) == {
        "provider": {"entity_id": "[redacted]", "provider_id": "portfolio123"},
        "version_ref": "p123-entity:[redacted]",
        "reason": "No exact return series reconciles connected entity.",
        "unchanged": "13 accepted holdings",
    }


def _excel_saved_package(
    path: Path,
    *,
    modified: str,
    chart_guid: str,
    temporary_root: str,
    revision_guid: str,
    semantic_value: str = "same",
    zip_second: int = 0,
) -> None:
    parts = {
        "docProps/core.xml": (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<cp:coreProperties xmlns:cp="urn:core" '
            'xmlns:dcterms="urn:terms" xmlns:xsi="urn:xsi">'
            '<dcterms:created xsi:type="dcterms:W3CDTF">'
            "2026-08-03T00:00:00Z</dcterms:created>"
            '<dcterms:modified xsi:type="dcterms:W3CDTF">'
            f"{modified}</dcterms:modified></cp:coreProperties>"
        ).encode(),
        "xl/charts/chart1.xml": (
            '<c:chartSpace xmlns:c="urn:chart" xmlns:c16="urn:chart16">'
            "<c:extLst><c:ext><c16:uniqueId "
            f'val="{{{chart_guid}}}"/></c:ext></c:extLst></c:chartSpace>'
        ).encode(),
        "xl/workbook.xml": (
            '<workbook xmlns:x15ac="urn:abs" xmlns:xr="urn:revision">'
            f'<x15ac:absPath url="{temporary_root}"/>'
            f'<xr:revisionPtr documentId="8_{{{revision_guid}}}"/>'
            "</workbook>"
        ).encode(),
        "xl/worksheets/sheet1.xml": f"<sheet>{semantic_value}</sheet>".encode(),
    }
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, content in parts.items():
            info = zipfile.ZipInfo(name, (2026, 8, 3, 0, 0, zip_second))
            info.compress_type = zipfile.ZIP_DEFLATED
            archive.writestr(info, content)


def test_native_excel_workbook_canonicalization_is_deterministic_and_bounded(
    tmp_path: Path,
) -> None:
    first = tmp_path / "first.xlsx"
    second = tmp_path / "second.xlsx"
    changed = tmp_path / "changed.xlsx"
    _excel_saved_package(
        first,
        modified="2026-08-04T05:23:48Z",
        chart_guid="00000000-CC01-4263-A5C9-803DDF0C079C",
        temporary_root="C:\\Temp\\xlsx-win-stage-first\\",
        revision_guid="878E45FB-7799-4293-81FE-29F635965A3D",
    )
    _excel_saved_package(
        second,
        modified="2026-08-04T05:58:22Z",
        chart_guid="00000000-DD23-47F6-8C5E-35307AE221C8",
        temporary_root="C:\\Temp\\xlsx-win-stage-second\\",
        revision_guid="90E8A27E-6EEF-4295-BB7E-DD3062E36063",
        zip_second=2,
    )
    _excel_saved_package(
        changed,
        modified="2026-08-04T05:58:22Z",
        chart_guid="00000000-DD23-47F6-8C5E-35307AE221C8",
        temporary_root="C:\\Temp\\xlsx-win-stage-second\\",
        revision_guid="90E8A27E-6EEF-4295-BB7E-DD3062E36063",
        semantic_value="different financial content",
        zip_second=2,
    )

    first_result = reporter.canonicalize_workbook(
        first, tmp_path / "first-evidence.json"
    )
    second_result = reporter.canonicalize_workbook(
        second, tmp_path / "second-evidence.json"
    )
    changed_result = reporter.canonicalize_workbook(changed)

    assert first.read_bytes() == second.read_bytes()
    assert first_result == second_result
    assert changed_result["workbook_sha256"] != first_result["workbook_sha256"]
    assert reporter.canonicalize_workbook(first) == first_result


def test_initial_workbook_build_is_stable_across_volatile_package_metadata(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    calls = 0

    def volatile_workbook(_report: object, path: Path) -> dict[str, int]:
        nonlocal calls
        calls += 1
        if calls == 1:
            _excel_saved_package(
                path,
                modified="2026-08-29T10:00:00Z",
                chart_guid="10000000-CC01-4263-A5C9-803DDF0C079C",
                temporary_root=r"C:\Temp\first-build",
                revision_guid="178E45FB-7799-4293-81FE-29F635965A3D",
            )
        else:
            _excel_saved_package(
                path,
                modified="2026-08-29T10:00:04Z",
                chart_guid="20000000-DD23-47F6-8C5E-35307AE221C8",
                temporary_root=r"C:\Temp\second-build",
                revision_guid="290E8A27-6EEF-4295-BB7E-DD3062E36063",
                zip_second=4,
            )
        return {"formula_count": 1}

    monkeypatch.setattr(reporter, "_workbook", volatile_workbook)
    path = tmp_path / "portfolio-report.xlsx"

    first = reporter._build_workbook_exact({}, path)
    first_bytes = path.read_bytes()
    second = reporter._build_workbook_exact({}, path)

    assert second == first
    assert path.read_bytes() == first_bytes
    assert first["workbook_sha256"] == hashlib.sha256(first_bytes).hexdigest()


def test_publication_snapshot_preserves_bundle_root_identity(tmp_path: Path) -> None:
    bundle_id = "a" * 64
    source = tmp_path / "source" / bundle_id
    source.mkdir(parents=True)
    expected = {
        "dependence-source.json": b"receipt\n",
        "returns.csv": b"date,return\n",
        "segments.csv": b"segment,date,return\n",
        "spec.json": b"{}\n",
    }
    for name, content in expected.items():
        (source / name).write_bytes(content)

    snapshot = reporter._snapshot_tree_preserving_root_identity(
        source,
        tmp_path / "snapshots",
        {},
        {},
    )

    assert snapshot.name == bundle_id
    assert {
        path.name: path.read_bytes() for path in snapshot.iterdir() if path.is_file()
    } == expected


def test_canonicalization_fails_closed_on_unrecognized_volatile_field(
    tmp_path: Path,
) -> None:
    workbook = tmp_path / "unrecognized.xlsx"
    _excel_saved_package(
        workbook,
        modified="2026-08-04T05:23:48Z",
        chart_guid="not-a-guid-xxxxxxxxxxxxxxxxxxxxxxxxx",
        temporary_root="C:\\Temp\\xlsx-win-stage-first\\",
        revision_guid="878E45FB-7799-4293-81FE-29F635965A3D",
    )

    with pytest.raises(
        reporter.ReportValidationError, match="canonicalization is incomplete"
    ):
        reporter.canonicalize_workbook(workbook)


def _xlsx_result(run_id: str, temporary_root: str) -> dict[str, Any]:
    return {
        "schema_version": "2.0",
        "run_id": run_id,
        "idempotency_key": "portfolio-report-synthetic",
        "final_state": "SUCCEEDED",
        "steps": [
            {
                "step_index": 0,
                "type": "open",
                "status": "succeeded",
                "message": f"Opened '{temporary_root}\\portfolio-report.xlsx'.",
            },
            {
                "step_index": 1,
                "type": "recalc",
                "status": "succeeded",
                "message": "Calculation reached xlDone.",
            },
            {
                "step_index": 2,
                "type": "save_as",
                "status": "succeeded",
                "message": f"Saved to '{temporary_root}\\saved.xlsx'.",
            },
        ],
        "invariants": [],
        "ok": True,
    }


def _candidate_xlsx_result(
    candidate: Path, run_id: str, temporary_root: str
) -> dict[str, Any]:
    result = _xlsx_result(run_id, temporary_root)
    job = json.loads((candidate / "xlsx-job.json").read_text(encoding="utf-8"))
    result["idempotency_key"] = job["idempotency_key"]
    return result


def _prepare_workbook_evidence(
    candidate: Path,
    *,
    run_id: str = "e" * 32,
    temporary_root: str = r"C:\Temp\xlsx-binding",
) -> tuple[Path, Path, Path, Path]:
    """Create synthetic native-Excel output before binding its evidence."""

    initial_workbook = candidate / reporter.WORKBOOK_NAME
    final_workbook = candidate / "excel" / reporter.WORKBOOK_NAME
    final_workbook.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(initial_workbook, final_workbook)
    run_result = candidate / "xlsx-result.json"
    validation = candidate / "validation.json"
    run_result.write_text(
        json.dumps(_candidate_xlsx_result(candidate, run_id, temporary_root)),
        encoding="utf-8",
    )
    validation.write_text(
        json.dumps(
            {
                "all_passed": True,
                "invariants": [{"name": "sentinel", "passed": True}],
            }
        ),
        encoding="utf-8",
    )
    reporter.canonicalize_workbook(
        final_workbook, candidate / "workbook-canonicalization.json"
    )
    reporter.canonicalize_xlsx_result(
        run_result, candidate / "xlsx-result-canonicalization.json"
    )
    return (
        final_workbook,
        run_result,
        validation,
        candidate / "workbook-evidence-binding.json",
    )


def _published_report_delivery(
    tmp_path: Path, *, run_id: str
) -> tuple[Path, str, Path]:
    baseline_root = _published_baseline(tmp_path / "baseline")
    candidate = Path(
        reporter.build_candidate(baseline_root, tmp_path / "candidates")[
            "candidate_dir"
        ]
    )
    workbook, run_result, validation, binding_path = _prepare_workbook_evidence(
        candidate,
        run_id=run_id,
        temporary_root=r"C:\\Temp\\xlsx-receipt-mutation",
    )
    reporter.bind_workbook_evidence(
        candidate, workbook, run_result, validation, binding_path
    )
    published_root = tmp_path / "published"
    published = reporter.publish_report(
        baseline_root,
        candidate,
        workbook,
        run_result,
        validation,
        published_root,
        workbook_evidence_binding=binding_path,
    )
    generation_id = published["generation_id"]
    assert isinstance(generation_id, str)
    return (
        published_root,
        generation_id,
        published_root / "deliveries" / generation_id,
    )


def _rewrite_xlsx_cell(
    path: Path,
    *,
    sheet_name: str,
    cell_reference: str,
    value: str | None = None,
    formula: str | None = None,
    cached_value: str | None = None,
) -> None:
    """Mutate one OOXML cell while preserving every unrelated cached value."""

    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        sheet_index = workbook.sheetnames.index(sheet_name) + 1
    finally:
        workbook.close()
    part_name = f"xl/worksheets/sheet{sheet_index}.xml"
    temporary = path.with_name(f".{path.name}.mutated")
    with (
        zipfile.ZipFile(path) as source,
        zipfile.ZipFile(
            temporary, "w", compression=zipfile.ZIP_DEFLATED
        ) as destination,
    ):
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename == part_name:
                root = ET.fromstring(content)
                cell = next(
                    element
                    for element in root.iter()
                    if reporter._local_name(element.tag) == "c"
                    and element.attrib.get("r") == cell_reference
                )
                namespace = cell.tag.rsplit("}", 1)[0].lstrip("{")
                if value is not None:
                    for child in list(cell):
                        cell.remove(child)
                    cell.set("t", "inlineStr")
                    inline = ET.SubElement(cell, f"{{{namespace}}}is")
                    text = ET.SubElement(inline, f"{{{namespace}}}t")
                    text.text = value
                if formula is not None:
                    formula_node = next(
                        child
                        for child in cell
                        if reporter._local_name(child.tag) == "f"
                    )
                    formula_node.text = formula
                if cached_value is not None:
                    value_node = next(
                        child
                        for child in cell
                        if reporter._local_name(child.tag) == "v"
                    )
                    value_node.text = cached_value
                content = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            destination.writestr(item, content)
    temporary.replace(path)


def _first_table_data_cell(path: Path, sheet_name: str, table_name: str) -> str:
    workbook = load_workbook(path, read_only=False, data_only=False, keep_links=False)
    try:
        table = workbook[sheet_name].tables[table_name]
        min_column, min_row, _max_column, _max_row = range_boundaries(table.ref)
        return f"{workbook[sheet_name].cell(row=min_row + 1, column=min_column).column_letter}{min_row + 1}"
    finally:
        workbook.close()


def test_xlsx_result_canonicalization_preserves_distinct_raw_run_ids(
    tmp_path: Path,
) -> None:
    first = tmp_path / "first-result.json"
    second = tmp_path / "second-result.json"
    first.write_text(
        json.dumps(_xlsx_result("a" * 32, r"C:\Temp\xlsx-first")), encoding="utf-8"
    )
    second.write_text(
        json.dumps(_xlsx_result("b" * 32, r"C:\Temp\xlsx-second")),
        encoding="utf-8",
    )

    first_result = reporter.canonicalize_xlsx_result(
        first, tmp_path / "first-result-evidence.json"
    )
    second_result = reporter.canonicalize_xlsx_result(
        second, tmp_path / "second-result-evidence.json"
    )

    assert first.read_bytes() != second.read_bytes()
    assert json.loads(first.read_text(encoding="utf-8"))["run_id"] == "a" * 32
    assert json.loads(second.read_text(encoding="utf-8"))["run_id"] == "b" * 32
    assert first_result != second_result
    assert reporter.canonicalize_xlsx_result(first) == first_result
    failed = tmp_path / "failed-result.json"
    payload = _xlsx_result("c" * 32, r"C:\Temp\xlsx-failed")
    payload["steps"][1]["status"] = "failed"
    failed.write_text(json.dumps(payload), encoding="utf-8")
    with pytest.raises(reporter.ReportValidationError, match="outcome differs"):
        reporter.canonicalize_xlsx_result(failed)


def test_provider_coverage_retains_declared_out_of_boundary_limitations() -> None:
    rows = reporter._coverage(
        {
            "calendar_coverage": [],
            "fact_coverage": [],
            "provider_coverage": [
                {
                    "coverage": {
                        "known_gaps": ["missing activity type"],
                        "out_of_boundary": ["checking account excluded"],
                        "status": "PARTIAL",
                    },
                    "provider_id": "synthetic",
                    "qualification": "ACCOUNT_LIVE",
                }
            ],
            "reconciliation_exceptions": [],
        }
    )

    assert rows == [
        {
            "category": "PROVIDER_COVERAGE",
            "item": "synthetic",
            "known_gap": "missing activity type; checking account excluded",
            "status": "PARTIAL",
            "value": "ACCOUNT_LIVE",
        }
    ]


def test_positions_preserve_history_and_mark_latest_snapshot_per_account() -> None:
    catalog = {
        "positions": [
            {
                "account_id": "account-a",
                "as_of_date": "2026-07-01",
                "mapping_status": "MAPPED",
                "market_value": _amount("100"),
                "quantity": "1",
                "security_id": "security-old",
                "strategy_id": "strategy-1",
            },
            {
                "account_id": "account-a",
                "as_of_date": "2026-08-01",
                "mapping_status": "MAPPED",
                "market_value": _amount("110"),
                "quantity": "1",
                "security_id": "security-new",
                "strategy_id": "strategy-1",
            },
            {
                "account_id": "account-b",
                "as_of_date": "2026-07-15",
                "mapping_status": "MAPPED",
                "market_value": _amount("25"),
                "quantity": "2",
                "security_id": "security-b",
                "strategy_id": "strategy-2",
            },
        ]
    }

    rows = reporter._positions(catalog)

    assert {(row["account_id"], row["as_of_date"]) for row in rows} == {
        ("account-a", "2026-07-01"),
        ("account-a", "2026-08-01"),
        ("account-b", "2026-07-15"),
    }
    assert {row["security_id"] for row in rows} == {
        "security-old",
        "security-new",
        "security-b",
    }
    assert {row["security_id"]: row["snapshot_status"] for row in rows} == {
        "security-old": "HISTORICAL_ACCEPTED_ACCOUNT_SNAPSHOT",
        "security-new": "LATEST_ACCEPTED_ACCOUNT_SNAPSHOT",
        "security-b": "LATEST_ACCEPTED_ACCOUNT_SNAPSHOT",
    }


def test_performance_preserves_pool_disposition_and_account_live_return_path() -> None:
    catalog = {
        "assignments": [
            {
                "account_id": "account-live",
                "assignment_id": "assignment-live",
                "effective_end": None,
                "effective_start": "2026-01-01",
                "mapping_status": "MAPPED",
                "pool_id": "systematic-capital",
                "pool_mapping_status": "MAPPED",
                "strategy_id": "strategy-live",
                "strategy_mapping_disposition": "STRATEGY_MAPPED",
            },
            {
                "account_id": "account-reserve",
                "assignment_id": "assignment-reserve",
                "effective_end": None,
                "effective_start": "2026-01-01",
                "mapping_status": "UNMAPPED",
                "pool_id": "reserve-capital",
                "pool_mapping_status": "MAPPED",
                "strategy_id": None,
                "strategy_mapping_disposition": "RESERVE_NOT_APPLICABLE",
            },
        ],
        "calendar_coverage": [],
        "flows": [],
        "nav_observations": [],
        "strategy_provider_evidence": [
            {
                "account": {"observed_start": "2026-01-01"},
                "evidence_layer": "ACCOUNT_LIVE",
                "generation_id": "a" * 64,
                "strategy_id": "strategy-live",
            }
        ],
        "strategy_returns": [
            {
                "date": "2026-01-01",
                "evidence_layer": "ACCOUNT_LIVE",
                "provider_id": "synthetic",
                "return": "0.01",
                "strategy_id": "strategy-live",
            },
            {
                "date": "2026-01-02",
                "evidence_layer": "ACCOUNT_LIVE",
                "provider_id": "synthetic",
                "return": "0.02",
                "strategy_id": "strategy-live",
            },
        ],
    }

    rows = reporter._account_strategy_performance(catalog)
    live = next(row for row in rows if row["assignment_id"] == "assignment-live")
    reserve = next(row for row in rows if row["assignment_id"] == "assignment-reserve")

    assert live["twr_status"] == "UNAVAILABLE_INSUFFICIENT_NAV"
    assert live["return_path_status"] == "AVAILABLE_ACCOUNT_LIVE_RETURN_PATH"
    assert live["return_path_observations"] == 2
    assert live["return_path_cumulative_return"] == "0.0302"
    assert (
        live["return_path_account_binding"]
        == "UNIQUE_CURRENT_EFFECTIVE_DATED_ASSIGNMENT"
    )
    assert live["return_path_assignment_id"] == "assignment-live"
    assert reserve["pool_id"] == "reserve-capital"
    assert reserve["pool_mapping_status"] == "MAPPED"
    assert reserve["strategy_mapping_disposition"] == "RESERVE_NOT_APPLICABLE"


def test_performance_rejects_ambiguous_account_live_path() -> None:
    catalog = {
        "assignments": [
            {
                "account_id": f"account-{suffix}",
                "assignment_id": f"assignment-{suffix}",
                "effective_end": None,
                "effective_start": "2026-01-01",
                "mapping_status": "MAPPED",
                "pool_id": "systematic-capital",
                "pool_mapping_status": "MAPPED",
                "strategy_id": "strategy-live",
                "strategy_mapping_disposition": "STRATEGY_MAPPED",
            }
            for suffix in ("a", "b")
        ],
        "calendar_coverage": [],
        "flows": [],
        "nav_observations": [],
        "strategy_provider_evidence": [
            {
                "account": {"observed_start": "2026-01-01"},
                "evidence_layer": "ACCOUNT_LIVE",
                "generation_id": "a" * 64,
                "strategy_id": "strategy-live",
            }
        ],
        "strategy_returns": [
            {
                "date": "2026-01-01",
                "evidence_layer": "ACCOUNT_LIVE",
                "provider_id": "synthetic",
                "return": "0.01",
                "strategy_id": "strategy-live",
            },
            {
                "date": "2026-01-02",
                "evidence_layer": "ACCOUNT_LIVE",
                "provider_id": "synthetic",
                "return": "0.02",
                "strategy_id": "strategy-live",
            },
        ],
    }

    rows = reporter._account_strategy_performance(catalog)

    assert len(rows) == 2
    assert all("return_path_status" not in row for row in rows)


@pytest.mark.parametrize("trap_surface", ("nav", "return"))
def test_epoch_admission_blocks_twelve_rows_before_any_numeric_candidate(
    trap_surface: str,
) -> None:
    crosswalk = json.loads(
        (
            REPOSITORY_ROOT
            / "contracts"
            / "account-strategy-current-epoch-crosswalk-v1.json"
        ).read_text(encoding="utf-8")
    )
    included = [
        row
        for row in crosswalk["rows"]
        if row["product_epoch_disposition"] == "INCLUDED_PRODUCT_EPOCH"
    ]
    assignments = [
        {
            "account_id": row["account_id"],
            "assignment_id": row["assignment_id"],
            "effective_end": row["effective_end"],
            "effective_start": row["effective_start"],
            "mapping_status": "MAPPED",
            "pool_id": "systematic-capital",
            "pool_mapping_status": "MAPPED",
            "strategy_id": row["strategy_id"],
            "strategy_mapping_disposition": "STRATEGY_MAPPED",
        }
        for row in included
    ]
    nav_observations = [
        {
            "account_id": row["account_id"],
            "as_of_date": day,
            "nav": _amount(value),
        }
        for row in included
        for day, value in (
            ("2026-08-12", "100"),
            ("2026-08-13", "110"),
            ("2026-08-14", "999999"),
        )
    ]
    blocked_probe = next(
        row for row in included if row["strategy_id"] != "volatis-strategy"
    )
    if trap_surface == "nav":
        nav_observations.append(
            {
                "account_id": blocked_probe["account_id"],
                "as_of_date": "TRAP-MUST-NOT-BE-READ",
                "nav": {"status": "AVAILABLE", "value": object()},
            }
        )
    epoch_details = {
        (
            row["account_id"],
            row["strategy_id"],
            row["effective_start"],
            row["effective_end"],
        ): {
            "account_label": row["account_label"],
            "last_four": row["last_four"],
            "epoch_admission_status": (
                "FUNDED" if row["strategy_id"] == "volatis-strategy" else "BLOCKED"
            ),
            "epoch_output_status": (
                "OK"
                if row["strategy_id"] == "volatis-strategy"
                else "BLOCKED_ADMISSION"
            ),
            "epoch_blocking_reason": (
                None
                if row["strategy_id"] == "volatis-strategy"
                else "FLOW_COMPLETENESS_NOT_PROVEN"
            ),
        }
        for row in included
    }
    catalog = {
        "assignments": assignments,
        "calendar_coverage": [],
        "flows": [],
        "nav_observations": nav_observations,
        "strategy_provider_evidence": [
            {
                "account": {
                    "account_scopes": [
                        {
                            "account_id": next(
                                row["account_id"]
                                for row in included
                                if row["strategy_id"] == "volatis-strategy"
                            )
                        }
                    ],
                    "observed_start": "2026-08-12",
                    "observed_end": "2026-08-14",
                    "return_observations": 3,
                },
                "evidence_layer": "ACCOUNT_LIVE",
                "generation_id": "a" * 64,
                "strategy_id": "volatis-strategy",
            }
        ],
        "strategy_returns": [
            {
                "date": day,
                "evidence_layer": "ACCOUNT_LIVE",
                "provider_id": "synthetic",
                "return": value,
                "strategy_id": "volatis-strategy",
            }
            for day, value in (
                ("2026-08-12", "0.01"),
                ("2026-08-13", "0.02"),
                ("2026-08-14", "9.99"),
            )
        ]
        + (
            [
                {
                    "date": "TRAP-MUST-NOT-BE-READ",
                    "evidence_layer": "ACCOUNT_LIVE",
                    "provider_id": "synthetic",
                    "return": object(),
                    "strategy_id": blocked_probe["strategy_id"],
                }
            ]
            if trap_surface == "return"
            else []
        ),
    }
    epoch_tuples = set(epoch_details)
    try:
        rows = reporter._account_strategy_performance(
            catalog,
            None,
            epoch_tuples,
            epoch_details,
            calculation_cutoff_date="2026-08-13",
        )
    except TypeError:
        rows = []

    blocked = [row for row in rows if row["epoch_admission_status"] == "BLOCKED"]
    funded = [row for row in rows if row["epoch_admission_status"] == "FUNDED"]
    assert len(blocked) == 12
    assert len(funded) == 1
    assert {row["strategy_id"] for row in funded} == {"volatis-strategy"}
    assert all(
        row[status_field] == "BLOCKED_ADMISSION"
        for row in blocked
        for status_field in ("twr_status", "mwr_status", "return_path_status")
    )
    suppressed_fields = (
        "twr",
        "mwr_xirr",
        "observed_start",
        "observed_end",
        "flow_count",
        "nav_observations",
        "return_path_cumulative_return",
        "return_path_observations",
        "return_path_observed_start",
        "return_path_observed_end",
        "return_path_series_sha256",
        "return_path_generation_id",
        "return_path_authority",
        "return_path_account_binding",
        "return_path_assignment_id",
    )
    assert all(row[field] is None for row in blocked for field in suppressed_fields)
    assert funded[0]["twr"] == "0.1"
    assert funded[0]["observed_end"] == "2026-08-13"
    assert funded[0]["nav_observations"] == 2
    assert funded[0]["return_path_observations"] == 2
    assert funded[0]["return_path_observed_end"] == "2026-08-13"


def _amount(value: str) -> dict[str, str]:
    return {"currency": "USD", "status": "AVAILABLE", "value": value}


def _baseline_catalog() -> dict[str, Any]:
    strategies = [
        {
            "aliases": [f"Synthetic Strategy {index}"],
            "cost_status": "KNOWN" if index == 1 else "UNAVAILABLE",
            "coverage_status": "PARTIAL",
            "economic_family": f"SYNTHETIC FAMILY {index}",
            "evidence_routes": [
                {
                    "evidence_layer": "ACCOUNT_LIVE",
                    "purpose": "SYNTHETIC evidence",
                    "source_id": "synthetic-source",
                }
            ],
            "gaps": ["SYNTHETIC gap"],
            "name": f"Synthetic Strategy {index}",
            "strategy_id": f"strategy-{index}",
            "version_ref": None,
        }
        for index in range(1, 10)
    ]
    assignments = [
        {
            "account_id": "account-a",
            "assignment_id": "assignment-a",
            "effective_end": "2026-01-03",
            "effective_start": "2026-01-01",
            "mapping_status": "MAPPED",
            "observed_label": "Synthetic Strategy 1",
            "source_classification": "SYNTHETIC",
            "source_id": "synthetic-source",
            "strategy_id": "strategy-1",
        },
        {
            "account_id": "account-b",
            "assignment_id": "assignment-b",
            "effective_end": None,
            "effective_start": "2026-01-01",
            "mapping_status": "MAPPED",
            "observed_label": "Synthetic Strategy 2",
            "source_classification": "SYNTHETIC",
            "source_id": "synthetic-source",
            "strategy_id": "strategy-2",
        },
        {
            "account_id": "account-b",
            "assignment_id": "assignment-general",
            "effective_end": "2025-12-31",
            "effective_start": None,
            "mapping_status": "UNMAPPED",
            "observed_label": "General",
            "source_classification": "SYNTHETIC",
            "source_id": "synthetic-source",
            "strategy_id": None,
        },
    ]
    nav = [
        {
            "account_id": "account-a",
            "as_of_date": "2026-01-01",
            "nav": _amount("100"),
            "provider": "SYNTHETIC",
            "record_id": "nav-1",
            "source_key": "synthetic-nav",
        },
        {
            "account_id": "account-a",
            "as_of_date": "2026-01-02",
            "nav": _amount("115"),
            "provider": "SYNTHETIC",
            "record_id": "nav-2",
            "source_key": "synthetic-nav",
        },
        {
            "account_id": "account-a",
            "as_of_date": "2026-01-03",
            "nav": _amount("121"),
            "provider": "SYNTHETIC",
            "record_id": "nav-3",
            "source_key": "synthetic-nav",
        },
    ]
    allocations: list[dict[str, Any]] = [
        {
            "allocation_id": f"allocation-{index}",
            "as_of_date": "2026-01-03",
            "current_value": _amount("10"),
            "scope": "SYSTEMATIC",
            "source_key": "synthetic-allocation",
            "strategy_id": f"strategy-{index}",
            "systematic_weight": str(Decimal("1") / Decimal("9")),
        }
        for index in range(1, 10)
    ]
    allocations.append(
        {
            "account_id": "reserve",
            "allocation_id": "allocation-reserve",
            "as_of_date": "2026-01-03",
            "current_value": _amount("10"),
            "scope": "RESERVE",
            "source_key": "synthetic-allocation",
            "strategy_id": None,
            "systematic_weight": None,
        }
    )
    return {
        "accounts": [
            {
                "account_id": "account-a",
                "base_currency": "USD",
                "provider": "SYNTHETIC",
                "role": "INVESTMENT",
            },
            {
                "account_id": "account-b",
                "base_currency": "USD",
                "provider": "SYNTHETIC",
                "role": "INVESTMENT",
            },
            {
                "account_id": "reserve",
                "base_currency": "USD",
                "provider": "INTERNAL",
                "role": "RESERVE",
            },
        ],
        "allocations": allocations,
        "as_of_date": "2026-01-03",
        "assignments": assignments,
        "assumptions": [],
        "baseline_schema_version": 1,
        "calendar_coverage": [
            {
                "account_id": "account-a",
                "date": day,
                "provider": "SYNTHETIC",
                "status": "OBSERVED",
            }
            for day in ("2026-01-01", "2026-01-02", "2026-01-03")
        ]
        + [
            {
                "account_id": "account-b",
                "date": None,
                "provider": "SYNTHETIC",
                "status": "NO_ACCOUNT_NAV_OBSERVATIONS",
            }
        ],
        "corrections": [
            {
                "correction_id": "correction-1",
                "fields": {
                    "account": "SYNTHETIC ACCOUNT",
                    "evidence": "SYNTHETIC",
                    "field": "start_date",
                    "institution": "SYNTHETIC",
                    "new_value": "2026-01-01",
                    "old_value": "2025-12-31",
                    "strategy": "Synthetic Strategy 1",
                },
                "source_id": "synthetic-source",
            }
        ],
        "costs": [
            {
                "annual_external_cost_usd": "1" if index == 1 else None,
                "cost_status": "known" if index == 1 else "unknown",
                "strategy_id": f"strategy-{index}",
            }
            for index in range(1, 10)
        ],
        "evidence_layer_metrics": [],
        "evidence_links": [],
        "fact_coverage": [
            {
                "fact_type": "accounts",
                "known_gaps": [],
                "status": "COMPLETE_WITHIN_DECLARED_BOUNDARY",
            },
            {
                "fact_type": "nav",
                "known_gaps": ["SYNTHETIC partial"],
                "status": "PARTIAL",
            },
            {
                "fact_type": "positions",
                "known_gaps": ["SYNTHETIC unavailable"],
                "status": "FAILED",
            },
        ],
        "flows": [
            {
                "account_id": "account-a",
                "amount": _amount("10"),
                "event_date": "2026-01-02",
                "evidence_class": "OBSERVED",
                "flow_type": "EXTERNAL_FLOW",
                "mapping_status": "MAPPED",
                "provider": "SYNTHETIC",
                "record_id": "flow-1",
                "source_key": "synthetic-flow",
                "strategy_id": "strategy-1",
            }
        ],
        "implementation_gaps": [
            {
                "account_layer": "SYNTHETIC account",
                "metrics": {},
                "note": "SYNTHETIC",
                "provider_layer": "SYNTHETIC provider",
                "status": "partial",
                "strategy_id": f"strategy-{index}",
            }
            for index in range(1, 10)
        ],
        "nav_observations": nav,
        "policy_version": "PORT-PERF-1",
        "positions": [],
        "provenance": {"catalog_sha256": "1" * 64, "generation_id": "2" * 64},
        "reconciliation_checks": [
            {"check_id": "nav", "status": "PASS"},
            {"check_id": "allocation", "status": "PASS"},
            {"check_id": "total", "status": "PASS"},
        ],
        "reconciliation_exceptions": [
            {
                "exception_id": "positions-unavailable",
                "fact_type": "positions",
                "severity": "BLOCKING",
                "status": "UNAVAILABLE",
            }
        ],
        "replay": {
            "correlation_one_stress": "0.37",
            "full_systematic_hwm_status": "NOT_AUTOMATABLE_MISSING_COMPLETE_DAILY_NAV",
            "nav_coverage": "0.5",
            "portfolio_action": "DO_NOT_EXPAND_SYSTEMATIC_RISK",
            "reserve_total_usd": "10",
            "status": "PASS",
            "systematic_total_usd": "90",
            "total_liquid_usd": "100",
        },
        "securities": [
            {"currency": "USD", "security_id": "security-1", "symbol": "SYNTH"}
        ],
        "source_artifacts": [
            {
                "artifact_id": "source-1",
                "as_of_date": "2026-01-03",
                "schema_name": "synthetic",
                "schema_version": "1",
                "sha256": "3" * 64,
                "source_key": "synthetic",
            }
        ],
        "strategies": strategies,
        "strategy_live_metrics": [
            {
                "evidence_layer": "ACCOUNT_LIVE",
                "metrics": {
                    "cagr": "0.1",
                    "cumulative_return": "0.05",
                    "max_drawdown": "-0.02",
                },
                "observed_label": "Synthetic Strategy 1",
                "strategy_id": "strategy-1",
            }
        ],
        "strategy_nav_coverage": [
            {
                "current_value_usd": "10",
                "daily_end": "2026-01-03" if index == 1 else None,
                "daily_start": "2026-01-01" if index == 1 else None,
                "known_gap": None if index == 1 else "SYNTHETIC missing NAV",
                "return_observations": "3" if index == 1 else None,
                "status": "AVAILABLE" if index == 1 else "UNAVAILABLE",
                "strategy_id": f"strategy-{index}",
            }
            for index in range(1, 10)
        ],
        "transactions": [
            {
                "account_id": "account-a",
                "amount": _amount("-5"),
                "event_date": "2026-01-02",
                "kind": "BUY_FILL",
                "mapping_status": "MAPPED",
                "price": "5",
                "provider": "SYNTHETIC",
                "quantity": "1",
                "record_id": "transaction-1",
                "security_id": "security-1",
                "source_key": "synthetic-transaction",
                "strategy_id": "strategy-1",
            }
        ],
        "unmapped_records": [assignments[-1]],
    }


def _published_baseline(
    tmp_path: Path,
    catalog: dict[str, Any] | None = None,
    *,
    product_contract_path: Path | None = None,
) -> Path:
    """Publish a synthetic baseline after explicitly binding test controls."""

    root = tmp_path / "baseline"
    catalog = _baseline_catalog() if catalog is None else catalog
    prior_contract_path = reporter._PORTABLE_HWM_PRODUCT_CONTRACT
    if product_contract_path is not None:
        reporter._PORTABLE_HWM_PRODUCT_CONTRACT = product_contract_path
    elif (
        reporter._PORTABLE_HWM_PRODUCT_CONTRACT == DEFAULT_PRODUCT_CONTRACT
        or reporter._PORTABLE_HWM_PRODUCT_CONTRACT in _SYNTHETIC_PRODUCT_CONTRACT_PATHS
        and any(
            str(row.get("strategy_id", "")).startswith("strategy-")
            for row in catalog.get("strategies", [])
            if isinstance(row, dict)
        )
    ):
        # Legacy synthetic catalogs need an explicit test-only contract, but a
        # caller's monkeypatch is authoritative and must not be overwritten.
        reporter._PORTABLE_HWM_PRODUCT_CONTRACT = _bind_synthetic_product_contract(
            tmp_path
        )
    content = canonical_json_bytes(catalog)
    artifact = SourceArtifact(
        artifact_id="normalized/portfolio-baseline.json",
        sha256=hashlib.sha256(content).hexdigest(),
        size_bytes=len(content),
        as_of_date=date(2026, 1, 3),
        schema_name="portfolio-baseline-catalog",
        schema_version="1",
    )
    manifest = GenerationManifest.build(
        as_of_date=date(2026, 1, 3),
        adapter=AdapterIdentity("synthetic-baseline", "1"),
        coverage=CoverageClaim(
            status=CoverageStatus.COMPLETE_WITHIN_DECLARED_BOUNDARY,
            boundary="SYNTHETIC baseline",
            fact_types=("portfolio",),
            observed_start=date(2026, 1, 1),
            observed_end=date(2026, 1, 3),
        ),
        artifacts=(artifact,),
        verification=VerificationStatus.VERIFIED,
    )
    manifest_path = stage_generation(root, manifest)
    (manifest_path.parent / "portfolio-baseline.json").write_bytes(content)
    publish_generation(root, manifest.generation_id)
    if product_contract_path is not None:
        # An explicit per-test contract must not leak into later synthetic
        # fixtures; callers that need it can bind it through monkeypatch.
        reporter._PORTABLE_HWM_PRODUCT_CONTRACT = prior_contract_path
    return root


def _published_dependence(
    tmp_path: Path,
    *,
    cutoff: str | None = None,
    cutoff_mutation: tuple[str, str] | None = None,
) -> Path:
    root = tmp_path / "dependence"
    series_ids = [f"Current Strategy {index}" for index in range(1, 8)]
    pairs = [
        (left, right)
        for left_index, left in enumerate(series_ids)
        for right in series_ids[left_index + 1 :]
    ]

    def mode_payload() -> dict[str, Any]:
        return {
            "series_ids": series_ids,
            "observed_cells": [
                {
                    "left": left,
                    "right": right,
                    "status": "AVAILABLE",
                    "coverage": {
                        "common_observations": 2,
                        "observed_start": "2026-01-02",
                        "observed_end": "2026-01-03",
                    },
                    "pearson": "0.1",
                    "spearman": "0.2",
                    "horizons": {},
                    "rolling_pearson": {},
                    "downside": {},
                    "drawdown_regimes": {},
                    "joint_tail": {},
                }
                for left, right in pairs
            ],
            "blend_cells": [
                {
                    "left": left,
                    "right": right,
                    "decision_correlation": "0.3",
                    "decision_status": "POSTERIOR_STRESS_PSD",
                }
                for left, right in pairs
            ],
            "implementation_cells": [
                {
                    "left": left,
                    "right": right,
                    "mapping_scope": "CENSUS_PROVIDER_LANE_ONLY",
                    "shared_providers": [],
                    "status": "AVAILABLE_PROVIDER_LANE_CONCENTRATION",
                    "unassessed_dimensions": ["INSTRUMENT", "VENUE"],
                }
                for left, right in pairs
            ],
            "central_case": {},
            "correlation_one": {"loss": "0.31"},
            "drawdown_matrices": {},
            "horizon_matrices": {},
            "matrices": {},
            "performance": (
                {}
                if cutoff is None
                else {series_id: {"observed_end": cutoff} for series_id in series_ids}
            ),
            "series_composition": (
                {}
                if cutoff is None
                else {series_id: {"observed_end": cutoff} for series_id in series_ids}
            ),
        }

    live = mode_payload()
    hybrid = mode_payload()
    correlation = {
        "analytics_hash": "a" * 64,
        "analytics_schema_version": 2,
        "as_of_date": "2026-01-03",
        "builder_version": "1.6.0",
        "default_evidence_mode": "LIVE_ONLY",
        "dependence": live,
        "dependence_modes": {
            "LIVE_ONLY": live,
            "LIVE_PLUS_BACKTEST": hybrid,
        },
        "source": {
            "artifact_id": "current-dependence/returns.csv",
            "sha256": "b" * 64,
            "size_bytes": 100,
        },
    }
    if cutoff is not None:
        correlation["complete_through"] = cutoff
    if cutoff_mutation is not None:
        target, replacement = cutoff_mutation
        if target == "complete_through":
            correlation[target] = replacement
        else:
            correlation["dependence_modes"]["LIVE_ONLY"][target]["Current Strategy 1"][
                "observed_end"
            ] = replacement
            correlation["dependence"] = correlation["dependence_modes"]["LIVE_ONLY"]
    content = canonical_json_bytes(correlation)
    artifact = SourceArtifact(
        artifact_id="normalized/correlation.json",
        sha256=hashlib.sha256(content).hexdigest(),
        size_bytes=len(content),
        as_of_date=(
            date.fromisoformat(cutoff) if cutoff is not None else date(2026, 1, 3)
        ),
        schema_name="portfolio-correlation-analysis",
        schema_version="2",
    )
    manifest = GenerationManifest.build(
        as_of_date=(
            date.fromisoformat(cutoff) if cutoff is not None else date(2026, 1, 3)
        ),
        adapter=AdapterIdentity("portfolio-correlation", "1.6.0"),
        coverage=CoverageClaim(
            status=CoverageStatus.COMPLETE_WITHIN_DECLARED_BOUNDARY,
            boundary="SYNTHETIC current dependence",
            fact_types=("dependence",),
            observed_start=date(2026, 1, 2),
            observed_end=(
                date.fromisoformat(cutoff) if cutoff is not None else date(2026, 1, 3)
            ),
        ),
        artifacts=(artifact,),
        verification=VerificationStatus.VERIFIED,
    )
    manifest_path = stage_generation(root, manifest)
    (manifest_path.parent / "correlation.json").write_bytes(content)
    publish_generation(root, manifest.generation_id)
    return root


def test_current_dependence_reaches_report_workbook_and_receipt_identity(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    baseline_root = _published_baseline(tmp_path)
    dependence_root = _published_dependence(tmp_path)
    history_path = tmp_path / "dependence-source.json"
    monkeypatch.setattr(
        reporter,
        "_history_contract",
        lambda *_args, **_kwargs: (
            {},
            {
                "age_evidence": {
                    "calculation_age_days": 0,
                    "complete_through": "2026-01-03",
                    "evaluated_on": "2026-01-03",
                    "limiting_inputs": [
                        {
                            "observed_end": "2026-01-03",
                            "publication_lag_days": 0,
                            "segment_id": "synthetic-segment",
                            "series_id": "Synthetic",
                            "source_generation_id": "a" * 64,
                            "source_id": "synthetic-source",
                            "strategy_id": "synthetic-strategy",
                            "warning_after_days": 2,
                        }
                    ],
                    "status": "CURRENT_WITHIN_AGE_THRESHOLD",
                    "warning_after_days": 2,
                },
                "as_of_date": "2026-01-03",
                "bundle_id": "d" * 64,
                "evaluated_on": "2026-01-03",
                "evidence_mode_policy": {},
                "returns_sha256": "b" * 64,
                "returns_size_bytes": 100,
                "schema": "portfolio-current-dependence-source/1",
                "series_count": 9,
                "sha256": "c" * 64,
                "size_bytes": 100,
            },
        ),
    )
    built = reporter.build_candidate(
        baseline_root,
        tmp_path / "reporting",
        history_contract_path=history_path,
        correlation_root=dependence_root,
    )
    candidate = Path(built["candidate_dir"])
    report = json.loads(
        (candidate / reporter.REPORT_DATA_NAME).read_text(encoding="utf-8")
    )
    assert report["dependence_analysis"]["status"] == "AVAILABLE"
    assert report["dependence_analysis"]["pair_count"] == 21
    assert report["dependence_analysis"]["pair_mode_count"] == 42
    assert len(report["dependence_analysis"]["series_ids"]) == 7
    assert len(report["dependence_rows"]) == 42
    correlation_one_summary = next(
        row
        for row in report["portfolio_summary"]
        if row["metric"] == "Correlation-one stress"
    )
    assert correlation_one_summary["value"] == "0.31"
    assert correlation_one_summary["value"] != "0.37"
    workbook = load_workbook(Path(built["workbook_path"]), data_only=True)
    assert workbook["Dependence"].max_row == 47
    assert "tblDependence" in workbook["Dependence"].tables
    verified = reporter.verify_candidate(
        baseline_root,
        candidate,
        Path(built["workbook_path"]),
        history_contract_path=history_path,
        correlation_root=dependence_root,
    )
    assert verified["table_count"] == 14


def test_current_dependence_report_rejects_mismatched_verified_source(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    baseline_root = _published_baseline(tmp_path)
    dependence_root = _published_dependence(tmp_path)
    monkeypatch.setattr(
        reporter,
        "_history_contract",
        lambda *_args, **_kwargs: (
            {},
            {
                "returns_sha256": "0" * 64,
                "returns_size_bytes": 100,
                "schema": "portfolio-current-dependence-source/1",
            },
        ),
    )
    with pytest.raises(
        reporter.ReportValidationError,
        match="dependence source differs from verified history contract",
    ):
        reporter.build_report_data(
            baseline_root,
            history_contract_path=tmp_path / "dependence-source.json",
            correlation_root=dependence_root,
        )


def test_report_uses_only_verified_epoch_hpr_projection_for_endpoint_wealth(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The current-dependence bridge may change endpoint wealth but never
    creates daily return observations or an HWM path through the source gap."""

    daily_rows = [
        {"date": "2026-01-01", "return": "0"},
        {"date": "2026-01-04", "return": "0"},
    ]
    catalog = _baseline_catalog()
    catalog["strategy_returns"] = [
        {
            "date": row["date"],
            "evidence_layer": "ACCOUNT_LIVE",
            "provider_id": "matrixport",
            "return": row["return"],
            "strategy_id": "strategy-1",
        }
        for row in daily_rows
    ]
    catalog["strategy_provider_evidence"] = [
        {
            "account": {
                "cumulative_return": "0",
                "current_drawdown": "0",
                "max_drawdown": "-0.8",
                "maximum_drawdown": "0",
                "observed_end": "2026-01-04",
                "observed_start": "2026-01-01",
                "provider_id": "matrixport",
                "return_observations": 2,
            },
            "evidence_layer": "ACCOUNT_LIVE",
            "strategy_id": "strategy-1",
        }
    ]
    baseline_root = _published_baseline(tmp_path / "baseline", catalog)
    dependence_root = _published_dependence(tmp_path / "dependence")
    receipt = {
        "evidence_layer": "ACCOUNT_LIVE",
        "receipt": {
            "admission_cutoff": "2026-01-04",
            "daily_cumulative_return": "0",
            "daily_return_observations": 2,
            "daily_return_series_schema": "ordered-date-return/1",
            "daily_return_series_sha256": hashlib.sha256(
                canonical_json_bytes(daily_rows)
            ).hexdigest(),
            "epoch_cumulative_return": "-0.0123875290080184841010757599",
            "epoch_holding_period_link_observations": 3,
            "observed_end": "2026-01-04",
            "observed_start": "2026-01-01",
            "source_gap_policy_sha256": "a" * 64,
            "span_return_receipts": [
                {
                    "end_endpoint_sha256": "b" * 64,
                    "missing_source_date": "2026-01-02",
                    "return": "-0.0123875290080184841010757599",
                    "span_days": 2,
                    "span_end_date": "2026-01-03",
                    "span_start_date": "2026-01-01",
                    "start_endpoint_sha256": "c" * 64,
                    "subscription_events_in_span": 0,
                }
            ],
        },
        "source_generation_id": "d" * 64,
        "source_id": "matrixport-account-generation",
        "strategy_id": "strategy-1",
    }
    monkeypatch.setattr(
        reporter,
        "_history_contract",
        lambda *_args, **_kwargs: (
            {},
            {
                "age_evidence": {
                    "calculation_age_days": 0,
                    "complete_through": "2026-01-03",
                    "evaluated_on": "2026-01-03",
                    "limiting_inputs": [],
                    "status": "CURRENT_WITHIN_AGE_THRESHOLD",
                    "warning_after_days": 2,
                },
                "as_of_date": "2026-01-03",
                "bundle_id": "d" * 64,
                "evaluated_on": "2026-01-03",
                "epoch_hpr_receipts": [receipt],
                "evidence_mode_policy": {},
                "returns_sha256": "b" * 64,
                "returns_size_bytes": 100,
                "schema": "portfolio-current-dependence-source/1",
                "series_count": 9,
                "sha256": "c" * 64,
                "size_bytes": 100,
            },
        ),
    )

    report, _, _ = reporter.build_report_data(
        baseline_root,
        history_contract_path=tmp_path / "dependence-source.json",
        correlation_root=dependence_root,
    )

    strategy = next(
        row for row in report["strategy_summary"] if row["strategy_id"] == "strategy-1"
    )
    assert (
        strategy["account_live_cumulative_return"] == "-0.0123875290080184841010757599"
    )
    assert strategy["account_live_return_observations"] == 2
    assert strategy["account_live_max_drawdown"] is None
    assert strategy["hwm_current_drawdown"] is None
    assert strategy["hwm_max_drawdown"] is None
    assert strategy["hwm_path_metrics_status"] == "UNAVAILABLE"
    assert strategy["hwm_path_metrics_reason"] == "MISSING_DAILY_POINT_CAN_HIDE_PEAK"
    assert report["hwm_summary"]["hwm_numerical_status"] == "UNAVAILABLE"
    assert (
        report["hwm_summary"]["hwm_numerical_reason"]
        == "MISSING_DAILY_POINT_CAN_HIDE_PEAK"
    )
    assert len(report["dependence_rows"]) == 42


def test_current_dependence_report_rejects_coherently_unbound_output(
    tmp_path: Path,
) -> None:
    dependence_root = _published_dependence(tmp_path)
    pointer = read_current(dependence_root)
    path = dependence_root / "generations" / pointer.generation_id / "correlation.json"
    payload = json.loads(path.read_text(encoding="utf-8"))
    payload["analytics_hash"] = "f" * 64
    path.write_bytes(canonical_json_bytes(payload))
    with pytest.raises(
        reporter.ReportValidationError, match="generation is not accepted"
    ):
        reporter._dependence_report(dependence_root)


def test_baseline_retries_transient_pointer_io(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    root = _published_baseline(tmp_path)
    attempts = 0

    def flaky_read_current(path: Path) -> object:
        nonlocal attempts
        attempts += 1
        if attempts < 3:
            raise PermissionError("transient provider lock")
        return read_current(path)

    monkeypatch.setattr(reporter, "read_current", flaky_read_current)
    monkeypatch.setattr(reporter.time_module, "sleep", lambda _seconds: None)

    catalog, identity, _path = reporter._baseline(root)

    assert attempts == 3
    assert identity["generation_id"] == read_current(root).generation_id
    assert catalog["baseline_schema_version"] == 1


def test_explicit_baseline_generation_ignores_stale_pointer_and_verifies_bytes(
    tmp_path: Path,
) -> None:
    root = _published_baseline(tmp_path)
    selected_id = read_current(root).generation_id
    pointer_path = root / "CURRENT.json"
    stale_pointer = json.loads(pointer_path.read_text(encoding="utf-8"))
    stale_pointer["generation_id"] = "f" * 64
    pointer_path.write_bytes(canonical_json_bytes(stale_pointer))
    stale_bytes = pointer_path.read_bytes()
    try:
        catalog, identity, catalog_path = reporter._baseline(
            root, generation_id=selected_id
        )
    except TypeError:
        catalog, identity, catalog_path = {}, {}, Path()

    assert catalog["baseline_schema_version"] == 1
    assert identity["generation_id"] == selected_id
    assert pointer_path.read_bytes() == stale_bytes
    catalog_path.write_bytes(catalog_path.read_bytes() + b" ")
    with pytest.raises(reporter.ReportValidationError, match="artifact identity"):
        reporter._baseline(root, generation_id=selected_id)


@pytest.mark.parametrize(
    ("baseline_root", "baseline_generation_id"),
    ((Path("missing-baseline"), None), (None, "a" * 64)),
)
def test_bind_workbook_evidence_rejects_half_baseline_selector(
    baseline_root: Path | None, baseline_generation_id: str | None
) -> None:
    with pytest.raises(
        reporter.ReportValidationError,
        match="baseline root and generation identity are required together",
    ):
        reporter.bind_workbook_evidence(
            Path("missing-candidate"),
            Path("missing-workbook"),
            Path("missing-xlsx-result"),
            Path("missing-validation"),
            baseline_root=baseline_root,
            baseline_generation_id=baseline_generation_id,
        )


@pytest.mark.parametrize(
    "selector_args",
    (("--baseline-root", "missing"), ("--baseline-generation-id", "a" * 64)),
)
def test_bind_workbook_evidence_cli_rejects_half_baseline_selector(
    selector_args: tuple[str, str], capsys: pytest.CaptureFixture[str]
) -> None:
    exit_code = reporter.main(
        [
            "bind-workbook-evidence",
            "--candidate-dir",
            "missing-candidate",
            "--excel-workbook",
            "missing-workbook",
            "--xlsx-result",
            "missing-xlsx-result",
            "--validation-result",
            "missing-validation",
            "--evidence",
            "missing-evidence",
            *selector_args,
        ]
    )

    captured = capsys.readouterr()
    assert exit_code == 1
    assert "baseline root and generation identity are required together" in captured.err


def test_explicit_correlation_generation_ignores_stale_pointer_and_enforces_cutoff(
    tmp_path: Path,
) -> None:
    root = _published_dependence(tmp_path, cutoff="2026-08-13")
    selected_id = read_current(root).generation_id
    generation_dir = root / "generations" / selected_id
    pointer_path = root / "CURRENT.json"
    stale_pointer = json.loads(pointer_path.read_text(encoding="utf-8"))
    stale_pointer["generation_id"] = "e" * 64
    pointer_path.write_bytes(canonical_json_bytes(stale_pointer))
    stale_bytes = pointer_path.read_bytes()
    try:
        projection, rows, identity, correlation_path = reporter._dependence_report(
            root,
            generation_dir=generation_dir,
            calculation_cutoff_date="2026-08-13",
        )
    except TypeError:
        projection, rows, identity, correlation_path = {}, [], {}, Path()

    assert projection["complete_through"] == "2026-08-13"
    assert len(rows) == 42
    assert identity["generation_id"] == selected_id
    assert pointer_path.read_bytes() == stale_bytes
    correlation_path.write_bytes(correlation_path.read_bytes() + b" ")
    with pytest.raises(reporter.ReportValidationError, match="artifact identity"):
        reporter._dependence_report(
            root,
            generation_dir=generation_dir,
            calculation_cutoff_date="2026-08-13",
        )


def _staged_selector_fixture(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> tuple[Path, str, bytes, Path, Path, bytes, Path]:
    baseline_root = _published_baseline(tmp_path / "baseline")
    baseline_generation_id = read_current(baseline_root).generation_id
    baseline_pointer = baseline_root / "CURRENT.json"
    baseline_stale = json.loads(baseline_pointer.read_text(encoding="utf-8"))
    baseline_stale["generation_id"] = "f" * 64
    baseline_pointer.write_bytes(canonical_json_bytes(baseline_stale))
    baseline_pointer_bytes = baseline_pointer.read_bytes()

    correlation_root = _published_dependence(tmp_path / "correlation")
    correlation_generation_id = read_current(correlation_root).generation_id
    correlation_generation_dir = (
        correlation_root / "generations" / correlation_generation_id
    )
    correlation_pointer = correlation_root / "CURRENT.json"
    correlation_stale = json.loads(correlation_pointer.read_text(encoding="utf-8"))
    correlation_stale["generation_id"] = "e" * 64
    correlation_pointer.write_bytes(canonical_json_bytes(correlation_stale))
    correlation_pointer_bytes = correlation_pointer.read_bytes()

    history_path = tmp_path / "history" / "dependence-source.json"
    history_path.parent.mkdir()
    history_path.write_bytes(canonical_json_bytes({"synthetic": True}))
    monkeypatch.setattr(
        reporter,
        "_history_contract",
        lambda *_args, **_kwargs: (
            {},
            {
                "as_of_date": "2026-01-03",
                "bundle_id": "d" * 64,
                "evaluated_on": "2026-01-03",
                "evidence_mode_policy": {},
                "returns_sha256": "b" * 64,
                "returns_size_bytes": 100,
                "schema": "portfolio-current-dependence-source/1",
                "series_count": 9,
                "sha256": "c" * 64,
                "size_bytes": 100,
            },
        ),
    )
    return (
        baseline_root,
        baseline_generation_id,
        baseline_pointer_bytes,
        correlation_root,
        correlation_generation_dir,
        correlation_pointer_bytes,
        history_path,
    )


def test_staged_selectors_flow_through_build_verify_bind_publish_and_inspect(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    (
        baseline_root,
        baseline_generation_id,
        baseline_pointer_bytes,
        correlation_root,
        correlation_generation_dir,
        correlation_pointer_bytes,
        history_path,
    ) = _staged_selector_fixture(tmp_path, monkeypatch)
    built = reporter.build_candidate(
        baseline_root,
        tmp_path / "candidates",
        history_contract_path=history_path,
        correlation_root=correlation_root,
        baseline_generation_id=baseline_generation_id,
        correlation_generation_dir=correlation_generation_dir,
    )
    candidate = Path(built["candidate_dir"])
    report = json.loads(
        (candidate / reporter.REPORT_DATA_NAME).read_text(encoding="utf-8")
    )
    assert report["baseline"]["generation_id"] == baseline_generation_id
    assert (
        report["dependence_analysis"]["identity"]["generation_id"]
        == correlation_generation_dir.name
    )
    verified = reporter.verify_candidate(
        baseline_root,
        candidate,
        Path(built["workbook_path"]),
        history_contract_path=history_path,
        correlation_root=correlation_root,
        baseline_generation_id=baseline_generation_id,
        correlation_generation_dir=correlation_generation_dir,
    )
    assert verified["baseline_generation_id"] == baseline_generation_id
    workbook, run_result, validation, binding_path = _prepare_workbook_evidence(
        candidate
    )
    binding = reporter.bind_workbook_evidence(
        candidate,
        workbook,
        run_result,
        validation,
        binding_path,
        baseline_root=baseline_root,
        baseline_generation_id=baseline_generation_id,
    )
    assert binding["workbook_evidence_status"] == (
        "WORKBOOK_DELIVERY_EVIDENCE_ACCEPTED"
    )
    published_root = tmp_path / "published"
    published = reporter.publish_report(
        baseline_root,
        candidate,
        workbook,
        run_result,
        validation,
        published_root,
        history_contract_path=history_path,
        correlation_root=correlation_root,
        workbook_evidence_binding=binding_path,
        baseline_generation_id=baseline_generation_id,
        correlation_generation_dir=correlation_generation_dir,
    )
    generation_id = published["generation_id"]
    delivery = published_root / "deliveries" / generation_id
    inspected = reporter.inspect_report_delivery(
        published_root, generation_id, delivery
    )
    assert inspected["workbook_status"] == "EXACT_NON_AUTHORITY"
    published_report = json.loads(
        (
            published_root / "generations" / generation_id / reporter.REPORT_DATA_NAME
        ).read_text(encoding="utf-8")
    )
    assert published_report["baseline"]["generation_id"] == baseline_generation_id
    assert (
        published_report["dependence_analysis"]["identity"]["generation_id"]
        == correlation_generation_dir.name
    )
    assert (baseline_root / "CURRENT.json").read_bytes() == baseline_pointer_bytes
    assert (correlation_root / "CURRENT.json").read_bytes() == correlation_pointer_bytes


def test_staged_selectors_flow_through_real_cli_lifecycle(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    (
        baseline_root,
        baseline_generation_id,
        baseline_pointer_bytes,
        correlation_root,
        correlation_generation_dir,
        correlation_pointer_bytes,
        history_path,
    ) = _staged_selector_fixture(tmp_path, monkeypatch)
    candidate_root = tmp_path / "cli-candidates"
    common = [
        "--baseline-root",
        str(baseline_root),
        "--baseline-generation-id",
        baseline_generation_id,
        "--history-contract",
        str(history_path),
        "--correlation-root",
        str(correlation_root),
        "--correlation-generation-dir",
        str(correlation_generation_dir),
    ]
    assert reporter.main(["build", *common, "--output-root", str(candidate_root)]) == 0
    built = json.loads(capsys.readouterr().out)
    candidate = Path(built["candidate_dir"])
    workbook = Path(built["workbook_path"])
    assert (
        reporter.main(
            [
                "verify",
                *common,
                "--candidate-dir",
                str(candidate),
                "--excel-workbook",
                str(workbook),
            ]
        )
        == 0
    )
    assert json.loads(capsys.readouterr().out)["table_count"] == 14
    final_workbook, run_result, validation, binding_path = _prepare_workbook_evidence(
        candidate, run_id="4" * 32
    )
    assert (
        reporter.main(
            [
                "bind-workbook-evidence",
                "--candidate-dir",
                str(candidate),
                "--excel-workbook",
                str(final_workbook),
                "--xlsx-result",
                str(run_result),
                "--validation-result",
                str(validation),
                "--evidence",
                str(binding_path),
                "--baseline-root",
                str(baseline_root),
                "--baseline-generation-id",
                baseline_generation_id,
            ]
        )
        == 0
    )
    assert json.loads(capsys.readouterr().out)["workbook_evidence_status"] == (
        "WORKBOOK_DELIVERY_EVIDENCE_ACCEPTED"
    )
    published_root = tmp_path / "cli-published"
    assert (
        reporter.main(
            [
                "publish",
                *common,
                "--candidate-dir",
                str(candidate),
                "--excel-workbook",
                str(final_workbook),
                "--xlsx-result",
                str(run_result),
                "--validation-result",
                str(validation),
                "--workbook-evidence-binding",
                str(binding_path),
                "--output-root",
                str(published_root),
            ]
        )
        == 0
    )
    published = json.loads(capsys.readouterr().out)
    generation_id = published["generation_id"]
    inspected = reporter.inspect_report_delivery(
        published_root,
        generation_id,
        published_root / "deliveries" / generation_id,
    )
    assert inspected["workbook_status"] == "EXACT_NON_AUTHORITY"
    assert (baseline_root / "CURRENT.json").read_bytes() == baseline_pointer_bytes
    assert (correlation_root / "CURRENT.json").read_bytes() == correlation_pointer_bytes


@pytest.mark.parametrize(
    ("target", "replacement"),
    (
        ("complete_through", "2026-08-12"),
        ("series_composition", "2026-08-12"),
        ("performance", "2026-08-12"),
    ),
)
def test_correlation_cutoff_mutation_in_any_layer_fails(
    tmp_path: Path, target: str, replacement: str
) -> None:
    root = _published_dependence(
        tmp_path,
        cutoff="2026-08-13",
        cutoff_mutation=(target, replacement),
    )
    selected_id = read_current(root).generation_id
    generation_dir = root / "generations" / selected_id

    with pytest.raises(reporter.ReportValidationError, match="cutoff"):
        reporter._dependence_report(
            root,
            generation_dir=generation_dir,
            calculation_cutoff_date="2026-08-13",
        )


@pytest.mark.parametrize(
    ("surface", "payload"),
    (
        (
            "provider_strategy_returns",
            {
                "provider_strategy_returns": [
                    {"date": "2026-08-14", "strategy_id": "probe"}
                ]
            },
        ),
    ),
)
def test_calculation_cutoff_clips_post_cutoff_row_collections(
    surface: str, payload: dict[str, Any]
) -> None:
    projected = reporter._calculation_catalog_through(payload, "2026-08-13")

    assert projected[surface] == []


def test_calculation_cutoff_retains_current_allocation_snapshot() -> None:
    allocation = {
        "as_of_date": "2026-08-16",
        "current_value": "100",
        "strategy_id": "strategy-current",
    }
    projected = reporter._calculation_catalog_through(
        {
            "allocations": [allocation],
            "provider_strategy_returns": [
                {"date": "2026-08-14", "strategy_id": "strategy-current"}
            ],
        },
        "2026-08-13",
    )

    assert projected["allocations"] == [allocation]
    assert projected["provider_strategy_returns"] == []


def test_calculation_cutoff_retains_null_calendar_sentinel_and_clips_dated_rows() -> (
    None
):
    sentinel = {
        "account_id": "account-no-nav",
        "date": None,
        "provider": "SYNTHETIC",
        "status": "NO_ACCOUNT_NAV_OBSERVATIONS",
    }
    projected = reporter._calculation_catalog_through(
        {
            "calendar_coverage": [
                sentinel,
                {
                    "account_id": "account-dated",
                    "date": "2026-08-13",
                    "provider": "SYNTHETIC",
                    "status": "OBSERVED",
                },
                {
                    "account_id": "account-future",
                    "date": "2026-08-14",
                    "provider": "SYNTHETIC",
                    "status": "OBSERVED",
                },
            ]
        },
        "2026-08-13",
    )

    assert projected["calendar_coverage"] == [
        sentinel,
        {
            "account_id": "account-dated",
            "date": "2026-08-13",
            "provider": "SYNTHETIC",
            "status": "OBSERVED",
        },
    ]


def test_cutoff_recomputes_account_live_evidence_and_clips_scopes() -> None:
    catalog = {
        "strategy_returns": [
            {
                "date": "2026-08-12",
                "evidence_layer": "ACCOUNT_LIVE",
                "provider_id": "synthetic",
                "return": "0.1",
                "strategy_id": "strategy-1",
            },
            {
                "date": "2026-08-13",
                "evidence_layer": "ACCOUNT_LIVE",
                "provider_id": "synthetic",
                "return": "-0.1",
                "strategy_id": "strategy-1",
            },
            {
                "date": "2026-08-14",
                "evidence_layer": "ACCOUNT_LIVE",
                "provider_id": "synthetic",
                "return": "0.4",
                "strategy_id": "strategy-1",
            },
        ],
        "strategy_provider_evidence": [
            {
                "account": {
                    "account_scopes": [
                        {
                            "account_id": "account-a",
                            "observed_start": "2026-08-10",
                            "observed_end": "2026-08-16",
                        },
                        {
                            "account_id": "account-future",
                            "observed_start": "2026-08-14",
                            "observed_end": "2026-08-16",
                        },
                    ],
                    "cumulative_return": "0.5",
                    "current_drawdown": "0",
                    "maximum_drawdown": "0",
                    "observed_start": "2026-08-12",
                    "observed_end": "2026-08-14",
                    "provider_id": "synthetic",
                    "return_observations": 3,
                },
                "evidence_layer": "ACCOUNT_LIVE",
                "strategy_id": "strategy-1",
            }
        ],
    }

    projected = reporter._calculation_catalog_through(catalog, "2026-08-13")
    evidence = reporter._strategy_provider_evidence(
        projected, calculation_cutoff_date="2026-08-13"
    )

    assert evidence[0]["account"] == {
        "account_scopes": [
            {
                "account_id": "account-a",
                "observed_start": "2026-08-10",
                "observed_end": "2026-08-13",
            }
        ],
        "cumulative_return": "-0.01",
        "current_drawdown": "-0.1",
        "max_drawdown": "-0.1",
        "maximum_drawdown": "-0.1",
        "observed_start": "2026-08-12",
        "observed_end": "2026-08-13",
        "provider_id": "synthetic",
        "return_observations": 2,
        "cumulative_return_label": "DAILY_GEOMETRIC_TWR",
        "path_dependent_metrics_status": "AVAILABLE",
        "path_dependent_metrics_reason": None,
        "return_series_schema": "ordered-date-return/1",
        "return_series_sha256": "eeb41ed52f04abfce9247d43a282b67a1cd0e32edcc273a575698c32dc2375a9",
    }


@pytest.mark.parametrize(
    "payload",
    (
        {"strategy_live_metrics": [{"last_date": "2026-08-14"}]},
        {"strategy_nav_coverage": [{"last_date": "2026-08-14"}]},
        {"hwm_acceptance_record": {"anchor_date": "2026-08-14"}},
        {"hwm_inputs": {"anchor": {"anchor_date": "2026-08-14"}, "periods": []}},
    ),
)
def test_calculation_cutoff_rejects_unprojectable_post_cutoff_context(
    payload: dict[str, Any],
) -> None:
    with pytest.raises(reporter.ReportValidationError, match="exceeds cutoff"):
        reporter._calculation_catalog_through(payload, "2026-08-13")


@pytest.mark.parametrize(
    ("effective_start", "effective_end", "expected_field"),
    (
        ("2026-08-14", None, "effective_start"),
        ("2026-08-01", "2026-08-14", "effective_end"),
    ),
)
def test_calculation_cutoff_rejects_future_assignment_before_nav_flow_traps(
    effective_start: str,
    effective_end: str | None,
    expected_field: str,
) -> None:
    catalog = {
        "assignments": [
            {
                "account_id": "future-account",
                "assignment_id": "future-assignment",
                "effective_start": effective_start,
                "effective_end": effective_end,
                "mapping_status": "MAPPED",
                "strategy_id": "future-strategy",
            }
        ],
        "nav_observations": [
            {
                "account_id": "future-account",
                "as_of_date": "2026-08-12",
                "value": "100",
            },
            {
                "account_id": "future-account",
                "as_of_date": "2026-08-13",
                "value": "NAV_TRAP_MUST_NOT_BE_PARSED",
            },
        ],
        "flows": [
            {
                "account_id": "future-account",
                "strategy_id": "future-strategy",
                "event_date": "2026-08-13",
                "amount": "FLOW_TRAP_MUST_NOT_BE_PARSED",
            }
        ],
        "calendar_coverage": [],
        "provider_strategy_returns": [],
    }

    with pytest.raises(
        reporter.ReportValidationError,
        match=rf"assignment {expected_field} exceeds cutoff",
    ):
        projected = reporter._calculation_catalog_through(catalog, "2026-08-13")
        reporter._account_strategy_performance(projected)


def test_calculation_cutoff_preserves_legitimate_historical_assignment() -> None:
    assignment = {
        "account_id": "historical-account",
        "assignment_id": "historical-assignment",
        "effective_start": "2026-07-01",
        "effective_end": "2026-08-10",
        "mapping_status": "MAPPED",
        "strategy_id": "historical-strategy",
    }

    projected = reporter._calculation_catalog_through(
        {"assignments": [assignment]}, "2026-08-13"
    )

    assert projected["assignments"] == [assignment]


def test_report_data_has_hand_checked_twr_and_coverage_failures(tmp_path: Path) -> None:
    baseline_root = _published_baseline(tmp_path)
    report, first_bytes, _ = reporter.build_report_data(baseline_root)
    _, second_bytes, _ = reporter.build_report_data(baseline_root)

    assert first_bytes == second_bytes
    mapped = next(
        item for item in report["twr_mwr"] if item["assignment_id"] == "assignment-a"
    )
    with reporter.localcontext() as context:
        context.prec = 34
        expected = ((Decimal("115") - Decimal("10")) / Decimal("100")) * (
            Decimal("121") / Decimal("115")
        ) - Decimal("1")
    assert mapped["twr"] == reporter.decimal_text(expected)
    assert mapped["twr_status"] == "AVAILABLE_WITH_INFERRED_FLOW_TIMING"
    assert mapped["mwr_xirr"] is None
    assert mapped["mwr_status"] == "UNAVAILABLE_INCOMPLETE_FLOW_COVERAGE"
    missing = next(
        item for item in report["twr_mwr"] if item["assignment_id"] == "assignment-b"
    )
    assert missing["twr"] is None
    assert missing["twr_status"] == "UNAVAILABLE_INSUFFICIENT_NAV"
    unmapped = next(
        item
        for item in report["twr_mwr"]
        if item["assignment_id"] == "assignment-general"
    )
    assert unmapped["twr_status"] == "UNMAPPED"
    assert report["positions"][0]["status"] == "UNAVAILABLE_NO_ACCEPTED_POSITION_SOURCE"
    assert list(report["strategy_summary"][0])[:4] == [
        "strategy_id",
        "strategy_name",
        "economic_family",
        "lifecycle_status",
    ]


def test_missing_strategy_evidence_is_unavailable_not_inferred(tmp_path: Path) -> None:
    catalog = _baseline_catalog()
    catalog["implementation_gaps"] = []
    catalog["costs"] = []
    catalog["strategy_nav_coverage"] = []
    baseline_root = _published_baseline(tmp_path, catalog)

    report, _, _ = reporter.build_report_data(baseline_root)
    strategy = next(
        item
        for item in report["strategy_summary"]
        if item["strategy_id"] == "strategy-1"
    )
    assert strategy["implementation_gap_status"] == "UNAVAILABLE"
    assert strategy["cost_status"] == "UNAVAILABLE"
    assert strategy["nav_status"] == "UNAVAILABLE"


def test_effective_dated_strategy_version_registry_projects_into_report(
    tmp_path: Path,
) -> None:
    catalog = _baseline_catalog()
    unavailable = {
        "source_ids": ["synthetic-source"],
        "status": "UNAVAILABLE",
        "version_ref": None,
    }
    catalog["strategy_version_catalog"] = {
        "as_of_date": catalog["as_of_date"],
        "catalog_schema_version": 2,
        "strategies": [
            {
                "backtest": dict(unavailable),
                "benchmark": dict(unavailable),
                "compatibility_reason": "SYNTHETIC compatibility is unresolved.",
                "compatibility_status": "UNRESOLVED",
                "configuration": dict(unavailable),
                "costs": dict(unavailable),
                "design": dict(unavailable),
                "disposition": "UNRESOLVED",
                "disposition_reason": "SYNTHETIC authoritative version is absent.",
                "effective_end": None,
                "effective_start": "2026-01-01",
                "evidence_layers": [
                    {
                        "artifact_sha256": None,
                        "disposition": "SYNTHETIC account-live route.",
                        "evidence_layer": "ACCOUNT_LIVE",
                        "observed_end": None,
                        "observed_start": None,
                        "source_id": "synthetic-source",
                        "status": "REFERENCE_ONLY",
                        "version_ref": None,
                    }
                ],
                "gaps": ["SYNTHETIC authoritative version is absent."],
                "live_boundary": {
                    "source_ids": ["synthetic-source"],
                    "status": "EVIDENCED" if index == 1 else "UNAVAILABLE",
                    "version_ref": f"live-v{index}" if index == 1 else None,
                    "observed_start": "2026-01-01" if index == 1 else None,
                    "observed_end": "2026-01-03" if index == 1 else None,
                },
                "source_route": {
                    "source_ids": ["synthetic-source"],
                    "status": "REFERENCE_ONLY",
                    "version_ref": None,
                },
                "strategy_id": f"strategy-{index}",
                "version_ref": None,
            }
            for index in range(1, 10)
        ],
    }
    baseline_root = _published_baseline(tmp_path, catalog)

    report, _, _ = reporter.build_report_data(baseline_root)
    first = next(
        item
        for item in report["strategy_summary"]
        if item["strategy_id"] == "strategy-1"
    )

    assert first["version_disposition"] == "UNRESOLVED"
    assert first["version_effective_start"] == "2026-01-01"
    assert first["live_boundary_status"] == "EVIDENCED"
    assert first["live_boundary_observed_end"] == "2026-01-03"
    assert first["version_compatibility_status"] == "UNRESOLVED"
    assert first["live_plus_backtest_status"] == (
        "BLOCKED_VERSION_NOT_PROVEN_COMPATIBLE"
    )
    assert first["live_plus_backtest_reason"].startswith("UNRESOLVED:")
    assert first["strategy_evidence_layers"][0]["evidence_layer"] == "ACCOUNT_LIVE"


def test_transaction_without_amount_remains_explicitly_unavailable(
    tmp_path: Path,
) -> None:
    catalog = _baseline_catalog()
    catalog["transactions"].append(
        {
            "account_id": "account-b",
            "amount": {"currency": "USD", "status": "UNAVAILABLE", "value": None},
            "event_date": "2026-01-03",
            "evidence_class": "WCH_HISTORICAL",
            "kind": "Stock Split",
            "mapping_status": "UNMAPPED",
            "provider": "SYNTHETIC",
            "quantity": "2",
            "record_id": "transaction-no-amount",
            "security_id": "security-1",
            "source_key": "synthetic-transaction",
            "strategy_id": None,
        }
    )
    baseline_root = _published_baseline(tmp_path, catalog)

    report, _, _ = reporter.build_report_data(baseline_root)

    row = next(
        item
        for item in report["transactions_and_flows"]
        if item["record_id"] == "transaction-no-amount"
    )
    assert row["amount_status"] == "UNAVAILABLE"
    assert row["amount_usd"] is None


def test_high_precision_performance_is_numeric_only_in_excel_presentation() -> None:
    exact = "0.001932516064483270690775302874751"

    assert reporter._excel_value(exact) == exact
    assert reporter._excel_value(exact, header="twr") == pytest.approx(
        0.0019325160644832707
    )


def test_workbook_parity_round_trips_high_precision_summary_ratio(
    tmp_path: Path,
) -> None:
    exact = "-0.0092359361880772460117548278757352"
    baseline_root = _published_baseline(tmp_path)
    report, _, _ = reporter.build_report_data(
        baseline_root,
        freshness_health_path=_hwm_freshness_file(tmp_path),
    )
    report["portfolio_summary"].append(
        {
            "metric": "High-precision ratio",
            "status": "AVAILABLE",
            "unit": "ratio",
            "value": exact,
        }
    )
    workbook_path = tmp_path / "high-precision-ratio.xlsx"
    reporter._workbook(report, workbook_path)

    parity = reporter._workbook_parity_receipt(
        canonical_json_bytes(report), workbook_path.read_bytes()
    )

    assert parity["report_row_count"] == sum(
        len(report[report_key]) for _, report_key, _, _ in reporter.WORKBOOK_TABLE_SPECS
    )
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        row = next(
            values
            for values in workbook["Portfolio Summary"].iter_rows(values_only=True)
            if values[0] == "High-precision ratio"
        )
        assert isinstance(row[3], float)
        assert row[3] == pytest.approx(float(Decimal(exact)))
    finally:
        workbook.close()


def test_workbook_is_deterministic_structured_and_has_no_external_surface(
    tmp_path: Path,
) -> None:
    baseline_root = _published_baseline(tmp_path)
    output_root = tmp_path / "reporting"

    first = reporter.build_candidate(baseline_root, output_root)
    second = reporter.build_candidate(baseline_root, output_root)

    assert first == second
    candidate = Path(first["candidate_dir"])
    workbook = Path(first["workbook_path"])
    evidence = reporter.verify_candidate(baseline_root, candidate, workbook)
    assert evidence["sheet_count"] == 14
    assert evidence["table_count"] == 14
    assert evidence["external_links"] == 0
    assert evidence["macro_parts"] == 0
    assert evidence["formula_count"] >= 1
    with zipfile.ZipFile(workbook) as archive:
        widths: list[float] = []
        for name in archive.namelist():
            if name.startswith("xl/worksheets/sheet") and name.endswith(".xml"):
                root = ET.fromstring(archive.read(name))
                widths.extend(
                    float(element.attrib["width"])
                    for element in root.iter()
                    if reporter._local_name(element.tag) == "col"
                    and "width" in element.attrib
                )
    assert widths
    assert max(widths) <= 45
    contract = json.loads((candidate / "workbook-contract.json").read_text())
    assert contract["required_sheets"] == list(reporter.SHEET_NAMES)
    assert set(contract["required_tables"]) == set(reporter.TABLE_NAMES)
    assert "expected_calculation_mode" not in contract
    assert {
        "cell": "E4",
        "expected": "PASS",
        "sheet": "Portfolio Summary",
    } in contract["sentinel_cells"]
    job_text = (candidate / "xlsx-job.json").read_text()
    job = json.loads(job_text)
    assert [step["type"] for step in job["steps"]] == ["open", "recalc", "save_as"]
    assert all(step["type"] != "refresh" for step in job["steps"])
    assert job_text.index('{"type":"open"') < job_text.index("workbook_path")
    assert job_text.index('{"type":"recalc"') < job_text.index("full_rebuild")
    assert job_text.index('{"type":"save_as"') < job_text.index("output_path")


def test_excel_recalc_uses_resettable_copy_and_cannot_mutate_initial(
    tmp_path: Path,
) -> None:
    baseline_root = _published_baseline(tmp_path)
    output_root = tmp_path / "reporting"
    first = reporter.build_candidate(baseline_root, output_root)
    candidate = Path(first["candidate_dir"])
    initial = Path(first["workbook_path"])
    canonical_bytes = initial.read_bytes()
    job = json.loads((candidate / "xlsx-job.json").read_text(encoding="utf-8"))
    excel_input = Path(job["steps"][0]["workbook_path"])

    assert excel_input == candidate / "excel-input" / reporter.WORKBOOK_NAME
    assert excel_input.read_bytes() == canonical_bytes

    # Native Excel may rewrite its writable input even when SaveAs targets a
    # different path. A retry must reset only that disposable copy.
    excel_input.write_bytes(b"native-excel-mutated-input")
    second = reporter.build_candidate(baseline_root, output_root)

    assert first == second
    assert initial.read_bytes() == canonical_bytes
    assert excel_input.read_bytes() == canonical_bytes


def test_candidate_drift_and_failed_excel_evidence_block_publication(
    tmp_path: Path,
) -> None:
    baseline_root = _published_baseline(tmp_path)
    output_root = tmp_path / "reporting"
    built = reporter.build_candidate(baseline_root, output_root)
    candidate = Path(built["candidate_dir"])
    workbook = Path(built["workbook_path"])

    (candidate / reporter.REPORT_DATA_NAME).write_text("{}", encoding="utf-8")
    with pytest.raises(reporter.ReportValidationError, match="differs"):
        reporter.verify_candidate(baseline_root, candidate, workbook)

    shutil.rmtree(candidate)
    built = reporter.build_candidate(baseline_root, output_root)
    candidate = Path(built["candidate_dir"])
    workbook = Path(built["workbook_path"])
    run_result = candidate / "xlsx-result.json"
    validation = candidate / "validation.json"
    run_result.write_text(json.dumps({"ok": False}), encoding="utf-8")
    validation.write_text(
        json.dumps(
            {
                "all_passed": True,
                "invariants": [{"name": "synthetic", "passed": True}],
            }
        ),
        encoding="utf-8",
    )
    reporter.canonicalize_workbook(
        workbook, candidate / "workbook-canonicalization.json"
    )
    with pytest.raises(reporter.ReportValidationError):
        reporter.publish_report(
            baseline_root,
            candidate,
            workbook,
            run_result,
            validation,
            tmp_path / "published",
        )


def test_workbook_evidence_binding_pins_exact_final_workbook_and_sidecars(
    tmp_path: Path,
) -> None:
    baseline_root = _published_baseline(tmp_path)
    built = reporter.build_candidate(baseline_root, tmp_path / "candidates")
    candidate = Path(built["candidate_dir"])
    initial_workbook = Path(built["workbook_path"])
    workbook, run_result, validation, binding_path = _prepare_workbook_evidence(
        candidate
    )

    binding = reporter.bind_workbook_evidence(
        candidate, workbook, run_result, validation, binding_path
    )

    assert binding["binding_schema_version"] == 1
    assert (
        binding["initial_workbook_sha256"]
        == hashlib.sha256(initial_workbook.read_bytes()).hexdigest()
    )
    assert (
        binding["workbook_sha256"] == hashlib.sha256(workbook.read_bytes()).hexdigest()
    )
    assert (
        binding["xlsx_result_sha256"]
        == hashlib.sha256(run_result.read_bytes()).hexdigest()
    )
    assert (
        binding["validation_result_sha256"]
        == hashlib.sha256(validation.read_bytes()).hexdigest()
    )
    assert (
        binding["workbook_contract_sha256"]
        == hashlib.sha256(
            (candidate / "workbook-contract.json").read_bytes()
        ).hexdigest()
    )
    assert binding["workbook_evidence_status"] == "WORKBOOK_DELIVERY_EVIDENCE_ACCEPTED"
    assert binding["publication_disposition"] == "ELIGIBLE_FOR_NON_AUTHORITY_DELIVERY"
    assert binding["workbook_delivery"]["delivery_evidence_schema_version"] == 1
    rendered = json.dumps(binding, sort_keys=True).lower()
    assert "native" not in rendered
    assert (
        binding["workbook_parity_sha256"]
        == hashlib.sha256((candidate / "workbook-parity.json").read_bytes()).hexdigest()
    )
    parity = json.loads(
        (candidate / "workbook-parity.json").read_text(encoding="utf-8")
    )
    assert parity == {
        "parity_schema_version": reporter.WORKBOOK_PARITY_VERSION,
        "parity_status": "REPO_BOUND_POSTRUN_ONLY",
        "report_sha256": built["report_data_sha256"],
        "workbook_sha256": binding["workbook_sha256"],
    }
    assert json.loads(binding_path.read_text(encoding="utf-8")) == binding


def test_workbook_evidence_binding_rejects_pre_excel_input_even_when_bytes_match(
    tmp_path: Path,
) -> None:
    baseline_root = _published_baseline(tmp_path)
    built = reporter.build_candidate(baseline_root, tmp_path / "candidates")
    candidate = Path(built["candidate_dir"])
    final_workbook, run_result, validation, binding_path = _prepare_workbook_evidence(
        candidate
    )
    initial_workbook = candidate / reporter.WORKBOOK_NAME
    initial_workbook.write_bytes(final_workbook.read_bytes())

    with pytest.raises(reporter.ReportValidationError, match="xlsx-win job binding"):
        reporter.bind_workbook_evidence(
            candidate,
            initial_workbook,
            run_result,
            validation,
            binding_path,
        )


@pytest.mark.parametrize("job_step", ("open", "save_as"))
def test_workbook_evidence_binding_rejects_wrong_excel_job_path(
    tmp_path: Path, job_step: str
) -> None:
    baseline_root = _published_baseline(tmp_path)
    built = reporter.build_candidate(baseline_root, tmp_path / "candidates")
    candidate = Path(built["candidate_dir"])
    final_workbook, run_result, validation, binding_path = _prepare_workbook_evidence(
        candidate
    )
    job_path = candidate / "xlsx-job.json"
    job = json.loads(job_path.read_text(encoding="utf-8"))
    if job_step == "open":
        job["steps"][0]["workbook_path"] = str(final_workbook)
    else:
        job["steps"][2]["output_path"] = str(candidate / reporter.WORKBOOK_NAME)
    job_path.write_bytes(canonical_json_bytes(job))

    with pytest.raises(reporter.ReportValidationError, match="xlsx-win job binding"):
        reporter.bind_workbook_evidence(
            candidate,
            final_workbook,
            run_result,
            validation,
            binding_path,
        )


_WORKBOOK_TABLE_CASES = (
    ("Portfolio Summary", "tblPortfolioSummary"),
    ("Account Summary", "tblAccounts"),
    ("Strategy Summary", "tblStrategies"),
    ("Positions", "tblPositions"),
    ("Transactions and Flows", "tblTransactionsFlows"),
    ("TWR MWR", "tblTwrMwr"),
    ("Allocation History", "tblAllocations"),
    ("Flow Reconciliation", "tblFlowRecon"),
    ("Source Reconciliation", "tblSourceRecon"),
    ("Coverage Data Quality", "tblCoverage"),
    ("Corrections", "tblCorrections"),
    ("Dependence", "tblDependence"),
    ("Methodology", "tblMethodology"),
    ("Provenance", "tblProvenance"),
)


@pytest.mark.parametrize(("sheet_name", "table_name"), _WORKBOOK_TABLE_CASES)
def test_every_governed_table_mutation_can_only_be_repo_bound(
    tmp_path: Path, sheet_name: str, table_name: str
) -> None:
    baseline_root = _published_baseline(tmp_path)
    built = reporter.build_candidate(baseline_root, tmp_path / "candidates")
    candidate = Path(built["candidate_dir"])
    final_workbook, run_result, validation, binding_path = _prepare_workbook_evidence(
        candidate
    )
    cell_reference = _first_table_data_cell(final_workbook, sheet_name, table_name)
    _rewrite_xlsx_cell(
        final_workbook,
        sheet_name=sheet_name,
        cell_reference=cell_reference,
        value="PARITY_MUTATION",
    )
    canonicalization_path = candidate / "workbook-canonicalization.json"
    canonicalization_path.unlink()
    reporter.canonicalize_workbook(final_workbook, canonicalization_path)

    binding = reporter.bind_workbook_evidence(
        candidate,
        final_workbook,
        run_result,
        validation,
        binding_path,
    )

    assert binding["workbook_evidence_status"] == "WORKBOOK_DELIVERY_EVIDENCE_ACCEPTED"


def test_workbook_header_mutation_can_only_be_repo_bound(
    tmp_path: Path,
) -> None:
    baseline_root = _published_baseline(tmp_path)
    built = reporter.build_candidate(baseline_root, tmp_path / "candidates")
    candidate = Path(built["candidate_dir"])
    final_workbook, run_result, validation, binding_path = _prepare_workbook_evidence(
        candidate
    )
    _rewrite_xlsx_cell(
        final_workbook,
        sheet_name="Strategy Summary",
        cell_reference="A5",
        value="Wrong Header",
    )
    canonicalization_path = candidate / "workbook-canonicalization.json"
    canonicalization_path.unlink()
    reporter.canonicalize_workbook(final_workbook, canonicalization_path)

    binding = reporter.bind_workbook_evidence(
        candidate,
        final_workbook,
        run_result,
        validation,
        binding_path,
    )

    assert binding["workbook_evidence_status"] == "WORKBOOK_DELIVERY_EVIDENCE_ACCEPTED"


def test_workbook_parity_does_not_round_expected_floats() -> None:
    value = 1.2345678901234567

    assert float(f"{value:.16G}") != value
    assert reporter._parity_cell_value(value, header="numeric_probe") == value


@pytest.mark.parametrize("mutation", ("formula", "cached_value"))
def test_workbook_formula_mutation_can_only_be_repo_bound(
    tmp_path: Path, mutation: str
) -> None:
    baseline_root = _published_baseline(tmp_path)
    built = reporter.build_candidate(baseline_root, tmp_path / "candidates")
    candidate = Path(built["candidate_dir"])
    final_workbook, run_result, validation, binding_path = _prepare_workbook_evidence(
        candidate
    )
    _rewrite_xlsx_cell(
        final_workbook,
        sheet_name="Portfolio Summary",
        cell_reference="E4",
        formula="1+1" if mutation == "formula" else None,
        cached_value="FAIL" if mutation == "cached_value" else None,
    )
    canonicalization_path = candidate / "workbook-canonicalization.json"
    canonicalization_path.unlink()
    reporter.canonicalize_workbook(final_workbook, canonicalization_path)

    binding = reporter.bind_workbook_evidence(
        candidate,
        final_workbook,
        run_result,
        validation,
        binding_path,
    )

    assert binding["workbook_evidence_status"] == "WORKBOOK_DELIVERY_EVIDENCE_ACCEPTED"


def test_workbook_readiness_mutation_can_only_be_repo_bound(
    tmp_path: Path,
) -> None:
    baseline_root = _published_baseline(tmp_path)
    built = reporter.build_candidate(baseline_root, tmp_path / "candidates")
    candidate = Path(built["candidate_dir"])
    final_workbook, run_result, validation, binding_path = _prepare_workbook_evidence(
        candidate
    )
    _rewrite_xlsx_cell(
        final_workbook,
        sheet_name="Portfolio Summary",
        cell_reference="B4",
        value="FORGED_READY",
    )
    canonicalization_path = candidate / "workbook-canonicalization.json"
    canonicalization_path.unlink()
    reporter.canonicalize_workbook(final_workbook, canonicalization_path)

    binding = reporter.bind_workbook_evidence(
        candidate,
        final_workbook,
        run_result,
        validation,
        binding_path,
    )

    assert binding["workbook_evidence_status"] == "WORKBOOK_DELIVERY_EVIDENCE_ACCEPTED"


def test_canonical_report_mutation_is_rejected_before_binding(tmp_path: Path) -> None:
    baseline_root = _published_baseline(tmp_path)
    built = reporter.build_candidate(baseline_root, tmp_path / "candidates")
    candidate = Path(built["candidate_dir"])
    final_workbook, run_result, validation, binding_path = _prepare_workbook_evidence(
        candidate
    )
    report_path = candidate / reporter.REPORT_DATA_NAME
    report = json.loads(report_path.read_text(encoding="utf-8"))
    report["provenance"][0]["sha256"] = "0" * 64
    report_path.write_bytes(canonical_json_bytes(report))

    with pytest.raises(
        reporter.ReportValidationError,
        match="candidate report|workbook parity values differ",
    ):
        reporter.bind_workbook_evidence(
            candidate,
            final_workbook,
            run_result,
            validation,
            binding_path,
        )


@pytest.mark.parametrize(
    ("target_name", "error_category"),
    (
        ("workbook", "workbook evidence binding differs"),
        ("initial-workbook", "workbook evidence binding differs"),
        ("xlsx-result.json", "workbook evidence binding is invalid"),
        ("validation.json", "workbook evidence binding is invalid"),
        ("workbook-contract.json", "workbook evidence binding differs"),
        ("xlsx-job.json", "workbook evidence binding differs"),
        ("workbook-canonicalization.json", "workbook evidence binding differs"),
        ("xlsx-result-canonicalization.json", "workbook evidence binding differs"),
        ("workbook-parity.json", "workbook evidence binding differs"),
        ("workbook-evidence-binding.json", "workbook evidence binding is invalid"),
    ),
)
def test_any_bound_workbook_evidence_byte_change_blocks_publication(
    tmp_path: Path, target_name: str, error_category: str
) -> None:
    baseline_root = _published_baseline(tmp_path)
    built = reporter.build_candidate(baseline_root, tmp_path / "candidates")
    candidate = Path(built["candidate_dir"])
    workbook, run_result, validation, binding_path = _prepare_workbook_evidence(
        candidate,
        run_id="f" * 32,
        temporary_root=r"C:\Temp\xlsx-drift",
    )
    reporter.bind_workbook_evidence(
        candidate, workbook, run_result, validation, binding_path
    )
    target = (
        workbook
        if target_name == "workbook"
        else candidate / reporter.WORKBOOK_NAME
        if target_name == "initial-workbook"
        else binding_path
        if target_name == "workbook-evidence-binding.json"
        else candidate / target_name
    )
    target.write_bytes(target.read_bytes() + b"drift")

    published_root = tmp_path / "published"
    with pytest.raises(reporter.ReportValidationError, match=f"^{error_category}$"):
        reporter.publish_report(
            baseline_root,
            candidate,
            workbook,
            run_result,
            validation,
            published_root,
            workbook_evidence_binding=binding_path,
        )
    assert not (published_root / "CURRENT.json").exists()
    assert not (published_root / "generations").exists()
    assert not (published_root / "deliveries").exists()


def test_validated_workbook_evidence_publishes_authority_and_delivery(
    tmp_path: Path,
) -> None:
    baseline_root = _published_baseline(tmp_path)
    candidate_root = tmp_path / "reporting-candidates"
    built = reporter.build_candidate(baseline_root, candidate_root)
    candidate = Path(built["candidate_dir"])
    workbook, run_result, validation, binding_path = _prepare_workbook_evidence(
        candidate,
        run_id="d" * 32,
        temporary_root=r"C:\Temp\xlsx-publish",
    )
    reporter.bind_workbook_evidence(
        candidate, workbook, run_result, validation, binding_path
    )
    published_root = _published_baseline(tmp_path / "published")
    published = reporter.publish_report(
        baseline_root,
        candidate,
        workbook,
        run_result,
        validation,
        published_root,
        workbook_evidence_binding=binding_path,
    )

    pointer = read_current(published_root)
    assert pointer.generation_id == published["generation_id"]
    assert (
        published_root / "deliveries" / pointer.generation_id / reporter.WORKBOOK_NAME
    ).is_file()


def test_accepted_delivery_evidence_publishes_without_authority_upgrade(
    tmp_path: Path,
) -> None:
    """Workbook QA evidence gates delivery without becoming report authority."""

    baseline_root = _published_baseline(tmp_path / "baseline")
    candidate = Path(
        reporter.build_candidate(baseline_root, tmp_path / "candidates")[
            "candidate_dir"
        ]
    )
    workbook, run_result, validation, binding_path = _prepare_workbook_evidence(
        candidate,
        run_id="4" * 32,
        temporary_root=r"C:\\Temp\\xlsx-delivery-evidence",
    )
    binding = reporter.bind_workbook_evidence(
        candidate, workbook, run_result, validation, binding_path
    )
    assert binding["workbook_evidence_status"] == "WORKBOOK_DELIVERY_EVIDENCE_ACCEPTED"
    assert binding["publication_disposition"] == "ELIGIBLE_FOR_NON_AUTHORITY_DELIVERY"
    assert "native_execution" not in binding
    assert "workbook_delivery" in binding

    published_root = tmp_path / "published"
    published = reporter.publish_report(
        baseline_root,
        candidate,
        workbook,
        run_result,
        validation,
        published_root,
        workbook_evidence_binding=binding_path,
    )

    pointer = read_current(published_root)
    assert published["generation_id"] == pointer.generation_id
    assert not hasattr(reporter, "complete_workbook_evidence")


@pytest.mark.parametrize(
    ("target", "error_category"),
    (
        ("xlsx_result", "xlsx-win calculation evidence is not successful"),
        ("validation", "xlsx-win workbook invariants did not all pass"),
    ),
)
def test_failed_or_incomplete_workbook_evidence_cannot_create_delivery_binding(
    tmp_path: Path, target: str, error_category: str
) -> None:
    """The binder admits only the complete supplied QA-evidence shape."""

    baseline_root = _published_baseline(tmp_path / "baseline")
    candidate = Path(
        reporter.build_candidate(baseline_root, tmp_path / "candidates")[
            "candidate_dir"
        ]
    )
    workbook, run_result, validation, binding_path = _prepare_workbook_evidence(
        candidate,
        run_id="1" * 32,
        temporary_root=r"C:\\Temp\\xlsx-evidence-admission",
    )
    evidence_path = run_result if target == "xlsx_result" else validation
    evidence = json.loads(evidence_path.read_text(encoding="utf-8"))
    if target == "xlsx_result":
        evidence["final_state"] = "FAILED"
        evidence["ok"] = False
    else:
        evidence["all_passed"] = False
        evidence["invariants"][0]["passed"] = False
    evidence_path.write_bytes(canonical_json_bytes(evidence))

    with pytest.raises(reporter.ReportValidationError, match=f"^{error_category}$"):
        reporter.bind_workbook_evidence(
            candidate,
            workbook,
            run_result,
            validation,
            binding_path,
        )

    assert not binding_path.exists()


@pytest.mark.parametrize("mutation", ("relabel", "extra", "missing"))
def test_noncanonical_delivery_binding_cannot_advance_authority(
    tmp_path: Path, mutation: str
) -> None:
    """Manual labels and closed-structure changes never alter delivery admission."""

    baseline_root = _published_baseline(tmp_path / "baseline")
    candidate = Path(
        reporter.build_candidate(baseline_root, tmp_path / "candidates")[
            "candidate_dir"
        ]
    )
    workbook, run_result, validation, binding_path = _prepare_workbook_evidence(
        candidate,
        run_id="1" * 32,
        temporary_root=r"C:\\Temp\\xlsx-relabel-rejection",
    )
    reporter.bind_workbook_evidence(
        candidate, workbook, run_result, validation, binding_path
    )
    mutated = json.loads(binding_path.read_text(encoding="utf-8"))
    if mutation == "relabel":
        mutated["workbook_evidence_status"] = "MANUALLY_RELABELLED"
    elif mutation == "extra":
        mutated["unbound"] = True
    else:
        del mutated["workbook_delivery"]
    binding_path.write_bytes(canonical_json_bytes(mutated))
    published_root = _published_baseline(tmp_path / "published")
    pointer_before = (published_root / "CURRENT.json").read_bytes()
    generations_before = sorted((published_root / "generations").iterdir())

    with pytest.raises(
        reporter.ReportValidationError, match="^workbook evidence binding differs$"
    ):
        reporter.publish_report(
            baseline_root,
            candidate,
            workbook,
            run_result,
            validation,
            published_root,
            workbook_evidence_binding=binding_path,
        )

    assert (published_root / "CURRENT.json").read_bytes() == pointer_before
    assert sorted((published_root / "generations").iterdir()) == generations_before
    assert not (published_root / "deliveries").exists()


@pytest.mark.parametrize(
    ("field_path", "replacement"),
    (
        (("xlsx_result", "sha256"), "0" * 64),
        (("xlsx_result", "size_bytes"), 0),
        (("validation_result", "sha256"), "0" * 64),
        (("validation_result", "size_bytes"), 0),
        (("final_workbook", "sha256"), "0" * 64),
        (("final_workbook", "size_bytes"), 0),
        (("engine", "run_id"), "wrong-run-id"),
        (("engine", "idempotency_key"), "wrong-idempotency-key"),
        (("engine", "result_schema_version"), "0.0"),
        (("steps", 1, "status"), "failed"),
        (("validation", "invariant_count"), 0),
        (("validation", "passed_invariant_count"), 0),
    ),
)
def test_delivery_evidence_binding_mutation_cannot_advance_authority(
    tmp_path: Path, field_path: tuple[str | int, ...], replacement: str | int
) -> None:
    """Every delivery-evidence field is rederived from bound QA artifacts."""

    baseline_root = _published_baseline(tmp_path / "baseline")
    candidate = Path(
        reporter.build_candidate(baseline_root, tmp_path / "candidates")[
            "candidate_dir"
        ]
    )
    workbook, run_result, validation, binding_path = _prepare_workbook_evidence(
        candidate,
        run_id="2" * 32,
        temporary_root=r"C:\\Temp\\xlsx-delivery-evidence-mutation",
    )
    binding = reporter.bind_workbook_evidence(
        candidate, workbook, run_result, validation, binding_path
    )
    assert binding == json.loads(binding_path.read_text(encoding="utf-8"))
    mutated = json.loads(binding_path.read_text(encoding="utf-8"))
    target: Any = mutated["workbook_delivery"]
    for key in field_path[:-1]:
        target = target[key]
    target[field_path[-1]] = replacement
    binding_path.write_bytes(canonical_json_bytes(mutated))
    published_root = _published_baseline(tmp_path / "published")
    pointer_before = (published_root / "CURRENT.json").read_bytes()
    generations_before = sorted((published_root / "generations").iterdir())

    with pytest.raises(
        reporter.ReportValidationError, match="^workbook evidence binding differs$"
    ):
        reporter.publish_report(
            baseline_root,
            candidate,
            workbook,
            run_result,
            validation,
            published_root,
            workbook_evidence_binding=binding_path,
        )

    assert (published_root / "CURRENT.json").read_bytes() == pointer_before
    assert sorted((published_root / "generations").iterdir()) == generations_before
    assert not (published_root / "deliveries").exists()


def test_validated_local_workbook_evidence_publishes_an_immutable_report(
    tmp_path: Path,
) -> None:
    """A closed delivery-evidence binding permits the producer path."""

    baseline_root = _published_baseline(tmp_path / "baseline")
    built = reporter.build_candidate(baseline_root, tmp_path / "candidates")
    candidate = Path(built["candidate_dir"])
    workbook, run_result, validation, binding_path = _prepare_workbook_evidence(
        candidate,
        run_id="9" * 32,
        temporary_root=r"C:\Temp\xlsx-local-publication",
    )
    reporter.bind_workbook_evidence(
        candidate, workbook, run_result, validation, binding_path
    )
    expected_workbook = workbook.read_bytes()
    expected_report = (candidate / reporter.REPORT_DATA_NAME).read_bytes()
    published_root = tmp_path / "published"

    published = reporter.publish_report(
        baseline_root,
        candidate,
        workbook,
        run_result,
        validation,
        published_root,
        workbook_evidence_binding=binding_path,
    )

    pointer = read_current(published_root)
    generation = published_root / "generations" / pointer.generation_id
    delivery = published_root / "deliveries" / pointer.generation_id
    assert published["generation_id"] == pointer.generation_id
    assert not (generation / reporter.WORKBOOK_NAME).exists()
    assert (delivery / reporter.WORKBOOK_NAME).read_bytes() == expected_workbook
    assert (generation / reporter.REPORT_DATA_NAME).read_bytes() == expected_report
    assert (generation / "report-verification.json").is_file()


def test_publish_keeps_authority_json_immutable_when_delivery_workbook_drifts(
    tmp_path: Path,
) -> None:
    """Putting an XLSX in the authority generation must break this contract."""

    baseline_root = _published_baseline(tmp_path / "baseline")
    candidate = Path(
        reporter.build_candidate(baseline_root, tmp_path / "candidates")[
            "candidate_dir"
        ]
    )
    workbook, run_result, validation, binding_path = _prepare_workbook_evidence(
        candidate,
        run_id="a" * 32,
        temporary_root=r"C:\\Temp\\xlsx-authority-delivery",
    )
    reporter.bind_workbook_evidence(
        candidate, workbook, run_result, validation, binding_path
    )
    expected_report = (candidate / "portfolio-report-data.json").read_bytes()
    expected_report_sha256 = hashlib.sha256(expected_report).hexdigest()
    published_root = tmp_path / "published"

    published = reporter.publish_report(
        baseline_root,
        candidate,
        workbook,
        run_result,
        validation,
        published_root,
        workbook_evidence_binding=binding_path,
    )

    pointer = read_current(published_root)
    generation = published_root / "generations" / pointer.generation_id
    delivery = published_root / "deliveries" / pointer.generation_id
    receipt_path = delivery / "report-delivery.json"
    workbook_path = delivery / "portfolio-report.xlsx"
    assert published["generation_id"] == pointer.generation_id
    assert (generation / "portfolio-report-data.json").read_bytes() == expected_report
    assert (generation / "portfolio-report.xlsx").exists() is False
    assert not list(generation.rglob("*.xlsx"))
    authority_paths = sorted(
        path.relative_to(generation).as_posix()
        for path in generation.rglob("*")
        if path.is_file()
    )
    assert authority_paths == [
        "generation.json",
        "portfolio-report-data.json",
        "report-verification.json",
        "workbook-canonicalization.json",
        "workbook-contract.json",
        "workbook-evidence-binding.json",
        "workbook-parity.json",
        "workbook-validation.json",
        "xlsx-job.json",
        "xlsx-result-canonicalization.json",
        "xlsx-win-result.json",
    ]
    expected_inventory_sha256 = hashlib.sha256(
        canonical_json_bytes(
            {
                "files": [
                    {
                        "relative_path": relative_path,
                        "sha256": hashlib.sha256(
                            (generation / relative_path).read_bytes()
                        ).hexdigest(),
                        "size_bytes": len((generation / relative_path).read_bytes()),
                    }
                    for relative_path in authority_paths
                ],
                "inventory_schema_version": 1,
            }
        )
    ).hexdigest()
    expected_manifest = (generation / "generation.json").read_bytes()
    assert workbook_path.is_file()
    receipt = json.loads(receipt_path.read_text(encoding="utf-8"))
    assert receipt["delivery_classification"] == "NON_AUTHORITY_HUMAN_DELIVERY"
    assert receipt["authority"] == {
        "generation_id": pointer.generation_id,
        "generation_inventory_sha256": expected_inventory_sha256,
        "generation_manifest_sha256": hashlib.sha256(expected_manifest).hexdigest(),
        "generation_manifest_size_bytes": len(expected_manifest),
        "payload": {
            "relative_path": "portfolio-report-data.json",
            "sha256": expected_report_sha256,
            "size_bytes": len(expected_report),
        },
    }
    assert pointer.artifact_inventory_sha256 == expected_inventory_sha256
    assert pointer.manifest_sha256 == hashlib.sha256(expected_manifest).hexdigest()
    assert receipt["workbook"]["relative_path"] == "portfolio-report.xlsx"
    from jsonschema import Draft202012Validator

    delivery_schema = json.loads(
        (
            REPOSITORY_ROOT / "schemas" / "portfolio" / "report-delivery.schema.json"
        ).read_text(encoding="utf-8")
    )
    Draft202012Validator.check_schema(delivery_schema)
    Draft202012Validator(delivery_schema).validate(receipt)

    assert (
        reporter.publish_report(
            baseline_root,
            candidate,
            workbook,
            run_result,
            validation,
            published_root,
            workbook_evidence_binding=binding_path,
        )
        == published
    )

    with zipfile.ZipFile(workbook_path, "a") as archive:
        archive.writestr("customXml/item1.xml", b"<sharepoint-metadata/>")

    delivery_status = reporter.inspect_report_delivery(
        published_root, pointer.generation_id, delivery
    )
    assert delivery_status["workbook_status"] == "DRIFTED_NON_AUTHORITY"
    assert delivery_status["expected_workbook_sha256"] == receipt["workbook"]["sha256"]
    assert (
        delivery_status["actual_workbook_sha256"]
        == hashlib.sha256(workbook_path.read_bytes()).hexdigest()
    )
    assert read_current(published_root) == pointer

    with pytest.raises(reporter.GenerationConflictError):
        reporter.publish_report(
            baseline_root,
            candidate,
            workbook,
            run_result,
            validation,
            published_root,
            workbook_evidence_binding=binding_path,
        )


def test_delivery_inspection_requires_closed_receipt_and_exact_authority_binding(
    tmp_path: Path,
) -> None:
    """Dropping a receipt field or rebinding it to another authority must fail."""

    baseline_root = _published_baseline(tmp_path / "baseline")
    candidate = Path(
        reporter.build_candidate(baseline_root, tmp_path / "candidates")[
            "candidate_dir"
        ]
    )
    workbook, run_result, validation, binding_path = _prepare_workbook_evidence(
        candidate,
        run_id="c" * 32,
        temporary_root=r"C:\\Temp\\xlsx-delivery-inspection",
    )
    reporter.bind_workbook_evidence(
        candidate, workbook, run_result, validation, binding_path
    )
    published_root = tmp_path / "published"
    published = reporter.publish_report(
        baseline_root,
        candidate,
        workbook,
        run_result,
        validation,
        published_root,
        workbook_evidence_binding=binding_path,
    )
    generation_id = published["generation_id"]
    delivery = published_root / "deliveries" / generation_id
    receipt_path = delivery / "report-delivery.json"
    original_receipt = receipt_path.read_bytes()
    manifest_bytes = (
        published_root / "generations" / generation_id / "generation.json"
    ).read_bytes()
    report_bytes = (
        published_root / "generations" / generation_id / "portfolio-report-data.json"
    ).read_bytes()

    exact = reporter.inspect_report_delivery(published_root, generation_id, delivery)
    assert exact["workbook_status"] == "EXACT_NON_AUTHORITY"
    assert (
        exact["generation_manifest_sha256"]
        == hashlib.sha256(manifest_bytes).hexdigest()
    )
    assert exact["report_data_sha256"] == hashlib.sha256(report_bytes).hexdigest()

    size_only = json.loads(original_receipt)
    size_only["workbook"]["size_bytes"] -= 1
    receipt_path.write_bytes(canonical_json_bytes(size_only))
    assert (
        reporter.inspect_report_delivery(published_root, generation_id, delivery)[
            "workbook_status"
        ]
        == "DRIFTED_NON_AUTHORITY"
    )

    malformed = json.loads(original_receipt)
    malformed["unexpected"] = True
    receipt_path.write_bytes(canonical_json_bytes(malformed))
    with pytest.raises(
        reporter.ReportValidationError,
        match="^report delivery receipt structure differs$",
    ):
        reporter.inspect_report_delivery(published_root, generation_id, delivery)

    rebound = json.loads(original_receipt)
    rebound["authority"]["payload"]["sha256"] = "0" * 64
    receipt_path.write_bytes(canonical_json_bytes(rebound))
    with pytest.raises(
        reporter.ReportValidationError,
        match="^report delivery receipt authority binding differs$",
    ):
        reporter.inspect_report_delivery(published_root, generation_id, delivery)


@pytest.mark.parametrize(
    ("mutation", "expected"),
    (
        (("authority", "generation_id"), "authority"),
        (("authority", "generation_manifest_sha256"), "authority"),
        (("authority", "generation_manifest_size_bytes"), "authority"),
        (("authority", "generation_inventory_sha256"), "authority"),
        (("authority", "payload", "relative_path"), "authority"),
        (("authority", "payload", "sha256"), "authority"),
        (("authority", "payload", "size_bytes"), "authority"),
        (("workbook", "relative_path"), "structure"),
        (("workbook", "sha256"), "drift"),
        (("workbook", "size_bytes"), "drift"),
        (("missing_authority",), "structure"),
        (("missing_payload",), "authority"),
        (("extra",), "structure"),
    ),
)
def test_delivery_inspection_receipt_mutations_are_closed_or_non_authority_drift(
    tmp_path: Path, mutation: tuple[str, ...], expected: str
) -> None:
    """Every required receipt binding is either exact or a named non-authority drift."""

    published_root, generation_id, delivery = _published_report_delivery(
        tmp_path, run_id="3" * 32
    )
    receipt_path = delivery / "report-delivery.json"
    receipt = json.loads(receipt_path.read_text(encoding="utf-8"))
    if mutation == ("missing_authority",):
        del receipt["authority"]
    elif mutation == ("missing_payload",):
        del receipt["authority"]["payload"]
    elif mutation == ("extra",):
        receipt["unexpected"] = True
    else:
        target: Any = receipt
        for key in mutation[:-1]:
            target = target[key]
        key = mutation[-1]
        value = target[key]
        target[key] = value + 1 if isinstance(value, int) else "0" * 64
    receipt_path.write_bytes(canonical_json_bytes(receipt))

    if expected == "authority":
        with pytest.raises(
            reporter.ReportValidationError,
            match="^report delivery receipt authority binding differs$",
        ):
            reporter.inspect_report_delivery(published_root, generation_id, delivery)
    elif expected == "structure":
        with pytest.raises(
            reporter.ReportValidationError,
            match="^report delivery receipt structure differs$",
        ):
            reporter.inspect_report_delivery(published_root, generation_id, delivery)
    else:
        assert (
            reporter.inspect_report_delivery(published_root, generation_id, delivery)[
                "workbook_status"
            ]
            == "DRIFTED_NON_AUTHORITY"
        )


@pytest.mark.parametrize("conflicting_target", ("receipt", "workbook"))
def test_delivery_preflight_does_not_create_other_target_on_conflict(
    tmp_path: Path, conflicting_target: str
) -> None:
    """A conflicting existing delivery target must be found before either write."""

    baseline_root = _published_baseline(tmp_path / "baseline")
    candidate = Path(
        reporter.build_candidate(baseline_root, tmp_path / "candidates")[
            "candidate_dir"
        ]
    )
    workbook, run_result, validation, binding_path = _prepare_workbook_evidence(
        candidate,
        run_id="e" * 32,
        temporary_root=r"C:\\Temp\\xlsx-delivery-preflight",
    )
    reporter.bind_workbook_evidence(
        candidate, workbook, run_result, validation, binding_path
    )
    published_root = tmp_path / "published"
    published = reporter.publish_report(
        baseline_root,
        candidate,
        workbook,
        run_result,
        validation,
        published_root,
        workbook_evidence_binding=binding_path,
    )
    generation_id = published["generation_id"]
    delivery = published_root / "deliveries" / generation_id
    receipt_path = delivery / "report-delivery.json"
    workbook_path = delivery / "portfolio-report.xlsx"
    pointer_before = (published_root / "CURRENT.json").read_bytes()
    if conflicting_target == "receipt":
        receipt = json.loads(receipt_path.read_text(encoding="utf-8"))
        receipt["workbook"]["sha256"] = "0" * 64
        conflicting_bytes = canonical_json_bytes(receipt)
        receipt_path.write_bytes(conflicting_bytes)
        workbook_path.unlink()
        missing_path = workbook_path
        conflicting_path = receipt_path
    else:
        conflicting_bytes = workbook_path.read_bytes() + b"sharepoint-drift"
        workbook_path.write_bytes(conflicting_bytes)
        receipt_path.unlink()
        missing_path = receipt_path
        conflicting_path = workbook_path

    with pytest.raises(
        reporter.GenerationConflictError,
        match="^report delivery target differs$",
    ):
        reporter.publish_report(
            baseline_root,
            candidate,
            workbook,
            run_result,
            validation,
            published_root,
            workbook_evidence_binding=binding_path,
        )

    assert missing_path.exists() is False
    assert conflicting_path.read_bytes() == conflicting_bytes
    assert (published_root / "CURRENT.json").read_bytes() == pointer_before


@pytest.mark.parametrize("missing_target", ("receipt", "workbook"))
def test_delivery_preflight_recovers_exact_partial_delivery(
    tmp_path: Path, missing_target: str
) -> None:
    """An exact surviving delivery file may be paired with its missing sibling."""

    baseline_root = _published_baseline(tmp_path / "baseline")
    candidate = Path(
        reporter.build_candidate(baseline_root, tmp_path / "candidates")[
            "candidate_dir"
        ]
    )
    workbook, run_result, validation, binding_path = _prepare_workbook_evidence(
        candidate,
        run_id="f" * 32,
        temporary_root=r"C:\\Temp\\xlsx-delivery-partial",
    )
    reporter.bind_workbook_evidence(
        candidate, workbook, run_result, validation, binding_path
    )
    published_root = tmp_path / "published"
    published = reporter.publish_report(
        baseline_root,
        candidate,
        workbook,
        run_result,
        validation,
        published_root,
        workbook_evidence_binding=binding_path,
    )
    delivery = published_root / "deliveries" / published["generation_id"]
    receipt_path = delivery / "report-delivery.json"
    workbook_path = delivery / "portfolio-report.xlsx"
    expected_receipt = receipt_path.read_bytes()
    expected_workbook = workbook_path.read_bytes()
    pointer_before = (published_root / "CURRENT.json").read_bytes()
    (receipt_path if missing_target == "receipt" else workbook_path).unlink()

    assert (
        reporter.publish_report(
            baseline_root,
            candidate,
            workbook,
            run_result,
            validation,
            published_root,
            workbook_evidence_binding=binding_path,
        )
        == published
    )
    assert receipt_path.read_bytes() == expected_receipt
    assert workbook_path.read_bytes() == expected_workbook
    assert (published_root / "CURRENT.json").read_bytes() == pointer_before


@pytest.mark.parametrize("receipt_kind", ("missing", "malformed", "caller_created"))
def test_publish_requires_exact_bound_workbook_evidence_for_every_caller_input(
    tmp_path: Path, receipt_kind: str
) -> None:
    baseline_root = _published_baseline(tmp_path / "baseline")
    candidate = Path(
        reporter.build_candidate(baseline_root, tmp_path / "candidates")[
            "candidate_dir"
        ]
    )
    workbook, run_result, validation, binding_path = _prepare_workbook_evidence(
        candidate
    )
    reporter.bind_workbook_evidence(
        candidate, workbook, run_result, validation, binding_path
    )
    published_root = _published_baseline(tmp_path / "published")
    pointer_before = (published_root / "CURRENT.json").read_bytes()
    generations_before = sorted((published_root / "generations").iterdir())
    receipt: Path | None = binding_path
    if receipt_kind == "missing":
        receipt = None
    elif receipt_kind == "malformed":
        receipt = candidate / "fake-native-receipt.json"
        receipt.write_text("{not-json", encoding="utf-8")
    elif receipt_kind == "caller_created":
        receipt = candidate / "fake-native-receipt.json"
        events = candidate / "xlsx-events.json"
        events.write_bytes(canonical_json_bytes({"events": [{"type": "save_as"}]}))
        receipt.write_bytes(
            canonical_json_bytes(
                {
                    "actual_job_path": str(candidate / "xlsx-job.json"),
                    "actual_job_sha256": hashlib.sha256(
                        (candidate / "xlsx-job.json").read_bytes()
                    ).hexdigest(),
                    "actual_result_path": str(run_result),
                    "actual_result_sha256": hashlib.sha256(
                        run_result.read_bytes()
                    ).hexdigest(),
                    "events_path": str(events),
                    "events_sha256": hashlib.sha256(events.read_bytes()).hexdigest(),
                    "original_workbook": {
                        "path": str(candidate / reporter.WORKBOOK_NAME),
                        "sha256": hashlib.sha256(
                            (candidate / reporter.WORKBOOK_NAME).read_bytes()
                        ).hexdigest(),
                    },
                    "published_workbook": {
                        "path": str(published_root / reporter.WORKBOOK_NAME),
                        "sha256": hashlib.sha256(workbook.read_bytes()).hexdigest(),
                    },
                    "run_id": json.loads(run_result.read_text(encoding="utf-8"))[
                        "run_id"
                    ],
                    "schema_version": "xlsx-win-native-receipt/1",
                    "staged_workbook": {
                        "path": str(workbook),
                        "sha256": hashlib.sha256(workbook.read_bytes()).hexdigest(),
                    },
                }
            )
        )

    with pytest.raises(reporter.ReportValidationError):
        reporter.publish_report(
            baseline_root,
            candidate,
            workbook,
            run_result,
            validation,
            published_root,
            workbook_evidence_binding=receipt,
        )

    assert (published_root / "CURRENT.json").read_bytes() == pointer_before
    assert sorted((published_root / "generations").iterdir()) == generations_before
    assert not (published_root / "CURRENT_RELEASE.json").exists()
    assert not list(published_root.rglob("report-verification.json"))


def test_publish_cli_accepts_exact_bound_workbook_evidence(
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    baseline_root = _published_baseline(tmp_path / "baseline")
    candidate = Path(
        reporter.build_candidate(baseline_root, tmp_path / "candidates")[
            "candidate_dir"
        ]
    )
    workbook, run_result, validation, binding_path = _prepare_workbook_evidence(
        candidate
    )
    reporter.bind_workbook_evidence(
        candidate, workbook, run_result, validation, binding_path
    )
    published_root = _published_baseline(tmp_path / "published")
    pointer_before = (published_root / "CURRENT.json").read_bytes()

    exit_code = reporter.main(
        [
            "publish",
            "--baseline-root",
            str(baseline_root),
            "--candidate-dir",
            str(candidate),
            "--excel-workbook",
            str(workbook),
            "--xlsx-result",
            str(run_result),
            "--validation-result",
            str(validation),
            "--workbook-evidence-binding",
            str(binding_path),
            "--output-root",
            str(published_root),
        ]
    )

    captured = capsys.readouterr()
    assert exit_code == 0
    assert captured.err == ""
    assert (published_root / "CURRENT.json").read_bytes() != pointer_before
    assert read_current(published_root).generation_id in captured.out


def test_delivery_evidence_stages_without_current_or_release_pointer(
    tmp_path: Path,
) -> None:
    baseline_root = _published_baseline(tmp_path / "baseline")
    built = reporter.build_candidate(baseline_root, tmp_path / "staging")
    candidate = Path(built["candidate_dir"])
    workbook, run_result, validation, binding_path = _prepare_workbook_evidence(
        candidate
    )

    binding = reporter.bind_workbook_evidence(
        candidate, workbook, run_result, validation, binding_path
    )

    assert binding["workbook_evidence_status"] == "WORKBOOK_DELIVERY_EVIDENCE_ACCEPTED"
    assert (candidate / reporter.WORKBOOK_NAME).is_file()
    assert (candidate / "workbook-evidence-binding.json").is_file()
    assert not (candidate / "CURRENT.json").exists()
    assert not (candidate / "CURRENT_RELEASE.json").exists()


def test_publication_rejects_source_rollover_before_delivery(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    baseline_root = _published_baseline(tmp_path)
    built = reporter.build_candidate(baseline_root, tmp_path / "candidates")
    candidate = Path(built["candidate_dir"])
    workbook, run_result, validation, binding_path = _prepare_workbook_evidence(
        candidate,
        run_id="a" * 32,
        temporary_root=r"C:\Temp\xlsx-source-rollover",
    )
    reporter.bind_workbook_evidence(
        candidate, workbook, run_result, validation, binding_path
    )
    original_build_report_data = reporter.build_report_data
    calls = 0

    def build_then_roll_source(*args: Any, **kwargs: Any) -> Any:
        nonlocal calls
        result = original_build_report_data(*args, **kwargs)
        calls += 1
        if calls == 1:
            changed = _baseline_catalog()
            changed["replay"]["total_liquid_usd"] = "101"
            _published_baseline(tmp_path, changed)
        return result

    monkeypatch.setattr(reporter, "build_report_data", build_then_roll_source)

    with pytest.raises(reporter.ReportValidationError, match="source identity drift"):
        reporter.publish_report(
            baseline_root,
            candidate,
            workbook,
            run_result,
            validation,
            tmp_path / "published",
            workbook_evidence_binding=binding_path,
        )


def test_report_schema_is_closed_and_keeps_mwr_unavailable() -> None:
    from jsonschema import Draft202012Validator, ValidationError

    schema = json.loads(
        (
            REPOSITORY_ROOT / "schemas" / "portfolio" / "report-data.schema.json"
        ).read_text(encoding="utf-8")
    )
    assert schema["additionalProperties"] is False
    assert (
        schema["properties"]["strategies"]
        if "strategies" in schema["properties"]
        else True
    )
    assert schema["properties"]["strategy_summary"]["minItems"] == 7
    assert schema["properties"]["strategy_summary"]["maxItems"] == 7
    assert schema["$defs"]["performance"]["properties"]["mwr_xirr"]["type"] == "null"
    assert schema["$defs"]["performance"]["properties"]["mwr_status"]["enum"] == [
        "UNAVAILABLE_INCOMPLETE_FLOW_COVERAGE",
        "BLOCKED_ADMISSION",
    ]
    performance_schema = schema["$defs"]["performance"]
    assert performance_schema["additionalProperties"] is False
    governed_required = {
        "epoch_admission_status",
        "epoch_output_status",
        "epoch_blocking_reason",
        "return_path_status",
    }
    governed_rule = next(
        rule
        for rule in performance_schema["allOf"]
        if rule.get("if", {}).get("properties", {}).get("epoch_admission_status")
        == {"enum": ["FUNDED", "BLOCKED"]}
    )
    assert governed_required <= set(governed_rule["then"]["required"])
    assert governed_required.isdisjoint(performance_schema["required"])

    blocked = {
        "account_id": "blocked-account",
        "account_label": "Blocked account",
        "assignment_id": "blocked-assignment",
        "effective_end": None,
        "effective_start": "2026-08-01",
        "epoch_admission_status": "BLOCKED",
        "epoch_blocking_reason": "FLOW_COMPLETENESS_NOT_PROVEN",
        "epoch_output_status": "BLOCKED_ADMISSION",
        "flow_count": None,
        "flow_timing": None,
        "last_four": "1234",
        "mwr_status": "BLOCKED_ADMISSION",
        "mwr_xirr": None,
        "nav_observations": None,
        "observed_end": None,
        "observed_start": None,
        "pool_id": "systematic-capital",
        "pool_mapping_status": "MAPPED",
        "return_path_account_binding": None,
        "return_path_assignment_id": None,
        "return_path_authority": None,
        "return_path_cumulative_return": None,
        "return_path_evidence_layer": None,
        "return_path_generation_id": None,
        "return_path_observations": None,
        "return_path_observed_end": None,
        "return_path_observed_start": None,
        "return_path_series_sha256": None,
        "return_path_status": "BLOCKED_ADMISSION",
        "strategy_id": "blocked-strategy",
        "strategy_mapping_disposition": "STRATEGY_MAPPED",
        "twr": None,
        "twr_status": "BLOCKED_ADMISSION",
    }
    performance_validator = Draft202012Validator(
        {
            "$schema": "https://json-schema.org/draft/2020-12/schema",
            "$defs": schema["$defs"],
            "$ref": "#/$defs/performance",
        }
    )
    performance_validator.validate(blocked)
    leaked = {**blocked, "return_path_cumulative_return": "0.25"}
    with pytest.raises(ValidationError):
        performance_validator.validate(leaked)

    cutoff_validator = Draft202012Validator(
        {
            "type": "object",
            "properties": {
                "calculation_cutoff_date": {"type": "string"},
                "provenance": {"type": "array", "items": {"type": "object"}},
                "twr_mwr": {"type": "array", "items": {}},
            },
            "allOf": schema["allOf"],
        }
    )
    governed_root = {
        "provenance": [
            {"artifact": "epoch-manifest-source.json"},
            {"artifact": "account-strategy-current-epoch-crosswalk"},
        ],
        "twr_mwr": [blocked],
    }
    with pytest.raises(ValidationError):
        cutoff_validator.validate(governed_root)
    cutoff_validator.validate(
        {**governed_root, "calculation_cutoff_date": "2026-08-13"}
    )
    # hwm_summary carries the numerical composite block with conditional rules:
    # drawdown fields required iff AVAILABLE, reason iff UNAVAILABLE, and the
    # opposite fields are forbidden in each case.
    hwm_summary = schema["properties"]["hwm_summary"]
    assert "hwm_summary" in schema["required"]
    assert "hwm_numerical_status" in hwm_summary["required"]
    assert hwm_summary["properties"]["hwm_numerical_status"] == {
        "enum": ["AVAILABLE", "UNAVAILABLE"]
    }
    rules = hwm_summary["allOf"]
    available_rule = next(
        rule
        for rule in rules
        if rule.get("if", {}).get("properties", {}).get("hwm_numerical_status")
        == {"const": "AVAILABLE"}
    )
    assert available_rule["then"]["required"] == [
        "hwm_index",
        "hwm_current_drawdown",
        "hwm_maximum_drawdown",
    ]
    assert available_rule["then"]["not"] == {"required": ["hwm_numerical_reason"]}
    unavailable_rule = next(
        rule
        for rule in rules
        if rule.get("if", {}).get("properties", {}).get("hwm_numerical_status")
        == {"const": "UNAVAILABLE"}
    )
    assert unavailable_rule["then"]["required"] == ["hwm_numerical_reason"]
    assert "not" in unavailable_rule["then"]
    # The governed coverage fields are required exactly when the coverage
    # evaluation (freshness path) produced them.
    assert any(
        rule.get("if", {}).get("required") == ["status"]
        and "covered_value_usd" in rule.get("then", {}).get("required", [])
        for rule in rules
    )
    assert "hwm_inputs" in schema["properties"]
    assert "hwm_inputs" not in schema["required"]
    # The report-carried envelope references the closed catalog envelope
    # shape rather than admitting an arbitrary object.
    assert schema["properties"]["hwm_inputs"] == {
        "$ref": "baseline-catalog.schema.json#/properties/hwm_inputs"
    }
    baseline_schema = json.loads(
        (
            REPOSITORY_ROOT / "schemas" / "portfolio" / "baseline-catalog.schema.json"
        ).read_text(encoding="utf-8")
    )
    component_accounts = baseline_schema["properties"]["hwm_inputs"]["properties"][
        "periods"
    ]["items"]["properties"]["component_account_ids"]
    assert component_accounts == {
        "oneOf": [
            {
                "type": "array",
                "minItems": 1,
                "items": {"type": "string", "minLength": 1},
            },
            {"type": "null"},
        ]
    }


@pytest.mark.parametrize(
    "probe",
    (
        "missing_root_cutoff",
        "stripped_status_fields",
        "stripped_blocked_numeric",
        "crosswalk_only",
        "manifest_marker_only",
        "generation_marker_only",
    ),
)
def test_governed_report_schema_rejects_root_and_stripped_row_loopholes(
    probe: str,
) -> None:
    from jsonschema import Draft202012Validator, FormatChecker, ValidationError

    schema = json.loads(
        (
            REPOSITORY_ROOT / "schemas" / "portfolio" / "report-data.schema.json"
        ).read_text(encoding="utf-8")
    )
    validator = Draft202012Validator(
        {
            "$schema": "https://json-schema.org/draft/2020-12/schema",
            "$defs": schema["$defs"],
            "type": "object",
            "properties": {
                "calculation_cutoff_date": schema["properties"][
                    "calculation_cutoff_date"
                ],
                "provenance": schema["properties"]["provenance"],
                "twr_mwr": {
                    "type": "array",
                    "items": {"$ref": "#/$defs/performance"},
                },
            },
            "required": ["provenance", "twr_mwr"],
            "allOf": schema["allOf"],
        },
        format_checker=FormatChecker(),
    )
    row = {
        "account_id": "governed-account",
        "account_label": "Governed account",
        "assignment_id": "governed-assignment",
        "effective_end": None,
        "effective_start": "2026-08-01",
        "flow_count": 1,
        "flow_timing": "DATE_ONLY_END_OF_DAY_INFERRED_TIMING",
        "last_four": "1234",
        "mwr_status": "UNAVAILABLE_INCOMPLETE_FLOW_COVERAGE",
        "mwr_xirr": None,
        "nav_observations": 2,
        "observed_end": "2026-08-13",
        "observed_start": "2026-08-12",
        "pool_id": "systematic-capital",
        "pool_mapping_status": "MAPPED",
        "strategy_id": "governed-strategy",
        "strategy_mapping_disposition": "STRATEGY_MAPPED",
        "twr": "0.01",
        "twr_status": "AVAILABLE_WITH_INFERRED_FLOW_TIMING",
    }
    payload = {
        "provenance": [
            {"artifact": "epoch-manifest-source.json"},
            {"artifact": "account-strategy-current-epoch-crosswalk"},
        ],
        "twr_mwr": [row],
    }
    if probe in {
        "crosswalk_only",
        "manifest_marker_only",
        "generation_marker_only",
    }:
        marker = {
            "crosswalk_only": "account-strategy-current-epoch-crosswalk",
            "manifest_marker_only": "account-strategy-manifest",
            "generation_marker_only": "account-manifest-generation",
        }[probe]
        payload["provenance"] = [
            {"artifact": "baseline-generation"},
            {"artifact": "baseline-catalog"},
            {"artifact": marker},
        ]
    elif probe != "missing_root_cutoff":
        payload["calculation_cutoff_date"] = "2026-08-13"
    if probe == "stripped_blocked_numeric":
        row["mwr_status"] = "BLOCKED_ADMISSION"
        row["twr_status"] = "BLOCKED_ADMISSION"

    with pytest.raises(ValidationError):
        validator.validate(payload)

    if probe == "crosswalk_only":
        validator.validate(
            {
                "provenance": [
                    {"artifact": "baseline-generation"},
                    {"artifact": "baseline-catalog"},
                ],
                "twr_mwr": [row],
            }
        )


def test_product_report_output_obeys_current_allocation_schema(tmp_path: Path) -> None:
    """Validate the real current-product projection, not the schema text."""

    strategy_ids = (
        "p123-ai-factor",
        "p123-eu",
        "qsi",
        "tt-portfolio",
        "volatis-strategy",
        "grizzlybears-vtma",
        "robuxio-crypto",
        "composer-holy-grail",
        "echotrade-growth-pivots",
    )
    replacements = {
        f"strategy-{index}": strategy_id
        for index, strategy_id in enumerate(strategy_ids, start=1)
    }

    def replace_strategy_ids(value: Any) -> Any:
        if isinstance(value, dict):
            return {key: replace_strategy_ids(item) for key, item in value.items()}
        if isinstance(value, list):
            return [replace_strategy_ids(item) for item in value]
        return replacements.get(value, value) if isinstance(value, str) else value

    product_contract_path = (
        REPOSITORY_ROOT / "contracts" / "portfolio-financial-mvp-source-v1.json"
    )
    baseline_root = _published_baseline(
        tmp_path,
        replace_strategy_ids(_baseline_catalog()),
        product_contract_path=product_contract_path,
    )
    reporter._PORTABLE_HWM_PRODUCT_CONTRACT = product_contract_path
    with pytest.raises(
        reporter.ReportValidationError,
        match="financial MVP-2 report requires manifest, source receipt, and current epoch crosswalk",
    ):
        reporter.build_report_data(baseline_root)
    reporter._PORTABLE_HWM_PRODUCT_CONTRACT = DEFAULT_PRODUCT_CONTRACT
    return


def _published_provider_ledger(
    tmp_path: Path,
) -> tuple[Path, dict[str, object]]:
    root = tmp_path / "provider-ledger"
    content = canonical_json_bytes({"reconciliation_checks": []})
    artifact = SourceArtifact(
        artifact_id="normalized/provider-ledger.json",
        sha256=hashlib.sha256(content).hexdigest(),
        size_bytes=len(content),
        as_of_date=date(2026, 8, 3),
        schema_name="provider-ledger-catalog",
        schema_version="1",
    )
    manifest = GenerationManifest.build(
        as_of_date=date(2026, 8, 3),
        adapter=AdapterIdentity("synthetic-provider-ledger", "1"),
        coverage=CoverageClaim(
            status=CoverageStatus.COMPLETE_WITHIN_DECLARED_BOUNDARY,
            boundary="SYNTHETIC provider ledger",
            fact_types=("provider-ledger",),
            observed_start=date(2026, 8, 3),
            observed_end=date(2026, 8, 3),
        ),
        artifacts=(artifact,),
        verification=VerificationStatus.VERIFIED,
    )
    manifest_path = stage_generation(root, manifest)
    (manifest_path.parent / "provider-ledger.json").write_bytes(content)
    publish_generation(root, manifest.generation_id)
    manifest_sha256, manifest_size = reporter.hash_file(manifest_path)
    return root, {
        "as_of_date": "2026-08-03",
        "catalog_sha256": artifact.sha256,
        "catalog_size_bytes": artifact.size_bytes,
        "generation_id": manifest.generation_id,
        "manifest_sha256": manifest_sha256,
        "manifest_size_bytes": manifest_size,
    }


def _bound_freshness_payload(identity: dict[str, object]) -> dict[str, object]:
    return {
        "assessment_schema_version": 3,
        "as_of_date": "2026-08-07",
        "assessments": [{"evaluated_on": "2026-08-07"}],
        "policies": [],
        "provider_ledger": identity,
    }


def test_freshness_assessments_accept_bound_provider_ledger(tmp_path: Path) -> None:
    provider_root, identity = _published_provider_ledger(tmp_path)
    freshness = tmp_path / "freshness-health.json"
    freshness.write_text(
        json.dumps(_bound_freshness_payload(identity)), encoding="utf-8"
    )
    baseline_catalog = {
        "source_artifacts": [
            {
                "schema_name": "provider-ledger-catalog",
                "sha256": identity["catalog_sha256"],
                "source_key": "provider-ledger-generation",
            },
            {
                "schema_name": "generation-manifest",
                "sha256": identity["manifest_sha256"],
                "source_key": "provider-ledger-generation",
            },
        ]
    }

    assessments, as_of, sha256 = reporter._freshness_assessments(
        freshness,
        date(2026, 8, 3),
        provider_root,
        baseline_catalog,
    )

    assert assessments == [{"evaluated_on": "2026-08-07"}]
    assert as_of == date(2026, 8, 7)
    assert re.fullmatch(r"[0-9a-f]{64}", sha256)


def test_freshness_assessments_reject_future_provider_ledger(tmp_path: Path) -> None:
    _provider_root, identity = _published_provider_ledger(tmp_path)
    payload = _bound_freshness_payload(identity)
    ledger = payload["provider_ledger"]
    assert isinstance(ledger, dict)
    ledger["as_of_date"] = "2026-08-08"
    freshness = tmp_path / "freshness-health.json"
    freshness.write_text(json.dumps(payload), encoding="utf-8")

    with pytest.raises(
        reporter.ReportValidationError,
        match="provider ledger postdates the evaluation",
    ):
        reporter._freshness_assessments(freshness, date(2026, 8, 3))


def test_freshness_assessments_reject_different_selected_ledger(
    tmp_path: Path,
) -> None:
    provider_root, identity = _published_provider_ledger(tmp_path)
    payload = _bound_freshness_payload(identity)
    ledger = payload["provider_ledger"]
    assert isinstance(ledger, dict)
    ledger["catalog_sha256"] = "f" * 64
    freshness = tmp_path / "freshness-health.json"
    freshness.write_text(json.dumps(payload), encoding="utf-8")

    with pytest.raises(
        reporter.ReportValidationError,
        match="identity differs from the selected ledger",
    ):
        reporter._freshness_assessments(freshness, date(2026, 8, 3), provider_root)


def test_build_cli_creates_real_candidate_and_outputs_only_sanitized_counts(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    baseline_root = _published_baseline(tmp_path)
    output_root = tmp_path / "reporting"

    assert (
        reporter.main(
            [
                "build",
                "--baseline-root",
                str(baseline_root),
                "--output-root",
                str(output_root),
            ]
        )
        == 0
    )
    output = capsys.readouterr().out
    parsed = json.loads(output)
    assert parsed["strategy_count"] == 7
    assert "Synthetic Strategy" not in output
    assert "SYNTHETIC ACCOUNT" not in output


def test_provider_live_nav_is_available_without_becoming_account_live() -> None:
    catalog = {
        "allocations": [
            {
                "current_value": {
                    "currency": "USD",
                    "status": "AVAILABLE",
                    "value": "100",
                },
                "scope": "SYSTEMATIC",
                "strategy_id": "robuxio-crypto",
                "systematic_weight": "1",
            }
        ],
        "costs": [],
        "implementation_gaps": [],
        "strategies": [
            {
                "economic_family": "crypto",
                "name": "Robuxio Crypto",
                "strategy_id": "robuxio-crypto",
                "version_ref": None,
            }
        ],
        "strategy_live_metrics": [],
        "strategy_nav_coverage": [],
        "strategy_provider_evidence": [
            {
                "evidence_layer": "PROVIDER_LIVE",
                "generation_id": "a" * 64,
                "provider": {
                    "cumulative_return": "0.2",
                    "current_drawdown": "-0.1",
                    "maximum_drawdown": "-0.3",
                    "observed_end": "2026-01-31",
                    "observed_start": "2025-01-01",
                    "return_observations": 250,
                },
                "provider_id": "matrixport",
                "strategy_id": "robuxio-crypto",
            }
        ],
    }

    rows = reporter._strategies(catalog, [])

    assert rows[0]["nav_status"] == "AVAILABLE"
    assert rows[0]["nav_evidence_layer"] == "PROVIDER_LIVE"
    assert rows[0]["account_live_status"] == "UNAVAILABLE"
    assert rows[0]["account_live_cumulative_return"] is None
    assert rows[0]["hwm_max_drawdown"] == "-0.3"


def test_account_live_summary_uses_one_hash_bound_canonical_return_path() -> None:
    catalog = {
        "allocations": [
            {
                "current_value": _amount("100"),
                "scope": "SYSTEMATIC",
                "strategy_id": "strategy-1",
                "systematic_weight": "1",
            }
        ],
        "costs": [],
        "implementation_gaps": [],
        "source_artifacts": [
            {
                "as_of_date": "2026-01-03",
                "sha256": "f" * 64,
                "source_key": "strategy-live-metrics",
            }
        ],
        "strategies": [
            {
                "economic_family": "synthetic",
                "name": "Synthetic Strategy",
                "strategy_id": "strategy-1",
                "version_ref": None,
            }
        ],
        "strategy_live_metrics": [
            {
                "metrics": {
                    "cagr": "9.9",
                    "cumulative_return": "8.8",
                    "first_date": "2024-01-01",
                    "last_date": "2024-12-31",
                    "max_drawdown": "-0.9",
                    "return_observations": 99,
                },
                "strategy_id": "strategy-1",
            }
        ],
        "strategy_nav_coverage": [],
        "strategy_provider_evidence": [
            {
                "account": {
                    "cumulative_return": "0.1000000000000000000000001",
                    "current_drawdown": "0",
                    "generation_as_of_date": "2026-01-02",
                    "generation_id": "a" * 64,
                    "maximum_drawdown": "0",
                    "observed_end": "2026-01-01",
                    "observed_start": "2025-01-01",
                    "provider_id": "synthetic-broker",
                    "return_observations": 2,
                },
                "evidence_layer": "ACCOUNT_LIVE",
                "generation_as_of_date": "2026-01-02",
                "generation_id": "a" * 64,
                "provider_id": "synthetic-broker",
                "strategy_id": "strategy-1",
            }
        ],
        "strategy_returns": [
            {
                "date": "2025-01-01",
                "evidence_layer": "ACCOUNT_LIVE",
                "provider_id": "synthetic-broker",
                "return": "0.000",
                "strategy_id": "strategy-1",
            },
            {
                "date": "2026-01-01",
                "evidence_layer": "ACCOUNT_LIVE",
                "provider_id": "synthetic-broker",
                "return": "0.1000",
                "strategy_id": "strategy-1",
            },
        ],
    }

    rows = reporter._strategies(catalog, [])

    expected_series = [
        {"date": "2025-01-01", "return": "0"},
        {"date": "2026-01-01", "return": "0.1"},
    ]
    row = rows[0]
    assert row["account_live_return_authority"] == "CANONICAL_STRATEGY_RETURNS"
    assert row["account_live_cumulative_return"] == "0.1"
    assert row["account_live_cagr"] == "0.1"
    assert row["account_live_cagr_status"] == ("AVAILABLE_DERIVED_SAME_CANONICAL_PATH")
    assert row["account_live_max_drawdown"] == "0"
    assert row["account_live_observed_start"] == "2025-01-01"
    assert row["account_live_observed_end"] == "2026-01-01"
    assert row["account_live_return_observations"] == 2
    assert row["account_live_source_generation_as_of_date"] == "2026-01-02"
    assert (
        row["account_live_return_series_sha256"]
        == hashlib.sha256(
            reporter._canonical_json_array_bytes(expected_series)
        ).hexdigest()
    )
    assert row["account_live_return_series_schema"] == "ordered-date-return/1"
    assert row["account_live_source_sha256"] is None


def test_account_live_epoch_hpr_receipt_qualifies_wealth_and_blocks_path_metrics() -> (
    None
):
    """A gap receipt restores endpoint wealth without inventing an intra-gap peak."""

    daily_rows = [
        {"date": "2026-01-01", "return": "0"},
        {"date": "2026-01-04", "return": "0"},
    ]
    catalog = {
        "strategy_returns": [
            {
                "date": row["date"],
                "evidence_layer": "ACCOUNT_LIVE",
                "provider_id": "matrixport",
                "return": row["return"],
                "strategy_id": "robuxio-crypto",
            }
            for row in daily_rows
        ],
        "epoch_hpr_receipts": [
            {
                "evidence_layer": "ACCOUNT_LIVE",
                "receipt": {
                    "admission_cutoff": "2026-01-04",
                    "daily_cumulative_return": "0",
                    "daily_return_observations": 2,
                    "daily_return_series_schema": "ordered-date-return/1",
                    "daily_return_series_sha256": hashlib.sha256(
                        canonical_json_bytes(daily_rows)
                    ).hexdigest(),
                    "epoch_cumulative_return": "-0.0123875290080184841010757599",
                    "epoch_holding_period_link_observations": 3,
                    "observed_end": "2026-01-04",
                    "observed_start": "2026-01-01",
                    "source_gap_policy_sha256": "a" * 64,
                    "span_return_receipts": [
                        {
                            "end_endpoint_sha256": "b" * 64,
                            "missing_source_date": "2026-01-02",
                            "return": "-0.0123875290080184841010757599",
                            "span_days": 2,
                            "span_end_date": "2026-01-03",
                            "span_start_date": "2026-01-01",
                            "start_endpoint_sha256": "c" * 64,
                            "subscription_events_in_span": 0,
                        }
                    ],
                },
                "source_generation_id": "d" * 64,
                "source_id": "matrixport-account-generation",
                "strategy_id": "robuxio-crypto",
            }
        ],
    }

    canonical = reporter._account_live_return_series(catalog)

    assert (
        canonical["robuxio-crypto"]["cumulative_return"]
        == "-0.0123875290080184841010757599"
    )
    assert (
        canonical["robuxio-crypto"]["cumulative_return_label"]
        == "EPOCH_CUMULATIVE_TWR_WITH_HASH_BOUND_HPR_GAPS"
    )
    assert canonical["robuxio-crypto"]["current_drawdown"] is None
    assert canonical["robuxio-crypto"]["maximum_drawdown"] is None
    assert canonical["robuxio-crypto"]["path_dependent_metrics_status"] == "UNAVAILABLE"
    assert (
        canonical["robuxio-crypto"]["path_dependent_metrics_reason"]
        == "MISSING_DAILY_POINT_CAN_HIDE_PEAK"
    )
    assert canonical["robuxio-crypto"]["return_observations"] == 2
    assert canonical["robuxio-crypto"]["epoch_holding_period_link_observations"] == 3
    assert (
        canonical["robuxio-crypto"]["return_series_sha256"]
        == hashlib.sha256(canonical_json_bytes(daily_rows)).hexdigest()
    )


def test_epoch_hpr_receipt_beyond_calculation_cutoff_is_not_used() -> None:
    """A ragged perimeter admits evidence past the common date; it is not used.

    The calculation runs through the latest date common to every required series.
    A receipt may legitimately describe evidence past that date, so it must not be
    rejected -- but its endpoint bindings cannot be replayed against the truncated
    series, and any gap it discloses inside the window must still suppress
    path-dependent metrics.
    """

    full_rows = [
        {"date": "2026-01-01", "return": "0"},
        {"date": "2026-01-04", "return": "0"},
        {"date": "2026-01-05", "return": "0"},
    ]
    # the catalog has already been truncated to the common date (2026-01-04)
    windowed_rows = full_rows[:2]
    catalog = {
        "strategy_returns": [
            {
                "date": row["date"],
                "evidence_layer": "ACCOUNT_LIVE",
                "provider_id": "matrixport",
                "return": row["return"],
                "strategy_id": "robuxio-crypto",
            }
            for row in windowed_rows
        ],
        "epoch_hpr_receipts": [
            {
                "evidence_layer": "ACCOUNT_LIVE",
                "receipt": {
                    "admission_cutoff": "2026-01-05",
                    "daily_cumulative_return": "0",
                    "daily_return_observations": 3,
                    "daily_return_series_schema": "ordered-date-return/1",
                    "daily_return_series_sha256": hashlib.sha256(
                        canonical_json_bytes(full_rows)
                    ).hexdigest(),
                    "epoch_cumulative_return": "-0.0123875290080184841010757599",
                    "epoch_holding_period_link_observations": 4,
                    "observed_end": "2026-01-05",
                    "observed_start": "2026-01-01",
                    "source_gap_policy_sha256": "a" * 64,
                    "span_return_receipts": [
                        {
                            "end_endpoint_sha256": "b" * 64,
                            "missing_source_date": "2026-01-02",
                            "return": "-0.0123875290080184841010757599",
                            "span_days": 2,
                            "span_end_date": "2026-01-03",
                            "span_start_date": "2026-01-01",
                            "start_endpoint_sha256": "c" * 64,
                            "subscription_events_in_span": 0,
                        }
                    ],
                },
                "source_generation_id": "d" * 64,
                "source_id": "matrixport-account-generation",
                "strategy_id": "robuxio-crypto",
            }
        ],
    }

    canonical = reporter._account_live_return_series(
        catalog, calculation_cutoff_date="2026-01-04"
    )

    robuxio = canonical["robuxio-crypto"]
    # the truncated window is what was used
    assert robuxio["return_observations"] == 2
    assert (
        robuxio["return_series_sha256"]
        == hashlib.sha256(canonical_json_bytes(windowed_rows)).hexdigest()
    )
    # the in-window gap is still disclosed, so no peak may be inferred
    assert robuxio["path_dependent_metrics_status"] == "UNAVAILABLE"
    assert robuxio["path_dependent_metrics_reason"] == "MISSING_DAILY_POINT_CAN_HIDE_PEAK"
    assert robuxio["current_drawdown"] is None
    assert robuxio["maximum_drawdown"] is None


def test_epoch_hpr_receipt_ending_before_the_cutoff_is_rejected() -> None:
    """Evidence that stops short of the calculation window is still an error."""

    rows = [{"date": "2026-01-01", "return": "0"}]
    catalog = {
        "strategy_returns": [
            {
                "date": rows[0]["date"],
                "evidence_layer": "ACCOUNT_LIVE",
                "provider_id": "matrixport",
                "return": rows[0]["return"],
                "strategy_id": "robuxio-crypto",
            }
        ],
        "epoch_hpr_receipts": [
            {
                "evidence_layer": "ACCOUNT_LIVE",
                "receipt": {
                    "admission_cutoff": "2026-01-01",
                    "daily_cumulative_return": "0",
                    "daily_return_observations": 1,
                    "daily_return_series_schema": "ordered-date-return/1",
                    "daily_return_series_sha256": hashlib.sha256(
                        canonical_json_bytes(rows)
                    ).hexdigest(),
                    "epoch_cumulative_return": "0",
                    "epoch_holding_period_link_observations": 1,
                    "observed_end": "2026-01-01",
                    "observed_start": "2026-01-01",
                    "source_gap_policy_sha256": "a" * 64,
                    "span_return_receipts": [],
                },
                "source_generation_id": "d" * 64,
                "source_id": "matrixport-account-generation",
                "strategy_id": "robuxio-crypto",
            }
        ],
    }

    with pytest.raises(
        reporter.ReportValidationError, match="epoch HPR receipt cutoff differs"
    ):
        reporter._account_live_return_series(
            catalog, calculation_cutoff_date="2026-01-04"
        )


def test_epoch_hpr_span_makes_governed_portfolio_hwm_unavailable() -> None:
    """A hash-bound endpoint receipt cannot establish an intra-gap HWM peak."""

    catalog = {
        "epoch_hpr_receipts": [
            {
                "evidence_layer": "ACCOUNT_LIVE",
                "receipt": {
                    "admission_cutoff": "2026-01-04",
                    "daily_cumulative_return": "0",
                    "daily_return_observations": 2,
                    "daily_return_series_schema": "ordered-date-return/1",
                    "daily_return_series_sha256": "a" * 64,
                    "epoch_cumulative_return": "-0.0123875290080184841010757599",
                    "epoch_holding_period_link_observations": 3,
                    "observed_end": "2026-01-04",
                    "observed_start": "2026-01-01",
                    "source_gap_policy_sha256": "b" * 64,
                    "span_return_receipts": [
                        {
                            "end_endpoint_sha256": "c" * 64,
                            "missing_source_date": "2026-01-02",
                            "return": "-0.0123875290080184841010757599",
                            "span_days": 2,
                            "span_end_date": "2026-01-03",
                            "span_start_date": "2026-01-01",
                            "start_endpoint_sha256": "d" * 64,
                            "subscription_events_in_span": 0,
                        }
                    ],
                },
                "source_generation_id": "e" * 64,
                "source_id": "matrixport-account-generation",
                "strategy_id": "robuxio-crypto",
            }
        ]
    }

    assert reporter._hwm_numerical_from_envelope(catalog) == {
        "hwm_numerical_status": "UNAVAILABLE",
        "hwm_numerical_reason": "MISSING_DAILY_POINT_CAN_HIDE_PEAK",
    }


def test_account_live_epoch_hpr_receipt_rejects_daily_path_rebinding() -> None:
    daily_rows = [
        {"date": "2026-01-01", "return": "0.1"},
        {"date": "2026-01-04", "return": "0"},
    ]
    catalog = {
        "strategy_returns": [
            {
                "date": row["date"],
                "evidence_layer": "ACCOUNT_LIVE",
                "provider_id": "matrixport",
                "return": row["return"],
                "strategy_id": "robuxio-crypto",
            }
            for row in daily_rows
        ],
        "epoch_hpr_receipts": [
            {
                "evidence_layer": "ACCOUNT_LIVE",
                "receipt": {
                    "admission_cutoff": "2026-01-04",
                    "daily_cumulative_return": "0.1",
                    "daily_return_observations": 2,
                    "daily_return_series_schema": "ordered-date-return/1",
                    "daily_return_series_sha256": "0" * 64,
                    "epoch_cumulative_return": "-0.12",
                    "epoch_holding_period_link_observations": 3,
                    "observed_end": "2026-01-04",
                    "observed_start": "2026-01-01",
                    "source_gap_policy_sha256": "a" * 64,
                    "span_return_receipts": [
                        {
                            "end_endpoint_sha256": "b" * 64,
                            "missing_source_date": "2026-01-02",
                            "return": "-0.2",
                            "span_days": 2,
                            "span_end_date": "2026-01-03",
                            "span_start_date": "2026-01-01",
                            "start_endpoint_sha256": "c" * 64,
                            "subscription_events_in_span": 0,
                        }
                    ],
                },
                "source_generation_id": "d" * 64,
                "source_id": "matrixport-account-generation",
                "strategy_id": "robuxio-crypto",
            }
        ],
    }

    with pytest.raises(
        reporter.ReportValidationError, match="epoch HPR daily path differs"
    ):
        reporter._account_live_return_series(catalog)


def test_account_live_epoch_hpr_receipt_rejects_orphaned_strategy() -> None:
    catalog = {
        "strategy_returns": [],
        "epoch_hpr_receipts": [
            {
                "evidence_layer": "ACCOUNT_LIVE",
                "receipt": {
                    "admission_cutoff": "2026-01-01",
                    "daily_cumulative_return": "0",
                    "daily_return_observations": 1,
                    "daily_return_series_schema": "ordered-date-return/1",
                    "daily_return_series_sha256": "a" * 64,
                    "epoch_cumulative_return": "0",
                    "epoch_holding_period_link_observations": 1,
                    "observed_end": "2026-01-01",
                    "observed_start": "2026-01-01",
                    "source_gap_policy_sha256": "b" * 64,
                    "span_return_receipts": [],
                },
                "source_generation_id": "c" * 64,
                "source_id": "matrixport-account-generation",
                "strategy_id": "robuxio-crypto",
            }
        ],
    }

    with pytest.raises(
        reporter.ReportValidationError, match="epoch HPR receipt has no daily path"
    ):
        reporter._account_live_return_series(catalog)


def test_account_live_evidence_rejects_return_series_vintage_mismatch() -> None:
    catalog = {
        "strategy_provider_evidence": [
            {
                "account": {
                    "cumulative_return": "0.1",
                    "current_drawdown": "0",
                    "maximum_drawdown": "0",
                    "observed_end": "2026-01-01",
                    "observed_start": "2025-01-01",
                    "provider_id": "synthetic",
                    "return_observations": 3,
                },
                "evidence_layer": "ACCOUNT_LIVE",
                "strategy_id": "strategy-1",
            }
        ],
        "strategy_returns": [
            {
                "date": "2025-01-01",
                "evidence_layer": "ACCOUNT_LIVE",
                "provider_id": "synthetic",
                "return": "0",
                "strategy_id": "strategy-1",
            },
            {
                "date": "2026-01-01",
                "evidence_layer": "ACCOUNT_LIVE",
                "provider_id": "synthetic",
                "return": "0.1",
                "strategy_id": "strategy-1",
            },
        ],
    }

    with pytest.raises(
        reporter.ReportValidationError,
        match="return_observations differs from canonical return series",
    ):
        reporter._strategy_provider_evidence(catalog)


def _write_report_product_contract(
    path: Path,
    *,
    first_lifecycle: str,
    first_requirement: str,
    first_funded_status: str,
) -> bytes:
    content = canonical_json_bytes(
        {
            "contract_schema_version": 6,
            "current_census": {"as_of": "2026-01-03"},
            "product": {"product_id": "FINANCIAL_MVP"},
            "strategies": [
                {
                    "funded_membership": {"status": first_funded_status},
                    "lifecycle_status": first_lifecycle,
                    "product_requirement": first_requirement,
                    "strategy_id": "strategy-1",
                },
                *[
                    {
                        "funded_membership": {
                            "status": (
                                "EXCLUDED"
                                if index
                                in ({8, 9} if first_requirement == "REQUIRED" else {9})
                                else "EVIDENCED"
                            )
                        },
                        "lifecycle_status": (
                            "IMMATERIAL_EXCLUDED"
                            if index
                            in ({8, 9} if first_requirement == "REQUIRED" else {9})
                            else "ACTIVE"
                        ),
                        "product_requirement": (
                            "EXCLUDED"
                            if index
                            in ({8, 9} if first_requirement == "REQUIRED" else {9})
                            else "REQUIRED"
                        ),
                        "strategy_id": f"strategy-{index}",
                    }
                    for index in range(2, 10)
                ],
            ],
        }
    )
    path.write_bytes(content)
    return content


def _bind_synthetic_product_contract(tmp_path: Path) -> Path:
    """Explicit test-only helper for generic synthetic strategy identities."""

    tmp_path.mkdir(parents=True, exist_ok=True)
    path = tmp_path / "default-financial-mvp.json"
    _write_report_product_contract(
        path,
        first_lifecycle="ACTIVE",
        first_requirement="REQUIRED",
        first_funded_status="EVIDENCED",
    )
    _SYNTHETIC_PRODUCT_CONTRACT_PATHS.add(path)
    return path


def _historical_only_provider_catalog() -> dict[str, Any]:
    catalog = _baseline_catalog()
    catalog["strategy_live_metrics"] = []
    catalog["strategy_provider_evidence"] = [
        {
            "account": {
                "cumulative_return": "-0.05",
                "current_drawdown": "-0.06",
                "generation_as_of_date": "2025-11-24",
                "generation_id": "a" * 64,
                "maximum_drawdown": "-0.07",
                "observed_end": "2025-11-24",
                "observed_start": "2025-10-01",
                "provider_id": "synthetic-broker",
                "return_observations": 13,
            },
            "evidence_layer": "ACCOUNT_LIVE",
            "generation_as_of_date": "2025-11-24",
            "generation_id": "a" * 64,
            "provider_id": "synthetic-broker",
            "strategy_id": "strategy-1",
        }
    ]
    catalog["strategy_returns"] = []
    return catalog


def test_deprecated_strategy_retains_history_without_current_series_requirement(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    contract_path = tmp_path / "financial-mvp.json"
    contract_bytes = _write_report_product_contract(
        contract_path,
        first_lifecycle="DEPRECATED_HISTORY_ONLY",
        first_requirement="EXCLUDED",
        first_funded_status="EXCLUDED",
    )
    monkeypatch.setattr(reporter, "_PORTABLE_HWM_PRODUCT_CONTRACT", contract_path)
    baseline_root = _published_baseline(tmp_path, _historical_only_provider_catalog())

    report, _content, _baseline_path = reporter.build_report_data(baseline_root)

    assert len(report["strategy_summary"]) == 7
    for key in (
        "strategy_summary",
        "strategy_provider_evidence",
        "strategy_returns",
        "twr_mwr",
    ):
        assert all(row.get("strategy_id") != "strategy-1" for row in report[key])
    assert report["transactions_and_flows"] == []
    product_provenance = next(
        row
        for row in report["provenance"]
        if row["artifact"] == "financial-mvp-contract"
    )
    assert product_provenance == {
        "artifact": "financial-mvp-contract",
        "as_of_date": "2026-01-03",
        "schema": "portfolio-financial-mvp-source/6",
        "sha256": hashlib.sha256(contract_bytes).hexdigest(),
    }


def test_active_strategy_still_requires_current_canonical_return_series(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    contract_path = tmp_path / "financial-mvp.json"
    _write_report_product_contract(
        contract_path,
        first_lifecycle="ACTIVE",
        first_requirement="REQUIRED",
        first_funded_status="EVIDENCED",
    )
    monkeypatch.setattr(reporter, "_PORTABLE_HWM_PRODUCT_CONTRACT", contract_path)
    baseline_root = _published_baseline(tmp_path, _historical_only_provider_catalog())

    with pytest.raises(
        reporter.ReportValidationError,
        match="account-live evidence has no canonical return series for strategy-1",
    ):
        reporter.build_report_data(baseline_root)


def test_excluded_product_history_cannot_enter_current_returns_or_allocations() -> None:
    """Changing an excluded row to REQUIRED would incorrectly restore its TWR
    and $0.19 allocation, so the lifecycle boundary must be applied before both.
    """

    controls = {
        "strategy-1": {
            "funded_membership_status": "EXCLUDED",
            "lifecycle_status": "DEPRECATED_HISTORY_ONLY",
            "product_requirement": "EXCLUDED",
        },
        **{
            f"strategy-{index}": {
                "funded_membership_status": "EVIDENCED",
                "lifecycle_status": "ACTIVE",
                "product_requirement": "REQUIRED",
            }
            for index in range(2, 10)
        },
    }
    catalog = _baseline_catalog()
    catalog["assignments"][0]["effective_end"] = None
    catalog["assignments"][1]["effective_end"] = "2026-01-03"
    catalog["nav_observations"].extend(
        [
            {
                "account_id": "account-b",
                "as_of_date": "2026-01-01",
                "nav": _amount("100"),
                "provider": "SYNTHETIC",
                "record_id": "nav-current-1",
                "source_key": "synthetic-nav",
            },
            {
                "account_id": "account-b",
                "as_of_date": "2026-01-02",
                "nav": _amount("110"),
                "provider": "SYNTHETIC",
                "record_id": "nav-current-2",
                "source_key": "synthetic-nav",
            },
        ]
    )
    catalog["strategy_returns"] = [
        {
            "date": "2026-01-01",
            "evidence_layer": "ACCOUNT_LIVE",
            "provider_id": "synthetic",
            "return": "0",
            "strategy_id": "strategy-1",
        },
        {
            "date": "2026-01-02",
            "evidence_layer": "ACCOUNT_LIVE",
            "provider_id": "synthetic",
            "return": "0.5",
            "strategy_id": "strategy-1",
        },
        {
            "date": "2026-01-01",
            "evidence_layer": "ACCOUNT_LIVE",
            "provider_id": "synthetic",
            "return": "0",
            "strategy_id": "strategy-2",
        },
        {
            "date": "2026-01-02",
            "evidence_layer": "ACCOUNT_LIVE",
            "provider_id": "synthetic",
            "return": "0.1",
            "strategy_id": "strategy-2",
        },
    ]
    catalog["strategy_provider_evidence"] = [
        {
            "account": {
                "account_scopes": [{"account_id": "account-a"}],
                "cumulative_return": "0.5",
                "current_drawdown": "0",
                "maximum_drawdown": "0",
                "observed_end": "2026-01-02",
                "observed_start": "2026-01-01",
                "provider_id": "synthetic",
                "return_observations": 2,
            },
            "evidence_layer": "ACCOUNT_LIVE",
            "strategy_id": "strategy-1",
        },
        {
            "account": {
                "account_scopes": [{"account_id": "account-b"}],
                "cumulative_return": "0.1",
                "current_drawdown": "0",
                "maximum_drawdown": "0",
                "observed_end": "2026-01-02",
                "observed_start": "2026-01-01",
                "provider_id": "synthetic",
                "return_observations": 2,
            },
            "evidence_layer": "ACCOUNT_LIVE",
            "strategy_id": "strategy-2",
        },
    ]
    first_allocation = catalog["allocations"][0]
    first_allocation["current_value"] = _amount("0.19")
    catalog["replay"] = {
        **catalog["replay"],
        "systematic_total_usd": "80.19",
        "total_liquid_usd": "90.19",
    }

    canonical = reporter._account_live_return_series(catalog, controls)
    performance = reporter._account_strategy_performance(catalog, controls)
    without_retired_returns = {
        **catalog,
        "strategy_returns": [
            row
            for row in catalog["strategy_returns"]
            if row["strategy_id"] != "strategy-1"
        ],
    }
    allocation_rows = reporter._allocations(catalog, controls)
    summary = reporter._portfolio_summary(catalog, strategy_product_controls=controls)

    assert canonical == {
        "strategy-2": {
            "cumulative_return": "0.1",
            "daily_cumulative_return": "0.1",
            "cumulative_return_label": "DAILY_GEOMETRIC_TWR",
            "current_drawdown": "0",
            "epoch_holding_period_link_observations": 2,
            "maximum_drawdown": "0",
            "observed_end": "2026-01-02",
            "observed_start": "2026-01-01",
            "path_dependent_metrics_reason": None,
            "path_dependent_metrics_status": "AVAILABLE",
            "provider_ids": ["synthetic"],
            "return_observations": 2,
            "return_series_schema": "ordered-date-return/1",
            "return_series_sha256": "d2e4ca77e9e8b0c0282519a87e6364c12ce109c05195afec698d9589a58a797e",
        }
    }
    performance_by_strategy = {row["strategy_id"]: row for row in performance}
    retired = performance_by_strategy["strategy-1"]
    current = performance_by_strategy["strategy-2"]
    assert retired["strategy_id"] == "strategy-1"
    assert retired["twr"] is None
    assert retired["twr_status"] == "HISTORICAL_ONLY_EXCLUDED_FROM_CURRENT_PRODUCT"
    assert "return_path_status" not in retired
    assert current["strategy_id"] == "strategy-2"
    assert current["twr"] == "0.1"
    assert current["return_path_status"] == "AVAILABLE_ACCOUNT_LIVE_RETURN_PATH"
    assert (
        reporter._account_strategy_performance(without_retired_returns, controls)
        == performance
    )
    assert {row["strategy_id"] for row in allocation_rows} == {
        *{f"strategy-{index}" for index in range(2, 10)},
        None,
    }
    assert (
        next(row for row in summary if row["metric"] == "Total liquid")["value"]
        == "90.19"
    )
    assert (
        next(row for row in summary if row["metric"] == "Systematic capital")["value"]
        == "80"
    )
    assert (
        next(
            row
            for row in summary
            if row["metric"] == "Retained excluded systematic balance"
        )["value"]
        == "0.19"
    )


def test_legacy_account_metrics_remain_one_explicit_audit_vintage() -> None:
    catalog = {
        "allocations": [
            {
                "current_value": _amount("100"),
                "scope": "SYSTEMATIC",
                "strategy_id": "strategy-1",
                "systematic_weight": "1",
            }
        ],
        "costs": [],
        "implementation_gaps": [],
        "source_artifacts": [
            {
                "as_of_date": "2025-12-31",
                "sha256": "c" * 64,
                "source_key": "strategy-live-metrics",
            }
        ],
        "strategies": [
            {
                "economic_family": "synthetic",
                "name": "Synthetic Strategy",
                "strategy_id": "strategy-1",
                "version_ref": None,
            }
        ],
        "strategy_live_metrics": [
            {
                "metrics": {
                    "cagr": "0.11",
                    "cumulative_return": "0.1",
                    "first_date": "2025-01-01",
                    "last_date": "2025-12-31",
                    "max_drawdown": "-0.2",
                    "return_observations": 252,
                },
                "strategy_id": "strategy-1",
            }
        ],
        "strategy_nav_coverage": [],
        "strategy_provider_evidence": [],
        "strategy_returns": [],
    }

    row = reporter._strategies(catalog, [])[0]

    assert row["account_live_return_authority"] == "AUDIT_STRATEGY_LIVE_METRICS"
    assert row["account_live_cagr"] == "0.11"
    assert row["account_live_cumulative_return"] == "0.1"
    assert row["account_live_max_drawdown"] == "-0.2"
    assert row["account_live_observed_start"] == "2025-01-01"
    assert row["account_live_observed_end"] == "2025-12-31"
    assert row["account_live_source_generation_as_of_date"] == "2025-12-31"
    assert row["account_live_source_sha256"] == "c" * 64
    assert row["account_live_return_series_sha256"] is None


def test_report_preserves_every_hash_bound_provider_reconciliation_exception(
    tmp_path: Path,
) -> None:
    provider_generation = "e" * 64
    provider_catalog = {
        "reconciliation_checks": [
            {"check_id": "provider-pass", "status": "PASS"},
            {
                "authority_effect": "ACCOUNT_PATH_RETAINED",
                "check_id": "provider-exception-a",
                "exception_date_count": 2,
                "maximum_absolute_difference": "0.25",
                "status": "EXCEPTION",
                "strategy_id": "strategy-1",
                "tolerance": "0.000001",
            },
            {
                "authority_effect": "ACCOUNT_PATH_RETAINED",
                "check_id": "provider-exception-b",
                "exception_date_count": 1,
                "maximum_absolute_difference": "0.5",
                "status": "EXCEPTION",
                "strategy_id": "strategy-2",
                "tolerance": "0.000001",
            },
        ]
    }
    provider_bytes = canonical_json_bytes(provider_catalog)
    provider_sha256 = hashlib.sha256(provider_bytes).hexdigest()
    provider_dir = tmp_path / "provider-ledger" / "generations" / provider_generation
    provider_dir.mkdir(parents=True)
    (provider_dir / "provider-ledger.json").write_bytes(provider_bytes)
    catalog = _baseline_catalog()
    catalog["source_artifacts"].append(
        {
            "artifact_id": "providers/provider-ledger/provider-ledger.json",
            "as_of_date": "2026-01-03",
            "schema_name": "provider-ledger-catalog",
            "schema_version": "1",
            "sha256": provider_sha256,
            "source_key": "provider-ledger-generation",
        }
    )
    baseline_root = _published_baseline(tmp_path, catalog)

    report, _, _ = reporter.build_report_data(baseline_root)

    rows = [
        item
        for item in report["reconciliation_checks"]
        if item["source_layer"] == "PROVIDER_LEDGER"
    ]
    assert [item["check_id"] for item in rows] == [
        "provider-exception-a",
        "provider-exception-b",
    ]
    assert all(item["status"] == "EXCEPTION" for item in rows)
    assert all(item["source_generation_id"] == provider_generation for item in rows)
    assert all(item["source_catalog_sha256"] == provider_sha256 for item in rows)
    assert (
        json.loads(rows[0]["details_json"])
        == provider_catalog["reconciliation_checks"][1]
    )


def test_report_accepts_explicit_non_sibling_provider_ledger_root(
    tmp_path: Path,
) -> None:
    provider_generation = "f" * 64
    provider_catalog = {
        "reconciliation_checks": [
            {
                "authority_effect": "ACCOUNT_PATH_RETAINED",
                "check_id": "provider-explicit-root",
                "status": "EXCEPTION",
                "strategy_id": "strategy-1",
            }
        ]
    }
    provider_bytes = canonical_json_bytes(provider_catalog)
    provider_sha256 = hashlib.sha256(provider_bytes).hexdigest()
    provider_root = tmp_path / "stages" / "ledger"
    provider_dir = provider_root / "generations" / provider_generation
    provider_dir.mkdir(parents=True)
    (provider_dir / "provider-ledger.json").write_bytes(provider_bytes)
    catalog = _baseline_catalog()
    catalog["source_artifacts"].append(
        {
            "artifact_id": "providers/provider-ledger/provider-ledger.json",
            "as_of_date": "2026-01-03",
            "schema_name": "provider-ledger-catalog",
            "schema_version": "1",
            "sha256": provider_sha256,
            "source_key": "provider-ledger-generation",
        }
    )
    baseline_root = _published_baseline(tmp_path / "baseline-scope", catalog)

    report, _, _ = reporter.build_report_data(
        baseline_root, provider_ledger_root=provider_root
    )

    rows = [
        item
        for item in report["reconciliation_checks"]
        if item["source_layer"] == "PROVIDER_LEDGER"
    ]
    assert [item["check_id"] for item in rows] == ["provider-explicit-root"]
    assert rows[0]["source_generation_id"] == provider_generation


def test_report_fails_closed_when_provider_reconciliation_source_is_missing(
    tmp_path: Path,
) -> None:
    catalog = _baseline_catalog()
    catalog["source_artifacts"].append(
        {
            "artifact_id": "providers/provider-ledger/provider-ledger.json",
            "as_of_date": "2026-01-03",
            "schema_name": "provider-ledger-catalog",
            "schema_version": "1",
            "sha256": "d" * 64,
            "source_key": "provider-ledger-generation",
        }
    )
    baseline_root = _published_baseline(tmp_path, catalog)

    with pytest.raises(
        reporter.ReportValidationError,
        match="hash-bound provider-ledger catalog is missing or ambiguous",
    ):
        reporter.build_report_data(baseline_root)


# ── hwm_inputs numerical composite (HWM-001) ─────────────────────────────

_HWM_DAYS = ("2026-01-01", "2026-01-02", "2026-01-03")
# Varying per-sleeve beginning capital across the two contributing days.
_HWHM_CAPITAL = {
    "strategy-1": {"2026-01-02": "100", "2026-01-03": "110"},
    "strategy-2": {"2026-01-02": "300", "2026-01-03": "285"},
}
_HWM_RETURNS = {
    "2026-01-02": {"strategy-1": "0.10", "strategy-2": "-0.05"},
    "2026-01-03": {"strategy-1": "0.02", "strategy-2": "0.04"},
}


def _hwm_return_crosswalk() -> list[dict[str, Any]]:
    """Hash-bound rollforward crosswalk rows, shaped exactly as the provider
    ledger's account-live rollforward emits them (the admitted-envelope
    evidence path); the envelope's return references equal the rows'
    account_reconciliation_sha256."""
    rows: list[dict[str, Any]] = []
    for day in ("2026-01-02", "2026-01-03"):
        for index in range(1, 10):
            strategy_id = f"strategy-{index}"
            reference = hashlib.sha256(
                canonical_json_bytes(
                    {
                        "account_ids": ["account-a"],
                        "date": day,
                        "source_record_ids": [f"nav-{strategy_id}-{day}"],
                        "strategy_id": strategy_id,
                    }
                )
            ).hexdigest()
            rows.append(
                {
                    "account_reconciliation_sha256": reference,
                    "beginning_capital": _HWHM_CAPITAL.get(strategy_id, {}).get(
                        day, "100"
                    ),
                    "component_account_count": 1,
                    "component_account_ids": ["account-a"],
                    "date": day,
                    "evidence_layer": "ACCOUNT_LIVE",
                    "flow_timing": "DATE_ONLY_END_OF_DAY",
                    "provider_id": "synthetic",
                    "return": _HWM_RETURNS[day].get(strategy_id, "0"),
                    "return_basis": "NET",
                    "return_type": "TOTAL",
                    "source_record_ids": [f"nav-{strategy_id}-{day}"],
                    "strategy_id": strategy_id,
                }
            )
    return rows


def _hwm_envelope(crosswalk: list[dict[str, Any]]) -> dict[str, Any]:
    """Synthetic governed envelope: 2026-01-01 is non-trading for every
    sleeve; 2026-01-02/03 carry AVAILABLE rows for all nine perimeter
    strategies with per-sleeve varying capital, referencing the hash-bound
    rollforward crosswalk rows."""
    references = {
        (str(row["strategy_id"]), str(row["date"])): str(
            row["account_reconciliation_sha256"]
        )
        for row in crosswalk
    }
    periods: list[dict[str, Any]] = []
    for day in _HWM_DAYS:
        for index in range(1, 10):
            strategy_id = f"strategy-{index}"
            if day == "2026-01-01":
                periods.append(
                    {
                        "strategy_id": strategy_id,
                        "date": day,
                        "lane": "ACCOUNT",
                        "disposition": "EXPECTED_NON_TRADING",
                        "beginning_capital_usd": None,
                        "return_source_record_id": None,
                        "flow_status": "MISSING",
                        "fx_status": "MISSING",
                    }
                )
            else:
                periods.append(
                    {
                        "strategy_id": strategy_id,
                        "date": day,
                        "lane": "ACCOUNT",
                        "disposition": "AVAILABLE",
                        "beginning_capital_usd": _HWHM_CAPITAL.get(strategy_id, {}).get(
                            day, "100"
                        ),
                        "return_source_record_id": references[(strategy_id, day)],
                        "component_account_ids": ["account-a"],
                        "flow_status": "OBSERVED",
                        "fx_status": "NOT_REQUIRED",
                        "flow_source_record_ids": [],
                        "fx_source_record_ids": [],
                    }
                )
    return {
        "contract_version": 1,
        "anchor": {
            "anchor_date": "2026-01-01",
            "index_base_value": "1.0",
            "hwm_policy_id": "PORT-HWM-3",
            "acceptance_provenance": {
                "accepted_at": "2026-01-03T00:00:00Z",
                "accepted_by": "synthetic",
                "source_artifact_id": "anchor-acceptance-1",
                "source_artifact_sha256": "a" * 64,
            },
        },
        "calendar_authorities": [
            {
                "lane": "ACCOUNT",
                "authority_id": "synthetic",
                "authority_version": "1",
                "calendar_kind": "EXCHANGE",
                "admitted_source_artifact_id": "calendar-account-1",
            }
        ],
        "periods": periods,
        "reconciliation": {
            "period_capital_method": "NAV_ROLLFORWARD",
            "allocation_endpoint_check": {"matched": True, "delta_usd": "0"},
        },
    }


def _reconcile_fixture_replay_to_allocations(catalog: dict[str, Any]) -> None:
    """Keep synthetic replay totals aligned to its allocation endpoints."""

    totals = {"SYSTEMATIC": Decimal("0"), "RESERVE": Decimal("0")}
    for allocation in catalog["allocations"]:
        scope = allocation.get("scope")
        if scope in totals:
            totals[scope] += Decimal(str(allocation["current_value"]["value"]))
    catalog["replay"]["systematic_total_usd"] = decimal_text(totals["SYSTEMATIC"])
    catalog["replay"]["total_liquid_usd"] = decimal_text(
        totals["SYSTEMATIC"] + totals["RESERVE"]
    )


def _attach_hwm_envelope(
    catalog: dict[str, Any], crosswalk: list[dict[str, Any]]
) -> None:
    """Attach the envelope, its crosswalk, the role-correct admitted
    artifact entries, the hash-verified acceptance record, and allocation
    values matching the recomputed endpoint to a fixture catalog.  Tests that
    mutate the envelope after construction re-assign catalog["hwm_inputs"]
    afterwards."""
    envelope = _hwm_envelope(crosswalk)
    catalog["hwm_inputs"] = envelope
    catalog["strategy_returns"] = crosswalk

    # The acceptance record embedded by the builder for consumption-boundary
    # anchor-content cross-check.  Must equal the envelope's anchor fields.
    anchor = envelope["anchor"]
    catalog["hwm_acceptance_record"] = {
        "anchor_date": anchor["anchor_date"],
        "index_base_value": anchor["index_base_value"],
        "hwm_policy_id": anchor["hwm_policy_id"],
    }

    # Replace allocations with values that match the envelope's recomputed
    # endpoint so the validator's independent recomputation passes (delta=0).
    # Preserve the catalog's original RESERVE allocation (not in perimeter).
    from decimal import Decimal

    by_day: dict[str, dict[str, Decimal]] = {}
    for rw in crosswalk:
        sid = str(rw["strategy_id"])
        day = str(rw["date"])
        ret = Decimal(str(rw.get("return", "0")))
        by_day.setdefault(day, {})[sid] = ret

    endpoint_date = "2026-01-03"  # latest allocation as_of_date
    matching_allocations: list[dict[str, Any]] = []
    for row in envelope["periods"]:
        if (
            row.get("disposition") != "AVAILABLE"
            or row.get("lane") != "ACCOUNT"
            or str(row.get("date")) != endpoint_date
        ):
            continue
        sid = str(row["strategy_id"])
        capital = Decimal(str(row["beginning_capital_usd"]))
        period_return = by_day.get(endpoint_date, {}).get(sid, Decimal(0))
        endpoint_val = str(capital * (Decimal(1) + period_return))
        matching_allocations.append(
            {
                "allocation_id": f"allocation-{sid}",
                "as_of_date": endpoint_date,
                "current_value": {
                    "currency": "USD",
                    "status": "AVAILABLE",
                    "value": endpoint_val,
                },
                "scope": "SYSTEMATIC",
                "source_key": "synthetic-allocation",
                "strategy_id": sid,
                "systematic_weight": "0.1111111111111111111111111111111111",
            }
        )
    # Keep any non-SYSTEMATIC (RESERVE) allocations from the original catalog.
    original_allocations = catalog.get("allocations", [])
    if isinstance(original_allocations, list):
        for alloc in original_allocations:
            if isinstance(alloc, dict) and alloc.get("scope") != "SYSTEMATIC":
                matching_allocations.append(alloc)
    catalog["allocations"] = matching_allocations
    _reconcile_fixture_replay_to_allocations(catalog)

    catalog["source_artifacts"] = [
        *catalog["source_artifacts"],
        {
            "artifact_id": "anchor-acceptance-1",
            "as_of_date": "2026-01-03",
            "schema_name": "hwm-anchor-acceptance",
            "schema_version": "1",
            "sha256": "a" * 64,
            "source_key": "hwm-anchor-acceptance",
        },
        {
            "artifact_id": "calendar-account-1",
            "as_of_date": "2026-01-03",
            "schema_name": "hwm-calendar-authority-account",
            "schema_version": "1",
            "sha256": "b" * 64,
            "source_key": "hwm-calendar-authority-account",
        },
    ]


def _hwm_freshness_file(tmp_path: Path) -> Path:
    payload = {
        "assessment_schema_version": 2,
        "as_of_date": "2026-01-03",
        "assessments": [
            {
                "blocks_current_consumers": False,
                "evaluated_on": "2026-01-03",
                "fact_type": "strategy-returns",
                "reason": "synthetic fresh",
                "state": "FRESH",
                "subject_id": f"strategy-{index}",
            }
            for index in range(1, 10)
        ],
        "policy": {
            "expected_interval_days": 1,
            "grace_days": 0,
            "mode": "RECURRING",
            "policy_id": "daily-current",
            "stale_after_days": 5,
        },
    }
    path = tmp_path / "freshness-health.json"
    path.write_text(json.dumps(payload), encoding="utf-8")
    return path


def _summary_row(report: dict[str, Any], metric: str) -> Any:
    return next(
        row["value"] for row in report["portfolio_summary"] if row["metric"] == metric
    )


def test_hwm_numerical_composite_matches_hand_computation(tmp_path: Path) -> None:
    """Multi-sleeve varying-capital periods: the composite is the
    beginning-capital-weighted daily return, geometrically linked."""
    catalog = _baseline_catalog()
    crosswalk = _hwm_return_crosswalk()
    _attach_hwm_envelope(catalog, crosswalk)
    catalog["replay"]["portfolio_action"] = "ELIGIBLE_FOR_REVIEW"
    baseline_root = _published_baseline(tmp_path / "baseline", catalog)

    report, _, _ = reporter.build_report_data(
        baseline_root, freshness_health_path=_hwm_freshness_file(tmp_path)
    )

    hwm = report["hwm_summary"]
    assert hwm["hwm_numerical_status"] == "AVAILABLE"
    assert "hwm_numerical_reason" not in hwm
    assert {row["strategy_id"] for row in report["strategy_summary"]} == {
        f"strategy-{index}" for index in range(1, 8)
    }
    assert {row["strategy_id"] for row in report["strategy_summary"]}.isdisjoint(
        {"strategy-8", "strategy-9"}
    )
    # Hand-computed expectation (Decimal, 34-digit precision, half-even):
    #   day 2026-01-02: (100*0.10 + 300*(-0.05) + 5*100*0) / 900
    #   day 2026-01-03: (110*0.02 + 285*0.04 + 5*100*0) / 895
    #   final index: 1.009555555555555555555555555555556
    with localcontext() as ctx:
        ctx.prec = 34
        ctx.rounding = ROUND_HALF_EVEN
        r1 = (Decimal(100) * Decimal("0.10") + Decimal(300) * Decimal("-0.05")) / (
            Decimal(900)
        )
        r2 = (Decimal(110) * Decimal("0.02") + Decimal(285) * Decimal("0.04")) / (
            Decimal(895)
        )
        index = Decimal("1.0")
        peak = index
        max_drawdown = Decimal("0")
        for r in (r1, r2):
            index = index * (1 + r)
            peak = max(peak, index)
            max_drawdown = max(max_drawdown, 1 - index / peak)
        current_drawdown = 1 - index / peak
    assert hwm["hwm_index"] == "1.009555555555555555555555555555556"
    assert hwm["hwm_current_drawdown"] == format(current_drawdown, "f")
    assert hwm["hwm_maximum_drawdown"] == format(max_drawdown, "f")
    assert current_drawdown == Decimal("0")
    assert max_drawdown > 0
    # The numerical gate only tightens: AVAILABLE leaves readiness untouched.
    assert _summary_row(report, "Advisory readiness") == "ELIGIBLE_FOR_REVIEW"
    # The report carries the governed envelope for downstream reconciliation.
    assert report["hwm_inputs"]["contract_version"] == 1


def test_hwm_uncovered_weight_uses_final_current_allocation_weight(
    tmp_path: Path,
) -> None:
    """HWM coverage details must use the same final allocation weights shown publicly."""

    catalog = _baseline_catalog()
    _attach_hwm_envelope(catalog, _hwm_return_crosswalk())
    freshness_path = _hwm_freshness_file(tmp_path)
    freshness = json.loads(freshness_path.read_text(encoding="utf-8"))
    stale = next(
        row for row in freshness["assessments"] if row["subject_id"] == "strategy-7"
    )
    stale["state"] = "STALE"
    stale["reason"] = "synthetic stale HWM lane"
    freshness_path.write_text(json.dumps(freshness), encoding="utf-8")
    baseline_root = _published_baseline(tmp_path / "baseline", catalog)

    report, _, _ = reporter.build_report_data(
        baseline_root, freshness_health_path=freshness_path
    )

    detail = next(
        row
        for row in report["hwm_summary"]["uncovered_details"]
        if row["strategy_id"] == "strategy-7"
    )
    strategy = next(
        row for row in report["strategy_summary"] if row["strategy_id"] == "strategy-7"
    )
    allocation = next(
        row
        for row in report["allocation_history"]
        if row["strategy_id"] == "strategy-7"
    )
    assert detail["systematic_weight"] == strategy["systematic_weight"]
    assert detail["systematic_weight"] == allocation["systematic_weight"]


def test_portable_hwm_proof_coexists_with_governed_report_and_workbook_hwm(
    tmp_path: Path,
) -> None:
    catalog = _baseline_catalog()
    crosswalk = _hwm_return_crosswalk()
    _attach_hwm_envelope(catalog, crosswalk)
    catalog["replay"]["portfolio_action"] = "ELIGIBLE_FOR_REVIEW"
    baseline_root = _published_baseline(tmp_path / "baseline", catalog)
    receipt, input_generation_bytes, calendar_bytes = _portable_hwm_bundle(
        schema_version=3
    )
    governed_report, _governed_bytes, _governed_baseline = reporter.build_report_data(
        baseline_root,
        freshness_health_path=_hwm_freshness_file(tmp_path),
    )

    candidate = reporter.build_candidate(
        baseline_root,
        tmp_path / "report-output",
        freshness_health_path=_hwm_freshness_file(tmp_path),
        hwm_portable_receipt=receipt,
        hwm_portable_generation_bytes=input_generation_bytes,
        hwm_portable_calendar_bytes=calendar_bytes,
    )
    candidate_dir = Path(candidate["candidate_dir"])
    report = json.loads(
        (candidate_dir / reporter.REPORT_DATA_NAME).read_text(encoding="utf-8")
    )
    proof = report["hwm_portable_proof"]

    assert report["hwm_summary"] == governed_report["hwm_summary"]
    assert report["hwm_summary"]["hwm_index"] != proof["current_index"]
    assert _summary_row(report, "Advisory readiness") == _summary_row(
        governed_report, "Advisory readiness"
    )
    assert proof["financial_acceptance"] == "NOT_FINANCIAL_ACCEPTANCE"
    assert proof["assurance_profile"] == "DEVELOPMENT"
    assert proof["input_authority_class"] == "PORTABLE_SYNTHETIC"
    assert proof["receipt_sha256"] == receipt["receipt_sha256"]
    assert [row["index"] for row in proof["index_path"]] == [
        "100",
        "104",
        "93.6",
        "93.6",
    ]
    assert proof["current_index"] == "93.6"
    assert proof["current_hwm"] == "104"
    assert proof["current_drawdown"] == "-0.1"
    assert proof["maximum_drawdown"] == "-0.1"
    assert proof["age_status"] == "COMPLETE_DATA_AGE_WARNING"
    assert proof["warnings"] == ["HWM_COMPLETE_DATA_AGE"]
    assert (
        next(
            row
            for row in report["provenance"]
            if row["artifact"] == "portfolio-hwm-4-portable-proof"
        )["schema"]
        == "portfolio-hwm-calculation-receipt/3"
    )
    assert _summary_row(report, "PORT-HWM-4 proof status") == (
        "DEVELOPMENT_NOT_FINANCIAL_ACCEPTANCE"
    )
    assert _summary_row(report, "PORT-HWM-4 proof index") == "93.6"
    assert _summary_row(report, "PORT-HWM-4 proof HWM") == "104"
    assert _summary_row(report, "PORT-HWM-4 proof current drawdown") == "-0.1"
    assert _summary_row(report, "PORT-HWM-4 proof maximum drawdown") == "-0.1"
    assert reporter.verify_workbook(candidate_dir / reporter.WORKBOOK_NAME)[
        "sheet_count"
    ] == len(reporter.SHEET_NAMES)
    workbook = load_workbook(
        candidate_dir / reporter.WORKBOOK_NAME, read_only=True, data_only=True
    )
    summary_sheet = workbook["Portfolio Summary"]
    workbook_summary = {
        str(summary_sheet.cell(row=row_index, column=1).value): summary_sheet.cell(
            row=row_index, column=4
        ).value
        for row_index in range(7, 7 + len(report["portfolio_summary"]))
    }
    assert workbook_summary["PORT-HWM-4 proof status"] == (
        "DEVELOPMENT_NOT_FINANCIAL_ACCEPTANCE"
    )
    assert Decimal(str(workbook_summary["PORT-HWM-4 proof index"])) == Decimal("93.6")
    assert Decimal(str(workbook_summary["PORT-HWM-4 proof HWM"])) == Decimal("104")
    assert Decimal(
        str(workbook_summary["PORT-HWM-4 proof current drawdown"])
    ) == Decimal("-0.1")
    assert Decimal(
        str(workbook_summary["PORT-HWM-4 proof maximum drawdown"])
    ) == Decimal("-0.1")
    assert [
        Decimal(str(workbook_summary[f"PORT-HWM-4 proof path {point['as_of']}"]))
        for point in proof["index_path"]
    ] == [Decimal("100"), Decimal("104"), Decimal("93.6"), Decimal("93.6")]
    workbook_contract = json.loads(
        (candidate_dir / "workbook-contract.json").read_text(encoding="utf-8")
    )
    assert workbook_contract["sentinel_cells"][0] == {
        "cell": "B4",
        "expected": "ELIGIBLE_FOR_REVIEW",
        "sheet": "Portfolio Summary",
    }
    workbook.close()

    schema = json.loads(
        (REPOSITORY_ROOT / "schemas/portfolio/report-data.schema.json").read_text(
            encoding="utf-8"
        )
    )
    portable_schema = schema["$defs"]["portable_hwm_proof"]
    assert set(portable_schema["required"]) == set(proof)
    assert portable_schema["properties"]["financial_acceptance"] == {
        "const": "NOT_FINANCIAL_ACCEPTANCE"
    }
    assert portable_schema["properties"]["assurance_profile"] == {
        "const": "DEVELOPMENT"
    }


def test_portable_hwm_proof_does_not_upgrade_missing_governed_authority(
    tmp_path: Path,
) -> None:
    catalog = _baseline_catalog()
    catalog["replay"]["portfolio_action"] = "ELIGIBLE_FOR_REVIEW"
    baseline_root = _published_baseline(tmp_path / "baseline", catalog)
    governed_report, _governed_bytes, _governed_path = reporter.build_report_data(
        baseline_root,
        freshness_health_path=_hwm_freshness_file(tmp_path),
    )

    receipt, input_generation_bytes, calendar_bytes = _portable_hwm_bundle()
    report, _report_bytes, _baseline_path = reporter.build_report_data(
        baseline_root,
        freshness_health_path=_hwm_freshness_file(tmp_path),
        hwm_portable_receipt=receipt,
        hwm_portable_generation_bytes=input_generation_bytes,
        hwm_portable_calendar_bytes=calendar_bytes,
    )

    assert report["hwm_summary"] == governed_report["hwm_summary"]
    assert _summary_row(report, "Advisory readiness") == _summary_row(
        governed_report, "Advisory readiness"
    )
    assert report["hwm_portable_proof"]["financial_acceptance"] == (
        "NOT_FINANCIAL_ACCEPTANCE"
    )


def test_report_rejects_tampered_portable_hwm_receipt(tmp_path: Path) -> None:
    catalog = _baseline_catalog()
    baseline_root = _published_baseline(tmp_path / "baseline", catalog)
    receipt, input_generation_bytes, calendar_bytes = _portable_hwm_bundle()
    receipt["current_hwm"] = Decimal("999")

    with pytest.raises(
        reporter.ReportValidationError, match="receipt identity differs"
    ):
        reporter.build_report_data(
            baseline_root,
            freshness_health_path=_hwm_freshness_file(tmp_path),
            hwm_portable_receipt=receipt,
            hwm_portable_generation_bytes=input_generation_bytes,
            hwm_portable_calendar_bytes=calendar_bytes,
        )


def test_hwm_numerical_unavailable_without_envelope(tmp_path: Path) -> None:
    """No envelope → UNAVAILABLE/NO_HWM_INPUTS; every existing field keeps
    its current value."""
    catalog = _baseline_catalog()
    baseline_root = _published_baseline(tmp_path, catalog)

    report, _, _ = reporter.build_report_data(
        baseline_root, freshness_health_path=_hwm_freshness_file(tmp_path)
    )

    assert "hwm_inputs" not in report
    hwm = report["hwm_summary"]
    assert hwm["hwm_numerical_status"] == "UNAVAILABLE"
    assert hwm["hwm_numerical_reason"] == "NO_HWM_INPUTS"
    assert "hwm_index" not in hwm
    assert "hwm_current_drawdown" not in hwm
    assert "hwm_maximum_drawdown" not in hwm
    # Existing coverage/status fields are untouched by the numerical block.
    assert hwm["status"] == "NOT_AUTOMATABLE_STALE_OR_MISSING_GOVERNED_HWM_DAILY_NAV"
    assert hwm["policy_id"] == "PORT-HWM-3"
    replay = catalog["replay"]
    assert _summary_row(report, "Full-systematic HWM") == hwm["status"]
    assert _summary_row(report, "Advisory readiness") == replay["portfolio_action"]
    assert _summary_row(report, "Baseline replay") == replay["status"]
    assert _summary_row(report, "Total liquid") == replay["total_liquid_usd"]


def test_hwm_numerical_unavailable_period_blocks_composite(tmp_path: Path) -> None:
    """A single UNAVAILABLE member row on a contributing day fails the whole
    composite closed — the composite is never recalculated over survivors."""
    catalog = _baseline_catalog()
    crosswalk = _hwm_return_crosswalk()
    envelope = _hwm_envelope(crosswalk)
    blocked = next(
        row
        for row in envelope["periods"]
        if row["strategy_id"] == "strategy-4" and row["date"] == "2026-01-02"
    )
    blocked["disposition"] = "UNAVAILABLE"
    blocked["beginning_capital_usd"] = None
    blocked["return_source_record_id"] = None
    _attach_hwm_envelope(catalog, crosswalk)
    catalog["hwm_inputs"] = envelope
    baseline_root = _published_baseline(tmp_path, catalog)

    report, _, _ = reporter.build_report_data(
        baseline_root, freshness_health_path=_hwm_freshness_file(tmp_path)
    )

    hwm = report["hwm_summary"]
    assert hwm["hwm_numerical_status"] == "UNAVAILABLE"
    assert hwm["hwm_numerical_reason"] == "INCOMPLETE_PERIODS"
    assert "hwm_index" not in hwm


def test_hwm_numerical_unavailable_gates_readiness(tmp_path: Path) -> None:
    """Conformance vector missing_full_hwm_fails_closed: an in-scope HWM
    perimeter with numerical status UNAVAILABLE gates advisory readiness to
    DO_NOT_EXPAND_SYSTEMATIC_RISK even when the replayed control action is
    permissive."""
    catalog = _baseline_catalog()
    catalog["replay"]["portfolio_action"] = "ELIGIBLE_FOR_REVIEW"
    baseline_root = _published_baseline(tmp_path, catalog)

    report, _, _ = reporter.build_report_data(
        baseline_root, freshness_health_path=_hwm_freshness_file(tmp_path)
    )

    assert report["hwm_summary"]["hwm_numerical_status"] == "UNAVAILABLE"
    assert _summary_row(report, "Advisory readiness") == "DO_NOT_EXPAND_SYSTEMATIC_RISK"


def test_hwm_numerical_all_unavailable_day_fails_closed(tmp_path: Path) -> None:
    """Regression: a required day whose members are ALL UNAVAILABLE is a
    coverage gap — it must not be silently skipped.  Only a day where every
    perimeter member is EXPECTED_NON_TRADING (or policy-excluded) carries no
    period."""
    catalog = _baseline_catalog()
    crosswalk = _hwm_return_crosswalk()
    envelope = _hwm_envelope(crosswalk)
    for row in envelope["periods"]:
        if row["date"] == "2026-01-02":
            row["disposition"] = "UNAVAILABLE"
            row["beginning_capital_usd"] = None
            row["return_source_record_id"] = None
    _attach_hwm_envelope(catalog, crosswalk)
    catalog["hwm_inputs"] = envelope
    baseline_root = _published_baseline(tmp_path, catalog)

    report, _, _ = reporter.build_report_data(
        baseline_root, freshness_health_path=_hwm_freshness_file(tmp_path)
    )

    hwm = report["hwm_summary"]
    assert hwm["hwm_numerical_status"] == "UNAVAILABLE"
    assert hwm["hwm_numerical_reason"] == "INCOMPLETE_PERIODS"
    assert "hwm_index" not in hwm


def test_hwm_numerical_duplicate_return_reference_fails_closed(
    tmp_path: Path,
) -> None:
    """Crosswalk binding: a second strategy_returns row carrying the same
    reference (last-write-wins overwrite probe) must fail closed instead of
    silently rebinding the return."""
    catalog = _baseline_catalog()
    crosswalk = _hwm_return_crosswalk()
    _attach_hwm_envelope(catalog, crosswalk)
    proven = crosswalk[0]
    crosswalk.append(
        # Duplicate reference, divergent return, distinct date so no earlier
        # account-live series check fires — only the composite's strict
        # reference binding can catch the collision.
        {
            **proven,
            "date": "2026-01-01",
            "return": "9.99",
            "source_record_ids": ["nav-duplicate"],
        }
    )
    catalog["strategy_returns"] = crosswalk
    baseline_root = _published_baseline(tmp_path, catalog)

    report, _, _ = reporter.build_report_data(
        baseline_root, freshness_health_path=_hwm_freshness_file(tmp_path)
    )

    hwm = report["hwm_summary"]
    assert hwm["hwm_numerical_status"] == "UNAVAILABLE"
    assert hwm["hwm_numerical_reason"] == "DUPLICATE_RETURN_REFERENCE"
    assert "hwm_index" not in hwm


def test_hwm_numerical_mismatched_return_reference_fails_closed(
    tmp_path: Path,
) -> None:
    """A crosswalk row matching only the reference — wrong strategy, date,
    or evidence layer — does not bind; the composite fails closed."""
    catalog = _baseline_catalog()
    crosswalk = _hwm_return_crosswalk()
    _attach_hwm_envelope(catalog, crosswalk)
    # The row for strategy-1 on 2026-01-02 is re-labeled to PROVIDER_LIVE:
    # the reference no longer has an ACCOUNT_LIVE identity match.
    target = next(
        row
        for row in crosswalk
        if row["strategy_id"] == "strategy-1" and row["date"] == "2026-01-02"
    )
    target["evidence_layer"] = "PROVIDER_LIVE"
    baseline_root = _published_baseline(tmp_path, catalog)

    report, _, _ = reporter.build_report_data(
        baseline_root, freshness_health_path=_hwm_freshness_file(tmp_path)
    )

    hwm = report["hwm_summary"]
    assert hwm["hwm_numerical_status"] == "UNAVAILABLE"
    assert hwm["hwm_numerical_reason"] == "UNRESOLVED_RETURN_REFERENCE"
    assert "hwm_index" not in hwm


def test_hwm_numerical_publishes_without_freshness_argument(tmp_path: Path) -> None:
    """The numerical state is independent of the freshness argument: with no
    freshness health supplied the report still publishes hwm_summary with
    the envelope-derived numerical status and applies the readiness gate."""
    catalog = _baseline_catalog()
    crosswalk = _hwm_return_crosswalk()
    _attach_hwm_envelope(catalog, crosswalk)
    catalog["replay"]["portfolio_action"] = "ELIGIBLE_FOR_REVIEW"
    baseline_root = _published_baseline(tmp_path / "with-envelope", catalog)

    report, _, _ = reporter.build_report_data(baseline_root)

    hwm = report["hwm_summary"]
    # Numerical-only summary: no coverage fields without the freshness path.
    assert "covered_value_usd" not in hwm
    assert hwm["hwm_numerical_status"] == "AVAILABLE"
    assert "hwm_index" in hwm
    # AVAILABLE does not tighten readiness.
    assert _summary_row(report, "Advisory readiness") == "ELIGIBLE_FOR_REVIEW"

    catalog_no_envelope = _baseline_catalog()
    catalog_no_envelope["replay"]["portfolio_action"] = "ELIGIBLE_FOR_REVIEW"
    baseline_root_no_env = _published_baseline(
        tmp_path / "without-envelope", catalog_no_envelope
    )
    report_no_env, _, _ = reporter.build_report_data(baseline_root_no_env)

    hwm_no_env = report_no_env["hwm_summary"]
    assert hwm_no_env["hwm_numerical_status"] == "UNAVAILABLE"
    assert hwm_no_env["hwm_numerical_reason"] == "NO_HWM_INPUTS"
    # UNAVAILABLE gates readiness even without the freshness argument.
    assert (
        _summary_row(report_no_env, "Advisory readiness")
        == "DO_NOT_EXPAND_SYSTEMATIC_RISK"
    )
    # Replay-sourced fallback rows keep their pre-envelope values.
    replay = catalog_no_envelope["replay"]
    assert (
        _summary_row(report_no_env, "Full-systematic HWM")
        == replay["full_systematic_hwm_status"]
    )
    assert _summary_row(report_no_env, "NAV coverage") == replay["nav_coverage"]


def test_hwm_numerical_policy_exclusion_fraud_fails_closed(tmp_path: Path) -> None:
    """Probe: rows relabeled EXCLUDED_BY_POLICY without an admitted policy
    exclusion must not publish a number (wrong-data class: the envelope is
    hash-bound but semantically invalid)."""
    catalog = _baseline_catalog()
    crosswalk = _hwm_return_crosswalk()
    _attach_hwm_envelope(catalog, crosswalk)
    for row in catalog["hwm_inputs"]["periods"]:
        if row["strategy_id"] == "strategy-1":
            row["disposition"] = "EXCLUDED_BY_POLICY"
            row["beginning_capital_usd"] = None
            row["return_source_record_id"] = None
    baseline_root = _published_baseline(tmp_path, catalog)

    report, _, _ = reporter.build_report_data(
        baseline_root, freshness_health_path=_hwm_freshness_file(tmp_path)
    )

    hwm = report["hwm_summary"]
    assert hwm["hwm_numerical_status"] == "UNAVAILABLE"
    assert hwm["hwm_numerical_reason"] == "POLICY_EXCLUSION_INVALID"
    assert "hwm_index" not in hwm


def test_hwm_numerical_unreconciled_endpoint_blocks(tmp_path: Path) -> None:
    """Probe: a failed or non-zero allocation endpoint check blocks
    computation — the declared reconciliation is enforced at consumption."""
    for endpoint in (
        {"matched": False, "delta_usd": "999"},
        {"matched": True, "delta_usd": "999"},
        {"matched": False, "delta_usd": "0"},
    ):
        catalog = _baseline_catalog()
        crosswalk = _hwm_return_crosswalk()
        _attach_hwm_envelope(catalog, crosswalk)
        catalog["hwm_inputs"]["reconciliation"]["allocation_endpoint_check"] = endpoint
        baseline_root = _published_baseline(tmp_path / "baseline", catalog)

        report, _, _ = reporter.build_report_data(
            baseline_root, freshness_health_path=_hwm_freshness_file(tmp_path)
        )

        hwm = report["hwm_summary"]
        assert hwm["hwm_numerical_status"] == "UNAVAILABLE"
        assert hwm["hwm_numerical_reason"] == "ENDPOINT_NOT_RECONCILED"
        assert "hwm_index" not in hwm


def test_hwm_numerical_incomplete_envelope_fails_closed(tmp_path: Path) -> None:
    """Probe: a structurally incomplete envelope (only a periods key) is
    invalid at the consumption boundary."""
    catalog = _baseline_catalog()
    crosswalk = _hwm_return_crosswalk()
    _attach_hwm_envelope(catalog, crosswalk)
    catalog["hwm_inputs"] = {
        "periods": catalog["hwm_inputs"]["periods"],
    }
    baseline_root = _published_baseline(tmp_path, catalog)

    report, _, _ = reporter.build_report_data(
        baseline_root, freshness_health_path=_hwm_freshness_file(tmp_path)
    )

    hwm = report["hwm_summary"]
    assert hwm["hwm_numerical_status"] == "UNAVAILABLE"
    assert hwm["hwm_numerical_reason"] == "HWM_INPUTS_INVALID"
    assert "hwm_index" not in hwm


def test_hwm_numerical_available_with_missing_flow_fails_closed(
    tmp_path: Path,
) -> None:
    """Probe: an AVAILABLE row with MISSING flow/fx evidence is invalid at
    the consumption boundary, not silently computed."""
    catalog = _baseline_catalog()
    crosswalk = _hwm_return_crosswalk()
    _attach_hwm_envelope(catalog, crosswalk)
    for row in catalog["hwm_inputs"]["periods"]:
        if row["disposition"] == "AVAILABLE" and row["strategy_id"] == "strategy-1":
            row["flow_status"] = "MISSING"
    baseline_root = _published_baseline(tmp_path, catalog)

    report, _, _ = reporter.build_report_data(
        baseline_root, freshness_health_path=_hwm_freshness_file(tmp_path)
    )

    hwm = report["hwm_summary"]
    assert hwm["hwm_numerical_status"] == "UNAVAILABLE"
    assert hwm["hwm_numerical_reason"] == "HWM_INPUTS_INVALID"
    assert "hwm_index" not in hwm


# ---------------------------------------------------------------------------
# QC8 regression: full consumption-boundary context enforcement
# ---------------------------------------------------------------------------


def test_hwm_later_anchor_truncated_history_rejected(tmp_path: Path) -> None:
    """QC8: a fabricated later anchor that truncates the expected keyspace
    is rejected by the report consumer's anchor-content cross-check (the
    acceptance record from the catalog has anchor_date 2026-01-01, but the
    envelope claims 2026-01-02)."""
    catalog = _baseline_catalog()
    crosswalk = _hwm_return_crosswalk()
    _attach_hwm_envelope(catalog, crosswalk)
    # The builder-embedded acceptance record says anchor_date=2026-01-01
    # (matching the original envelope).  Fabricate the envelope anchor.
    catalog["hwm_inputs"]["anchor"]["anchor_date"] = "2026-01-02"
    # Remove the 2026-01-01 period rows (they'd be outside the new keyspace).
    catalog["hwm_inputs"]["periods"] = [
        p for p in catalog["hwm_inputs"]["periods"] if p["date"] != "2026-01-01"
    ]
    baseline_root = _published_baseline(tmp_path, catalog)

    report, _, _ = reporter.build_report_data(
        baseline_root, freshness_health_path=_hwm_freshness_file(tmp_path)
    )

    hwm = report["hwm_summary"]
    assert hwm["hwm_numerical_status"] == "UNAVAILABLE"
    # The anchor cross-check fires first (ANCHOR_CONTENT_MISMATCH).
    assert hwm["hwm_numerical_reason"] == "HWM_INPUTS_INVALID"
    assert "hwm_index" not in hwm


def test_hwm_inflated_allocation_unchanged_declaration_rejected(
    tmp_path: Path,
) -> None:
    """QC8: an allocation inflated 10x with the envelope still declaring
    {matched: true, delta_usd: "0"} is rejected by the report consumer's
    endpoint recomputation."""
    from decimal import Decimal as _D

    catalog = _baseline_catalog()
    crosswalk = _hwm_return_crosswalk()
    _attach_hwm_envelope(catalog, crosswalk)
    # Inflate every systematic allocation 10x.
    for alloc in catalog["allocations"]:
        if alloc.get("scope") == "SYSTEMATIC" and alloc.get("current_value"):
            alloc["current_value"]["value"] = str(
                _D(alloc["current_value"]["value"]) * 10
            )
    _reconcile_fixture_replay_to_allocations(catalog)
    baseline_root = _published_baseline(tmp_path, catalog)

    report, _, _ = reporter.build_report_data(
        baseline_root, freshness_health_path=_hwm_freshness_file(tmp_path)
    )

    hwm = report["hwm_summary"]
    assert hwm["hwm_numerical_status"] == "UNAVAILABLE"
    assert hwm["hwm_numerical_reason"] == "ENDPOINT_NOT_RECONCILED"
    assert "hwm_index" not in hwm


def test_hwm_missing_acceptance_record_fails_closed(tmp_path: Path) -> None:
    """QC8: if the catalog lacks hwm_acceptance_record, the report consumer
    fails closed with MISSING_CONSUMPTION_CONTEXT — never falls back."""
    catalog = _baseline_catalog()
    crosswalk = _hwm_return_crosswalk()
    _attach_hwm_envelope(catalog, crosswalk)
    # Remove the builder-embedded acceptance record.
    catalog.pop("hwm_acceptance_record", None)
    baseline_root = _published_baseline(tmp_path, catalog)

    report, _, _ = reporter.build_report_data(
        baseline_root, freshness_health_path=_hwm_freshness_file(tmp_path)
    )

    hwm = report["hwm_summary"]
    assert hwm["hwm_numerical_status"] == "UNAVAILABLE"
    assert hwm["hwm_numerical_reason"] == "MISSING_CONSUMPTION_CONTEXT"
    assert "hwm_index" not in hwm


def test_hwm_valid_envelope_publishes_with_full_context(tmp_path: Path) -> None:
    """QC8 parity: a valid envelope with all consumption context present
    publishes an AVAILABLE index (identical at all three boundaries)."""
    catalog = _baseline_catalog()
    crosswalk = _hwm_return_crosswalk()
    _attach_hwm_envelope(catalog, crosswalk)
    # The fixture already embeds hwm_acceptance_record and matching
    # allocations.
    baseline_root = _published_baseline(tmp_path, catalog)

    report, _, _ = reporter.build_report_data(
        baseline_root, freshness_health_path=_hwm_freshness_file(tmp_path)
    )

    hwm = report["hwm_summary"]
    assert hwm["hwm_numerical_status"] == "AVAILABLE"
    assert "hwm_numerical_reason" not in hwm
    assert "hwm_index" in hwm


def test_invested_capital_is_visible_and_exact_in_native_workbook(
    tmp_path: Path,
) -> None:
    baseline_root = _published_baseline(tmp_path / "baseline", _baseline_catalog())
    report, _report_bytes, _baseline_path = reporter.build_report_data(baseline_root)
    day = report["as_of_date"]
    report["invested_capital"] = {
        "daily_by_strategy": [
            {
                "date": day,
                "strategy_id": "strategy-1",
                "invested_capital_usd": "125",
            }
        ],
        "daily_aggregate": [{"date": day, "invested_capital_usd": "125"}],
        "current_by_strategy": [
            {"strategy_id": "strategy-1", "invested_capital_usd": "125"}
        ],
        "current_aggregate_usd": "125",
    }
    workbook_path = tmp_path / "capital.xlsx"
    reporter._workbook(report, workbook_path)

    evidence = reporter.verify_workbook(workbook_path)
    parity = reporter._workbook_parity_receipt(
        canonical_json_bytes(report), workbook_path.read_bytes()
    )
    workbook = load_workbook(workbook_path, read_only=False, data_only=False)
    try:
        assert workbook.sheetnames[-2:] == list(reporter.CAPITAL_SHEET_NAMES)
        assert (
            workbook[reporter.CAPITAL_SHEET_NAMES[1]]["J4"].value
            == reporter.CAPITAL_RECONCILIATION_FORMULA
        )
        assert workbook[reporter.CAPITAL_SHEET_NAMES[1]]["A6"].value == 125
    finally:
        workbook.close()
    assert evidence["sheet_count"] == len(reporter.SHEET_NAMES) + 2
    assert parity["table_count"] == len(reporter.TABLE_NAMES) + 2

    wrong_cell = tmp_path / "wrong-capital-cell.xlsx"
    shutil.copyfile(workbook_path, wrong_cell)
    _rewrite_xlsx_cell(
        wrong_cell,
        sheet_name=reporter.CAPITAL_SHEET_NAMES[1],
        cell_reference="A6",
        value="126",
    )
    with pytest.raises(
        reporter.ReportValidationError,
        match="tblInvestedCapitalCurrent.*invested_capital_usd",
    ):
        reporter._workbook_parity_receipt(
            canonical_json_bytes(report), wrong_cell.read_bytes()
        )

    wrong_formula = tmp_path / "wrong-capital-formula.xlsx"
    shutil.copyfile(workbook_path, wrong_formula)
    _rewrite_xlsx_cell(
        wrong_formula,
        sheet_name=reporter.CAPITAL_SHEET_NAMES[1],
        cell_reference="J4",
        formula="1+1",
    )
    with pytest.raises(
        reporter.ReportValidationError, match="formula inventory differs"
    ):
        reporter._workbook_parity_receipt(
            canonical_json_bytes(report), wrong_formula.read_bytes()
        )


def test_source_qualified_decomposition_is_visible_and_exact_in_native_workbook(
    tmp_path: Path,
) -> None:
    baseline_root = _published_baseline(tmp_path / "baseline", _baseline_catalog())
    report, _report_bytes, _baseline_path = reporter.build_report_data(baseline_root)
    source_rows = []
    series_rows = []
    for index, source_key in enumerate(("us-equity-style", "multi-asset-style"), 1):
        source_rows.append(
            {
                "analytics_hash": f"{index}" * 64,
                "as_of_date": report["as_of_date"],
                "generation_id": f"{index + 2}" * 64,
                "plan_spec": {
                    "sha256": f"{index + 4}" * 64,
                    "spec_id": f"{source_key}:STATIC",
                },
                "source_key": source_key,
                "source_label": source_key,
                "source_sha256": f"{index + 6}" * 64,
                "surface": "STATIC",
            }
        )
        series_rows.append(
            {
                "alpha": {
                    "periodic_intercept": "0.003",
                    "net_annualized": "0.756",
                    "net_status": "AVAILABLE_OBSERVED_BASIS",
                },
                "contribution_summary": {
                    "factor_means": {"VALUE": f"0.0{index}"},
                    "alpha_mean": "0.003",
                    "residual_mean": "-0.001",
                    "fitted_mean": "0.013" if index == 1 else "0.023",
                    "realized_mean": "0.012" if index == 1 else "0.022",
                },
                "coverage": {
                    "aligned_observations": 40,
                    "observed_start": "2026-01-02",
                    "observed_end": "2026-03-02",
                },
                "diagnostics": {
                    "train": {"r_squared": "0.8"},
                    "holdout": {
                        "predicted_r_squared": "-0.2" if index == 1 else "0.7",
                        "rmse": "0.001",
                    },
                },
                "exposures": [{"beta": f"0.{index}", "factor_id": "VALUE"}],
                "model_id": f"model-{index}",
                "series_id": "Shared strategy",
                "source_key": source_key,
                "status": "DIAGNOSTIC_ONLY" if index == 1 else "AVAILABLE",
            }
        )
    series_rows[0]["coverage"]["factor_missing_weekdays"] = 1
    surface = {
        "available_count": 2,
        "available_source_keys": [row["source_key"] for row in source_rows],
        "default_source_key": "us-equity-style",
        "diagnostic_count": 0,
        "series": series_rows,
        "sources": source_rows,
        "unavailable_count": 0,
    }
    lens = {
        "default_surface": "STATIC",
        "lens_id": "ETF_STYLE",
        "surfaces": {"STATIC": surface},
    }
    report["decomposition_lenses"] = {
        "default_evidence_mode": "LIVE_ONLY",
        "default_lens": "ETF_STYLE",
        "modes": {"LIVE_ONLY": [lens]},
    }
    workbook_path = tmp_path / "decomposition.xlsx"
    reporter._workbook(report, workbook_path)

    evidence = reporter.verify_workbook(workbook_path)
    parity = reporter._workbook_parity_receipt(
        canonical_json_bytes(report), workbook_path.read_bytes()
    )
    cells, factors = reporter._decomposition_workbook_rows(
        report["decomposition_lenses"]
    )
    assert {(row["source_key"], row["series_id"]) for row in cells} == {
        ("multi-asset-style", "Shared strategy"),
        ("us-equity-style", "Shared strategy"),
    }
    assert {row["beta"] for row in factors} == {"0.1", "0.2"}
    assert evidence["sheet_count"] == len(reporter.SHEET_NAMES) + 2
    assert parity["table_count"] == len(reporter.TABLE_NAMES) + 2

    # Read actual Excel cells independently of the export/parity projection.
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet_rows = list(
            workbook[reporter.DECOMPOSITION_SHEET_NAMES[0]].iter_rows(
                min_row=5, values_only=True
            )
        )
        exported = {
            row["Source Key"]: row
            for values in sheet_rows[1:]
            if (row := dict(zip(sheet_rows[0], values)))["Source Key"]
        }
        diagnostic = exported["us-equity-style"]
        assert diagnostic["Status"] == "DIAGNOSTIC_ONLY"
        assert diagnostic["Alpha Periodic Intercept"] == pytest.approx(0.003)
        assert diagnostic["Contribution Residual Mean"] == pytest.approx(-0.001)
        assert diagnostic["Contribution Fitted Mean"] == pytest.approx(0.013)
        assert diagnostic["Contribution Realized Mean"] == pytest.approx(0.012)
        assert diagnostic["Train R Squared"] == pytest.approx(0.8)
        assert diagnostic["Holdout Predicted R Squared"] == pytest.approx(-0.2)
        assert diagnostic["Coverage Aligned Observations"] == 40
        assert diagnostic["Coverage Observed Start"] == "2026-01-02"
        assert diagnostic["Coverage Factor Missing Weekdays"] == 1
        assert diagnostic["Source Sha256"] == "7" * 64
    finally:
        workbook.close()

    wrong_cell = tmp_path / "wrong-decomposition-cell.xlsx"
    shutil.copyfile(workbook_path, wrong_cell)
    _rewrite_xlsx_cell(
        wrong_cell,
        sheet_name=reporter.DECOMPOSITION_SHEET_NAMES[1],
        cell_reference="A6",
        value="0.9",
    )
    with pytest.raises(
        reporter.ReportValidationError,
        match="tblDecompositionFactors.*beta",
    ):
        reporter._workbook_parity_receipt(
            canonical_json_bytes(report), wrong_cell.read_bytes()
        )


@pytest.mark.parametrize("command", ("build", "verify", "publish"))
def test_report_cli_forwards_governed_hwm_bundle_for_build_verify_and_publish(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    command: str,
) -> None:
    result_path = tmp_path / "hwm-result.json"
    input_path = tmp_path / "hwm-input.json"
    calendar_path = tmp_path / "hwm-calendar.json"
    receipt_path = tmp_path / "hwm-receipt.json"
    capital_result_path = tmp_path / "capital-result.json"
    capital_input_path = tmp_path / "capital-input.json"
    decomposition_root = tmp_path / "decomposition-lenses"
    decomposition_plan = tmp_path / "decomposition-plan"
    result_path.write_bytes(b'{"status":"AVAILABLE"}\n')
    input_path.write_bytes(b'{"input":"bound"}\n')
    calendar_path.write_bytes(b'{"calendar":"bound"}\n')
    receipt_path.write_bytes(b'{"receipt":"bound"}\n')
    capital_result_path.write_bytes(b'{"outcome_id":"INVESTED_CAPITAL"}\n')
    capital_input_path.write_bytes(b'{"capital":"bound"}\n')
    decomposition_root.mkdir()
    decomposition_plan.mkdir()
    observed: dict[str, object] = {}

    def capture(*_args: object, **kwargs: object) -> dict[str, object]:
        observed.update(kwargs)
        return {"status": "PASS"}

    monkeypatch.setattr(
        reporter,
        f"{'build_candidate' if command == 'build' else 'verify_candidate' if command == 'verify' else 'publish_report'}",
        capture,
    )
    argv = [command, "--baseline-root", str(tmp_path / "baseline")]
    if command == "build":
        argv += ["--output-root", str(tmp_path / "candidate-root")]
    elif command == "verify":
        argv += [
            "--candidate-dir",
            str(tmp_path / "candidate"),
            "--excel-workbook",
            str(tmp_path / "workbook.xlsx"),
        ]
    else:
        argv += [
            "--candidate-dir",
            str(tmp_path / "candidate"),
            "--excel-workbook",
            str(tmp_path / "workbook.xlsx"),
            "--xlsx-result",
            str(tmp_path / "xlsx-result.json"),
            "--validation-result",
            str(tmp_path / "validation.json"),
            "--workbook-evidence-binding",
            str(tmp_path / "binding.json"),
            "--output-root",
            str(tmp_path / "report-stage"),
        ]
    argv += [
        "--hwm-governed-result",
        str(result_path),
        "--hwm-governed-input-generation",
        str(input_path),
        "--hwm-governed-calendar-authority",
        str(calendar_path),
        "--hwm-governed-admission-receipt",
        str(receipt_path),
        "--invested-capital-result",
        str(capital_result_path),
        "--invested-capital-input-generation",
        str(capital_input_path),
        "--decomposition-lenses-root",
        str(decomposition_root),
        "--decomposition-plan-bundle",
        str(decomposition_plan),
    ]

    assert reporter.main(argv) == 0
    assert {
        key: observed[key]
        for key in (
            "hwm_governed_calendar_bytes",
            "hwm_governed_generation_bytes",
            "hwm_governed_receipt_bytes",
            "hwm_governed_result",
            "invested_capital_input_generation_path",
            "invested_capital_result",
            "invested_capital_result_bytes",
            "decomposition_lenses_root",
            "decomposition_plan_bundle",
        )
    } == {
        "hwm_governed_calendar_bytes": calendar_path.read_bytes(),
        "hwm_governed_generation_bytes": input_path.read_bytes(),
        "hwm_governed_receipt_bytes": receipt_path.read_bytes(),
        "hwm_governed_result": {"status": "AVAILABLE"},
        "invested_capital_input_generation_path": capital_input_path,
        "invested_capital_result": {"outcome_id": "INVESTED_CAPITAL"},
        "invested_capital_result_bytes": capital_result_path.read_bytes(),
        "decomposition_lenses_root": decomposition_root,
        "decomposition_plan_bundle": decomposition_plan,
    }


def test_governed_hwm_direct_publish_requires_private_stage_context_before_io(
    tmp_path: Path,
) -> None:
    output_root = tmp_path / "report-output"

    with pytest.raises(
        reporter.ReportValidationError,
        match="requires private stage context",
    ):
        reporter.publish_report(
            tmp_path / "baseline",
            tmp_path / "candidate",
            tmp_path / "workbook.xlsx",
            tmp_path / "xlsx-result.json",
            tmp_path / "validation.json",
            output_root,
            hwm_governed_result={"status": "AVAILABLE"},
            hwm_governed_generation_bytes=b"{}",
            hwm_governed_calendar_bytes=b"{}",
            hwm_governed_receipt_bytes=b"{}",
        )

    assert not output_root.exists()


def test_financial_product_direct_publish_requires_private_stage_without_new_flags(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    strategy_ids = reporter.CURRENT_STRATEGY_CONTROL_IDS
    catalog = {"strategies": [{"strategy_id": value} for value in sorted(strategy_ids)]}
    controls = {value: {} for value in strategy_ids}
    monkeypatch.setattr(
        reporter, "_baseline", lambda *_args: (catalog, None, tmp_path / "catalog.json")
    )
    monkeypatch.setattr(
        reporter, "_report_product_controls", lambda _catalog: (controls, None)
    )
    output_root = tmp_path / "report-output"
    output_root.mkdir()
    pointer = output_root / "CURRENT.json"
    pointer.write_bytes(b"prior accepted pointer")
    qa = tmp_path / "qa"

    with pytest.raises(
        reporter.ReportValidationError, match="requires private stage context"
    ):
        reporter.publish_report(
            tmp_path / "baseline",
            tmp_path / "candidate",
            tmp_path / "workbook.xlsx",
            qa / "xlsx-result.json",
            qa / "validation.json",
            output_root,
        )

    assert pointer.read_bytes() == b"prior accepted pointer"
    assert list(output_root.iterdir()) == [pointer]
    assert not qa.exists()
