from __future__ import annotations

import csv
import json
import re
import sqlite3
import zipfile
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence

from openpyxl import load_workbook  # type: ignore[import-untyped]

from .catalog import ENDPOINTS, PRODUCTION_HOST, get_endpoint
from .exporters import (
    EXCEL_MAX_DATA_ROWS,
    _columns,
    _expanded_columns,
    _row_sort_key,
    _safe_filename,
    _unique_sheet_name,
    dataset_digest,
    partition_datasets,
    protect_csv_formula,
)
from .storage import COVERAGE_STATUSES
from .util import atomic_write, canonical_json, sha256_bytes, sha256_text


STATE_DATABASE = "matrixport.sqlite"
EXCEL_ERRORS = {"#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A"}
REQUIRED_CORE_PATHS = {
    "/strategy/api/v2/user/holding",
    "/strategy/api/v2/user/holding/volume",
    "/strategy/api/v2/user/trade_record",
    "/strategy/api/v2/product",
    "/strategy/api/v2/risk/nav_chart_data_v2",
    "/mapi/v1/wallet/balance",
    "/mapi/v1/wallet/currencies",
    "/mapi/v1/wallet/bills",
    "/mapi/v1/wallet/deposits",
    "/mapi/v1/wallet/withdrawals",
}
SOURCE_LIMITED_WALLET_HISTORY_PATHS = {
    "/mapi/v1/wallet/deposits",
    "/mapi/v1/wallet/withdrawals",
}

# Terminal atoms the wallet history walk may emit that verification accepts as
# source-limited coverage.  "enumerated_currency_failures" is the partial
# shape (accepted only when every failed currency carries strict rejection
# evidence; anything else fails closed) and
# "structured_parameter_rejections_dispositioned" is the complete shape where
# every provider-structured rejection was dispositioned.  These must stay a
# subset of the pipeline's WALLET_HISTORY_TERMINAL_REASONS.
TERMINAL_ENUMERATED_CURRENCY_FAILURES = "enumerated_currency_failures"
TERMINAL_STRUCTURED_REJECTIONS_DISPOSITIONED = (
    "structured_parameter_rejections_dispositioned"
)
WALLET_SOURCE_LIMITED_TERMINAL_REASONS = frozenset(
    {
        TERMINAL_ENUMERATED_CURRENCY_FAILURES,
        TERMINAL_STRUCTURED_REJECTIONS_DISPOSITIONED,
    }
)

_WALLET_PATH_ATOM_PREFIX: dict[str, str] = {
    "/mapi/v1/wallet/deposits": "wallet_deposit_currency",
    "/mapi/v1/wallet/withdrawals": "wallet_withdrawal_currency",
}


def _iter_file_bytes(path: Path) -> Iterable[tuple[str, bytes]]:
    if path.suffix.lower() == ".xlsx":
        with zipfile.ZipFile(path) as archive:
            for name in archive.namelist():
                yield f"{path}:{name}", archive.read(name)
    else:
        yield str(path), path.read_bytes()


def secret_scan(run_dir: Path, needles: Iterable[bytes]) -> list[str]:
    matches: list[str] = []
    normalized = [item for item in needles if item]
    if not normalized:
        return matches
    for path in run_dir.rglob("*"):
        if not path.is_file() or path.name.endswith(".tmp"):
            continue
        try:
            parts = _iter_file_bytes(path)
            for label, data in parts:
                if any(needle in data for needle in normalized):
                    matches.append(label)
        except (OSError, zipfile.BadZipFile):
            matches.append(f"{path}:unreadable")
    return sorted(set(matches))


def _state_path(run_dir: Path) -> Path:
    return run_dir / "state" / STATE_DATABASE


def _read_state(
    run_dir: Path,
) -> tuple[dict[str, Any] | None, dict[str, list[dict[str, Any]]], list[str]]:
    database = _state_path(run_dir)
    failures: list[str] = []
    datasets: dict[str, list[dict[str, Any]]] = defaultdict(list)
    if not database.is_file():
        return None, {}, [f"state/{STATE_DATABASE} missing"]
    extra = sorted(
        path.name for path in (run_dir / "state").glob("*.sqlite") if path != database
    )
    if extra:
        failures.append(f"unexpected SQLite files: {', '.join(extra)}")
    try:
        with sqlite3.connect(database) as connection:
            connection.row_factory = sqlite3.Row
            integrity = [
                str(row[0]) for row in connection.execute("PRAGMA integrity_check")
            ]
            foreign_keys = list(connection.execute("PRAGMA foreign_key_check"))
            if integrity != ["ok"]:
                failures.append(f"integrity_check={integrity!r}")
            if foreign_keys:
                failures.append(f"foreign_key_check={len(foreign_keys)} violation(s)")
            runs = [dict(row) for row in connection.execute("SELECT * FROM runs")]
            if len(runs) != 1:
                failures.append(f"expected one run row, found {len(runs)}")
                run = runs[0] if runs else None
            else:
                run = runs[0]
            if run is not None:
                if str(run.get("run_id")) != run_dir.name:
                    failures.append("run_id does not match output directory")
                if run.get("status") != "COMPLETED":
                    failures.append(
                        f"run status is {run.get('status')!r}, not COMPLETED"
                    )
                if run.get("completed_at_ms") is None:
                    failures.append("completed run has no completed_at_ms")
            for record in connection.execute(
                """
                SELECT dataset, stable_key, row_json, row_sha256
                FROM canonical_rows ORDER BY dataset, stable_key
                """
            ):
                try:
                    row = json.loads(record["row_json"])
                except json.JSONDecodeError:
                    failures.append(
                        f"invalid canonical JSON: {record['dataset']}/{record['stable_key']}"
                    )
                    continue
                if not isinstance(row, dict):
                    failures.append(
                        f"canonical row is not an object: {record['dataset']}/{record['stable_key']}"
                    )
                    continue
                if canonical_json(row) != record["row_json"]:
                    failures.append(
                        f"non-canonical row JSON: {record['dataset']}/{record['stable_key']}"
                    )
                if sha256_text(record["row_json"]) != record["row_sha256"]:
                    failures.append(
                        f"canonical row digest mismatch: {record['dataset']}/{record['stable_key']}"
                    )
                datasets[str(record["dataset"])].append(row)
    except (OSError, sqlite3.Error) as error:
        return None, {}, [f"state database unreadable:{type(error).__name__}"]
    return run, dict(datasets), failures


def verify_sqlite(run_dir: Path) -> dict[str, Any]:
    run, datasets, failures = _read_state(run_dir)
    return {
        "ok": not failures and run is not None,
        "path": f"state/{STATE_DATABASE}",
        "status": None if run is None else run.get("status"),
        "dataset_counts": {name: len(rows) for name, rows in sorted(datasets.items())},
        "failures": failures,
    }


def _read_manifest(run_dir: Path) -> tuple[list[dict[str, Any]], list[str]]:
    path = run_dir / "reports" / "request_manifest.jsonl"
    if not path.is_file():
        return [], ["request manifest missing"]
    rows: list[dict[str, Any]] = []
    failures: list[str] = []
    for line_number, line in enumerate(
        path.read_text(encoding="utf-8").splitlines(), start=1
    ):
        if not line.strip():
            continue
        try:
            row = json.loads(line)
        except json.JSONDecodeError:
            failures.append(f"line {line_number}:invalid JSON")
            continue
        if not isinstance(row, dict):
            failures.append(f"line {line_number}:not an object")
            continue
        row["_line_number"] = line_number
        rows.append(row)
    if not rows:
        failures.append("request manifest is empty")
    return rows, failures


def _safe_params(row: Mapping[str, Any]) -> dict[str, Any]:
    value = row.get("safe_params", {})
    if isinstance(value, str):
        value = json.loads(value)
    if not isinstance(value, dict):
        raise ValueError("safe_params is not an object")
    return dict(value)


def verify_requests(run_dir: Path) -> dict[str, Any]:
    rows, failures = _read_manifest(run_dir)
    checked = 0
    seen_hashes: set[str] = set()
    status_counts: dict[str, int] = defaultdict(int)
    for row in rows:
        line_number = int(row.get("_line_number", 0))
        method = str(row.get("method", "")).upper()
        path = str(row.get("path") or row.get("endpoint") or "")
        checked += 1
        try:
            params = _safe_params(row)
        except (ValueError, json.JSONDecodeError):
            failures.append(f"line {line_number}:invalid safe_params")
            continue
        try:
            endpoint = get_endpoint(path, method=method, host=PRODUCTION_HOST)
        except KeyError:
            failures.append(
                f"line {line_number}:endpoint is not in the production catalog: {method} {path}"
            )
            continue
        if method != "GET" or endpoint.method != "GET":
            failures.append(f"line {line_number}:non-GET request: {method} {path}")
        if not endpoint.live_enabled or endpoint.safety_class != "read_only":
            failures.append(
                f"line {line_number}:endpoint is not live read-only: {method} {path}"
            )
        names = set(params)
        unknown = names - endpoint.allowed_params
        missing = endpoint.required_params - names
        empty_required = {
            name for name in endpoint.required_params if params.get(name) in (None, "")
        }
        if unknown:
            failures.append(
                f"line {line_number}:unknown params: {', '.join(sorted(unknown))}"
            )
        if missing or empty_required:
            invalid = sorted(missing | empty_required)
            failures.append(
                f"line {line_number}:missing/empty required params: {', '.join(invalid)}"
            )
        if row.get("endpoint") not in (None, path):
            failures.append(f"line {line_number}:path and endpoint disagree")
        try:
            status_code = int(str(row.get("status_code")))
        except (TypeError, ValueError):
            status_code = 0
        if not 100 <= status_code <= 599:
            failures.append(f"line {line_number}:invalid HTTP status {status_code}")
        else:
            status_counts[str(status_code)] += 1
        request_hash = str(row.get("request_hash", ""))
        if not re.fullmatch(r"[0-9a-f]{64}", request_hash):
            failures.append(f"line {line_number}:invalid request_hash")
        elif request_hash in seen_hashes:
            failures.append(f"line {line_number}:duplicate request_hash")
        else:
            seen_hashes.add(request_hash)
    return {
        "ok": not failures and checked > 0,
        "checked": checked,
        "status_counts": dict(sorted(status_counts.items())),
        "failures": failures,
    }


def verify_raw(run_dir: Path) -> dict[str, Any]:
    raw_dir = run_dir / "raw"
    envelope_dir = raw_dir / "envelopes"
    response_dir = raw_dir / "responses"
    database = _state_path(run_dir)
    manifest, manifest_failures = _read_manifest(run_dir)
    failures = list(manifest_failures)
    if not envelope_dir.is_dir():
        failures.append("raw/envelopes missing")
    if not response_dir.is_dir():
        failures.append("raw/responses missing")
    if not database.is_file():
        failures.append(f"state/{STATE_DATABASE} missing")
        return {"ok": False, "checked": 0, "failures": failures}

    try:
        with sqlite3.connect(database) as connection:
            connection.row_factory = sqlite3.Row
            db_rows = {
                str(row["request_hash"]): dict(row)
                for row in connection.execute(
                    "SELECT * FROM requests ORDER BY request_hash"
                )
            }
    except sqlite3.Error as error:
        failures.append(f"requests table unreadable:{type(error).__name__}")
        return {"ok": False, "checked": 0, "failures": failures}

    manifest_rows = {str(row.get("request_hash", "")): row for row in manifest}
    envelopes: dict[str, dict[str, Any]] = {}
    for path in sorted(envelope_dir.glob("*.json")) if envelope_dir.is_dir() else []:
        try:
            envelope = json.loads(path.read_text(encoding="utf-8"))
            if not isinstance(envelope, dict):
                raise ValueError("not an object")
        except (OSError, ValueError, json.JSONDecodeError) as error:
            failures.append(f"{path.name}:invalid envelope:{type(error).__name__}")
            continue
        request_hash = str(envelope.get("request_hash", ""))
        if path.stem != request_hash or not re.fullmatch(r"[0-9a-f]{64}", request_hash):
            failures.append(f"{path.name}:filename/request_hash mismatch")
            continue
        if request_hash in envelopes:
            failures.append(f"{path.name}:duplicate request_hash")
            continue
        envelopes[request_hash] = envelope
        expected_raw = f"responses/{request_hash}.bin"
        expected_envelope = f"envelopes/{request_hash}.json"
        if envelope.get("raw_relpath") != expected_raw:
            failures.append(f"{path.name}:invalid raw_relpath")
        if envelope.get("envelope_relpath") != expected_envelope:
            failures.append(f"{path.name}:invalid envelope_relpath")
        response_path = response_dir / f"{request_hash}.bin"
        if not response_path.is_file():
            failures.append(f"{path.name}:response missing")
        else:
            digest = sha256_bytes(response_path.read_bytes())
            if digest != envelope.get("response_sha256"):
                failures.append(f"{path.name}:response digest mismatch")
        params = envelope.get("safe_params")
        if not isinstance(params, dict):
            failures.append(f"{path.name}:invalid safe_params")

    envelope_hashes = set(envelopes)
    response_hashes = (
        {path.stem for path in response_dir.glob("*.bin")}
        if response_dir.is_dir()
        else set()
    )
    db_hashes = set(db_rows)
    manifest_hashes = set(manifest_rows)
    if not (envelope_hashes == response_hashes == db_hashes == manifest_hashes):
        failures.append(
            "request evidence sets differ: "
            f"envelopes={len(envelope_hashes)},responses={len(response_hashes)},"
            f"sqlite={len(db_hashes)},manifest={len(manifest_hashes)}"
        )

    compare_fields = (
        "request_id",
        "request_hash",
        "method",
        "endpoint",
        "status_code",
        "elapsed_ms",
        "row_count",
        "response_sha256",
        "raw_relpath",
        "envelope_relpath",
        "requested_at_ms",
        "terminal_reason",
    )
    for request_hash in sorted(envelope_hashes & db_hashes & manifest_hashes):
        envelope = envelopes[request_hash]
        db_row = db_rows[request_hash]
        manifest_row = manifest_rows[request_hash]
        for field in compare_fields:
            if envelope.get(field) != db_row.get(field):
                failures.append(
                    f"{request_hash}:{field} differs between envelope and SQLite"
                )
            if db_row.get(field) != manifest_row.get(field):
                failures.append(
                    f"{request_hash}:{field} differs between SQLite and manifest"
                )
        try:
            envelope_params = canonical_json(envelope.get("safe_params"))
            manifest_params = canonical_json(_safe_params(manifest_row))
        except (ValueError, json.JSONDecodeError):
            failures.append(f"{request_hash}:invalid manifest safe_params")
            continue
        if envelope_params != db_row.get(
            "params_json"
        ) or manifest_params != db_row.get("params_json"):
            failures.append(f"{request_hash}:safe_params differ across evidence")
        envelope_bounds = (
            canonical_json(envelope.get("observed_bounds"))
            if envelope.get("observed_bounds") is not None
            else None
        )
        if envelope_bounds != db_row.get("observed_bounds_json"):
            failures.append(
                f"{request_hash}:observed bounds differ between envelope and SQLite"
            )
        if db_row.get("observed_bounds_json") != manifest_row.get(
            "observed_bounds_json"
        ):
            failures.append(
                f"{request_hash}:observed bounds differ between SQLite and manifest"
            )

    return {
        "ok": not failures and bool(envelopes),
        "checked": len(envelopes),
        "failures": failures,
    }


def verify_coverage(run_dir: Path) -> dict[str, Any]:
    database = _state_path(run_dir)
    failures: list[str] = []
    if not database.is_file():
        return {
            "ok": False,
            "checked": 0,
            "failures": [f"state/{STATE_DATABASE} missing"],
        }
    try:
        with sqlite3.connect(database) as connection:
            connection.row_factory = sqlite3.Row
            rows = [
                dict(row)
                for row in connection.execute(
                    "SELECT * FROM coverage ORDER BY endpoint"
                )
            ]
            request_rows = [
                dict(row)
                for row in connection.execute(
                    "SELECT endpoint, params_json, status_code, raw_relpath FROM requests ORDER BY requested_at_ms"
                )
            ]
    except sqlite3.Error as error:
        return {
            "ok": False,
            "checked": 0,
            "failures": [f"coverage unreadable:{type(error).__name__}"],
        }
    by_key = {str(row["endpoint"]): row for row in rows}
    expected = {
        f"{endpoint.method} https://{endpoint.host}{endpoint.path}": endpoint
        for endpoint in ENDPOINTS
    }
    missing = sorted(set(expected) - set(by_key))
    extra = sorted(set(by_key) - set(expected))
    if missing:
        failures.append(f"missing catalog coverage rows: {len(missing)}")
    if extra:
        failures.append(f"unexpected coverage rows: {len(extra)}")
    source_limited: list[str] = []
    for key in sorted(set(expected) & set(by_key)):
        endpoint = expected[key]
        row = by_key[key]
        status = str(row.get("status", ""))
        terminal = row.get("terminal_reason")
        parsed_fields: dict[str, dict[str, Any]] = {}
        if status not in COVERAGE_STATUSES:
            failures.append(f"{key}:invalid status {status!r}")
        if (
            status
            in {
                "complete_to_observed_boundary",
                "partial",
                "permission_denied",
                "not_applicable",
            }
            and not terminal
        ):
            failures.append(f"{key}:status {status} has no terminal reason")
        if endpoint.live_enabled and terminal == "not_attempted":
            failures.append(f"{key}:live endpoint was not attempted or dispositioned")
        for field in ("requested_bounds_json", "observed_bounds_json", "details_json"):
            value = row.get(field)
            if value is None:
                continue
            try:
                parsed = json.loads(value)
            except json.JSONDecodeError:
                failures.append(f"{key}:{field} is invalid JSON")
                continue
            if not isinstance(parsed, dict):
                failures.append(f"{key}:{field} is not an object")
            else:
                parsed_fields[field] = parsed
        # Wallet history lanes whose provider rejected parameters for whole
        # currencies need source-limit validation even when the walk completed:
        # "complete_to_observed_boundary" with dispositioned rejections still
        # names the rejected currencies as known gaps.
        dispositioned_complete = (
            endpoint.path in SOURCE_LIMITED_WALLET_HISTORY_PATHS
            and status == "complete_to_observed_boundary"
            and terminal == TERMINAL_STRUCTURED_REJECTIONS_DISPOSITIONED
        )
        if endpoint.path in REQUIRED_CORE_PATHS and (
            status != "complete_to_observed_boundary" or dispositioned_complete
        ):
            details = parsed_fields.get("details_json", {})
            failed_currencies = details.get("failed_currencies")
            currency_count = details.get("currency_count")
            failure_count = details.get("failed_currency_count")
            dispositions = details.get("dispositions")
            endpoint_requests = [
                request
                for request in request_rows
                if str(request.get("endpoint")) == endpoint.path
            ]
            request_currencies: set[str] = set()
            rejected_currencies: set[str] = set()
            rejection_evidence_ok = True
            for request in endpoint_requests:
                try:
                    params = json.loads(str(request["params_json"]))
                except (KeyError, TypeError, json.JSONDecodeError):
                    rejection_evidence_ok = False
                    continue
                currency = str(params.get("currency", ""))
                if currency:
                    request_currencies.add(currency)
                response_status = int(request.get("status_code", 0))
                try:
                    raw_root = (run_dir / "raw").resolve()
                    raw_path = (raw_root / str(request["raw_relpath"])).resolve()
                    if not raw_path.is_relative_to(raw_root):
                        raise ValueError("raw response path escapes run")
                    payload = json.loads(raw_path.read_bytes().decode("utf-8"))
                except (OSError, UnicodeDecodeError, json.JSONDecodeError, ValueError):
                    rejection_evidence_ok = False
                    continue
                api_code = payload.get("code") if isinstance(payload, dict) else None
                if response_status == 200 and api_code in (None, 0, "0"):
                    continue
                if (
                    response_status not in {200, 400}
                    or not currency
                    or str(api_code) != "18100410"
                    or payload.get("message") != "Invalid params"
                ):
                    rejection_evidence_ok = False
                    continue
                rejected_currencies.add(currency)
            if dispositioned_complete:
                disposition_items = (
                    dispositions if isinstance(dispositions, list) else []
                )
                disposition_currencies = [
                    item.get("currency")
                    for item in disposition_items
                    if isinstance(item, dict)
                ]
                valid_source_limit = (
                    bool(disposition_items)
                    and len(disposition_currencies) == len(disposition_items)
                    and all(
                        isinstance(currency, str) and currency
                        for currency in disposition_currencies
                    )
                    and len(set(disposition_currencies)) == len(disposition_currencies)
                    and failed_currencies == []
                    and failure_count == 0
                    and details.get("rejected_currency_count")
                    == len(disposition_currencies)
                    and isinstance(currency_count, int)
                    and currency_count == len(request_currencies)
                    and rejected_currencies == set(disposition_currencies)
                    and rejection_evidence_ok
                )
            else:
                valid_source_limit = (
                    endpoint.path in SOURCE_LIMITED_WALLET_HISTORY_PATHS
                    and status == "partial"
                    and terminal == TERMINAL_ENUMERATED_CURRENCY_FAILURES
                    and isinstance(failed_currencies, list)
                    and bool(failed_currencies)
                    and all(
                        isinstance(currency, str) and currency
                        for currency in failed_currencies
                    )
                    and len(set(failed_currencies)) == len(failed_currencies)
                    and failure_count == len(failed_currencies)
                    and isinstance(currency_count, int)
                    and currency_count == len(request_currencies)
                    and rejected_currencies == set(failed_currencies)
                    and rejection_evidence_ok
                )
            if valid_source_limit:
                prefix = _WALLET_PATH_ATOM_PREFIX.get(
                    endpoint.path,
                    endpoint.path.lstrip("/").replace("/", "_"),
                )
                for currency in sorted(rejected_currencies):
                    source_limited.append(f"{prefix}:{currency}")
            else:
                failures.append(f"{key}:required endpoint status is {status}")
    counts: dict[str, int] = defaultdict(int)
    for row in rows:
        counts[str(row.get("status"))] += 1
    return {
        "ok": not failures and len(rows) == len(expected),
        "status": "PARTIAL"
        if source_limited and not failures
        else ("PASS" if not failures else "FAIL"),
        "checked": len(rows),
        "status_counts": dict(sorted(counts.items())),
        "source_limited_reasons": source_limited,
        "failures": failures,
    }


def _dataset_manifest(
    run_dir: Path,
    datasets: Mapping[str, Sequence[Mapping[str, Any]]],
) -> tuple[dict[str, dict[str, str]], list[str]]:
    path = run_dir / "reports" / "dataset_manifest.csv"
    failures: list[str] = []
    if not path.is_file():
        return {}, ["dataset manifest missing"]
    try:
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.DictReader(handle))
    except (OSError, csv.Error) as error:
        return {}, [f"dataset manifest unreadable:{type(error).__name__}"]
    manifest: dict[str, dict[str, str]] = {}
    for row in rows:
        name = str(row.get("dataset", ""))
        if not name or name in manifest:
            failures.append(f"invalid or duplicate dataset manifest row: {name!r}")
            continue
        manifest[name] = {
            str(key): str(value) for key, value in row.items() if key is not None
        }
    if set(manifest) != set(datasets):
        failures.append(
            f"dataset manifest set differs: manifest={len(manifest)},sqlite={len(datasets)}"
        )
    for name in sorted(set(manifest) & set(datasets)):
        row = manifest[name]
        if row.get("row_count") != str(len(datasets[name])):
            failures.append(f"{name}:manifest row_count mismatch")
        if row.get("sha256") != dataset_digest(datasets[name]):
            failures.append(f"{name}:manifest digest mismatch")
    return manifest, failures


def verify_csvs(
    run_dir: Path,
    datasets: Mapping[str, Sequence[Mapping[str, Any]]] | None = None,
) -> dict[str, Any]:
    if datasets is None:
        _run, loaded, state_failures = _read_state(run_dir)
        datasets = loaded
    else:
        state_failures = []
    failures = list(state_failures)
    expected_files = {_safe_filename(name) + ".csv": name for name in datasets}
    if len(expected_files) != len(datasets):
        failures.append("CSV filename collision between datasets")
    actual_files = {path.name: path for path in (run_dir / "csv").glob("*.csv")}
    if set(actual_files) != set(expected_files):
        failures.append(
            f"CSV file set differs: expected={len(expected_files)},actual={len(actual_files)}"
        )
    details: dict[str, int] = {}
    for filename in sorted(set(actual_files) & set(expected_files)):
        name = expected_files[filename]
        rows = list(datasets[name])
        columns = _columns(rows)
        expected_rows = [
            {column: protect_csv_formula(row.get(column)) for column in columns}
            for row in sorted(rows, key=lambda item: _row_sort_key(item, columns))
        ]
        try:
            with actual_files[filename].open(
                "r", encoding="utf-8-sig", newline=""
            ) as handle:
                reader = csv.DictReader(handle)
                actual_rows = list(reader)
                actual_columns = list(reader.fieldnames or [])
        except (OSError, csv.Error) as error:
            failures.append(f"{filename}:unreadable:{type(error).__name__}")
            continue
        details[filename] = len(actual_rows)
        if actual_columns != columns:
            failures.append(f"{filename}:header mismatch")
        if actual_rows != expected_rows:
            failures.append(f"{filename}:content/digest parity mismatch")
    return {
        "ok": not failures and bool(details),
        "row_counts": details,
        "failures": failures,
    }


def _workbook_specs(
    rows_by_name: Mapping[str, Sequence[Mapping[str, Any]]],
) -> tuple[
    list[str], dict[str, list[tuple[str, Sequence[Mapping[str, Any]], Sequence[Any]]]]
]:
    used: set[str] = set()
    names = [_unique_sheet_name("README", used)]
    specs: dict[str, list[tuple[str, Sequence[Mapping[str, Any]], Sequence[Any]]]] = (
        defaultdict(list)
    )
    for dataset_name in sorted(rows_by_name):
        rows = list(rows_by_name[dataset_name])
        columns = _columns(rows)
        ordered = sorted(rows, key=lambda row: _row_sort_key(row, columns))
        expanded = _expanded_columns(ordered)
        if not expanded:
            ordered = [{"status": "no rows"}]
            expanded = _expanded_columns(ordered)
        for part_number, start in enumerate(
            range(0, len(ordered), EXCEL_MAX_DATA_ROWS), start=1
        ):
            part_rows = ordered[start : start + EXCEL_MAX_DATA_ROWS]
            split = len(ordered) > EXCEL_MAX_DATA_ROWS
            sheet_name = _unique_sheet_name(
                dataset_name, used, part_number if split else None
            )
            names.append(sheet_name)
            specs[dataset_name].append((sheet_name, part_rows, expanded))
    return names, dict(specs)


def _readme_run_id(workbook: Any) -> str | None:
    if "README" not in workbook.sheetnames:
        return None
    sheet = workbook["README"]
    for row in sheet.iter_rows(min_row=1, max_col=2, values_only=True):
        if row and row[0] == "Run ID":
            return None if len(row) < 2 or row[1] is None else str(row[1])
    return None


def verify_xlsx(
    run_dir: Path,
    datasets: Mapping[str, Sequence[Mapping[str, Any]]] | None = None,
) -> dict[str, Any]:
    if datasets is None:
        _run, loaded, state_failures = _read_state(run_dir)
        datasets = loaded
    else:
        state_failures = []
    failures = list(state_failures)
    groups = partition_datasets(datasets)
    expected_files = {
        f"matrixport_{kind}_{run_dir.name}.xlsx": kind
        for kind in ("strategy", "account", "audit")
    }
    actual_files = {path.name: path for path in (run_dir / "xlsx").glob("*.xlsx")}
    if set(actual_files) != set(expected_files):
        failures.append(
            f"XLSX file set differs: expected={len(expected_files)},actual={len(actual_files)}"
        )
    details: dict[str, Any] = {}
    for filename in sorted(set(actual_files) & set(expected_files)):
        kind = expected_files[filename]
        group = groups[kind]
        expected_sheets, specs = _workbook_specs(group)
        try:
            workbook = load_workbook(
                actual_files[filename], read_only=True, data_only=False
            )
        except Exception as error:
            failures.append(f"{filename}:unreadable:{type(error).__name__}")
            continue
        try:
            workbook_failures: list[str] = []
            if workbook.sheetnames != expected_sheets:
                workbook_failures.append("sheet set/order mismatch")
            if _readme_run_id(workbook) != run_dir.name:
                workbook_failures.append("README run ID mismatch")
            rebuilt: dict[str, list[dict[str, str]]] = defaultdict(list)
            for dataset_name, sheet_specs in specs.items():
                original_rows = list(group[dataset_name])
                original_columns = _columns(original_rows)
                placeholder = not original_rows
                for sheet_name, part_rows, expanded in sheet_specs:
                    if sheet_name not in workbook.sheetnames:
                        continue
                    sheet = workbook[sheet_name]
                    header = [
                        value
                        for value in next(
                            sheet.iter_rows(
                                min_row=1,
                                max_row=1,
                                max_col=len(expanded),
                                values_only=True,
                            )
                        )
                    ]
                    expected_header = [column.name for column in expanded]
                    if header != expected_header:
                        workbook_failures.append(f"{sheet_name}:header mismatch")
                        continue
                    values = list(
                        sheet.iter_rows(
                            min_row=2,
                            max_col=len(expanded),
                            values_only=False,
                        )
                    )
                    if len(values) != len(part_rows):
                        workbook_failures.append(
                            f"{sheet_name}:row count {len(values)} != {len(part_rows)}"
                        )
                    for excel_row in values:
                        for cell in excel_row:
                            if cell.data_type == "f":
                                workbook_failures.append(
                                    f"{sheet_name}!{cell.coordinate}:formula cell"
                                )
                            if cell.data_type == "e" or (
                                cell.data_type != "s" and cell.value in EXCEL_ERRORS
                            ):
                                workbook_failures.append(
                                    f"{sheet_name}!{cell.coordinate}:Excel error"
                                )
                        if placeholder:
                            continue
                        item: dict[str, str] = {}
                        for source in original_columns:
                            parts: list[tuple[int, str]] = []
                            for index, column in enumerate(expanded):
                                if column.kind == "text" and column.source == source:
                                    value = excel_row[index].value
                                    parts.append(
                                        (
                                            column.part,
                                            "" if value is None else str(value),
                                        )
                                    )
                            item[source] = "".join(
                                value for _part, value in sorted(parts)
                            )
                        rebuilt[dataset_name].append(item)
                if not placeholder and dataset_digest(
                    rebuilt[dataset_name]
                ) != dataset_digest(original_rows):
                    workbook_failures.append(
                        f"{dataset_name}:authoritative text digest mismatch"
                    )
            if workbook_failures:
                failures.extend(f"{filename}:{item}" for item in workbook_failures)
            details[filename] = {
                "sheets": workbook.sheetnames,
                "dataset_rows": {name: len(rows) for name, rows in rebuilt.items()},
                "failures": workbook_failures,
            }
        finally:
            workbook.close()
    return {
        "ok": not failures and len(details) == 3,
        "workbooks": details,
        "failures": failures,
    }


def verify_parity(run_dir: Path) -> dict[str, Any]:
    run, datasets, state_failures = _read_state(run_dir)
    _manifest, manifest_failures = _dataset_manifest(run_dir, datasets)
    csv_result = verify_csvs(run_dir, datasets)
    xlsx_result = verify_xlsx(run_dir, datasets)
    failures = [*state_failures, *manifest_failures]
    if run is None:
        failures.append("run state missing")
    if not csv_result["ok"]:
        failures.extend(f"CSV:{item}" for item in csv_result["failures"])
    if not xlsx_result["ok"]:
        failures.extend(f"XLSX:{item}" for item in xlsx_result["failures"])
    return {
        "ok": not failures,
        "dataset_count": len(datasets),
        "dataset_digests": {
            name: dataset_digest(rows) for name, rows in sorted(datasets.items())
        },
        "failures": failures,
    }


def _nonblank(value: Any) -> bool:
    return value is not None and str(value) != ""


def _nonzero(value: Any) -> bool:
    if not _nonblank(value):
        return False
    try:
        return Decimal(str(value)) != 0
    except InvalidOperation:
        return True


def verify_valuation(
    datasets: Mapping[str, Sequence[Mapping[str, Any]]],
) -> dict[str, Any]:
    """Evaluate valuation acceptance without conflating gaps with corruption.

    ``ok`` reports logical/contract integrity. ``status`` reports whether the
    factual valuation is complete (PASS), explicitly incomplete (PARTIAL), or
    unavailable/contradictory (FAIL).
    """

    structural_failures: list[str] = []
    acceptance_reasons: list[str] = []
    required = ("holding_reconciliation", "daily_holdings", "daily_portfolio_totals")
    missing = [name for name in required if name not in datasets]
    if missing:
        structural_failures.append(f"missing valuation datasets: {', '.join(missing)}")
        return {
            "ok": False,
            "status": "FAIL",
            "reconciliation_counts": {},
            "failures": structural_failures,
            "reasons": ["valuation datasets are missing"],
        }

    reconciliation = list(datasets["holding_reconciliation"])
    holdings = list(datasets["daily_holdings"])
    totals = list(datasets["daily_portfolio_totals"])
    by_product: dict[str, str] = {}
    counts: dict[str, int] = defaultdict(int)
    for index, row in enumerate(reconciliation, start=1):
        product_id = str(row.get("product_id", ""))
        classification = str(row.get("classification", ""))
        if not product_id or product_id in by_product:
            structural_failures.append(
                f"holding_reconciliation row {index}:missing/duplicate product_id"
            )
            continue
        if classification not in {"exact", "anchored_gap", "unavailable"}:
            structural_failures.append(
                f"holding_reconciliation {product_id}:unknown classification {classification!r}"
            )
            continue
        by_product[product_id] = classification
        counts[classification] += 1
        if classification == "anchored_gap":
            acceptance_reasons.append(f"{product_id}:anchored_gap")
        elif classification == "unavailable":
            acceptance_reasons.append(f"{product_id}:unavailable")

    missing_coverage = False
    for index, row in enumerate(holdings, start=1):
        product_id = str(row.get("product_id", ""))
        classification = str(row.get("reconciliation_class", ""))
        expected = by_product.get(product_id)
        if expected is None:
            structural_failures.append(
                f"daily_holdings row {index}:unknown product {product_id!r}"
            )
        elif classification != expected:
            structural_failures.append(
                f"daily_holdings row {index}:{product_id} reconciliation label {classification!r} != {expected!r}"
            )
        active = _nonzero(row.get("share"))
        nav_missing = active and str(row.get("nav_available", "")).lower() != "true"
        fx_missing = active and str(row.get("fx_available", "")).lower() != "true"
        strict_component = row.get("value_usd_strict_component")
        if nav_missing or fx_missing:
            missing_coverage = True
            if _nonblank(strict_component):
                structural_failures.append(
                    f"daily_holdings row {index}:{product_id} has strict value despite missing NAV/FX"
                )
            if not _nonblank(row.get("missing_reason")):
                structural_failures.append(
                    f"daily_holdings row {index}:{product_id} missing NAV/FX without missing_reason"
                )

    coverage_reasons: set[str] = set()
    for index, row in enumerate(totals, start=1):
        active = _nonzero(row.get("active_product_count"))
        nav_complete = str(row.get("nav_coverage_ratio", "")) == "1"
        fx_complete = str(row.get("fx_coverage_ratio", "")) == "1"
        declared_complete = str(row.get("strict_complete", "")).lower() == "true"
        strict_value = row.get("portfolio_value_usd_strict")
        reasons = str(row.get("missing_reasons") or "")
        utc_date = str(row.get("utc_date", ""))
        if active:
            if not nav_complete and utc_date:
                missing_coverage = True
                coverage_reasons.add(f"nav_missing_date:{utc_date}")
            if not fx_complete and utc_date:
                missing_coverage = True
                coverage_reasons.add(f"fx_missing_date:{utc_date}")
            if not declared_complete and utc_date:
                missing_coverage = True
                coverage_reasons.add(f"reconciliation_incomplete_date:{utc_date}")
        incomplete = active and (
            not nav_complete or not fx_complete or not declared_complete
        )
        if incomplete:
            if _nonblank(strict_value):
                structural_failures.append(
                    f"daily_portfolio_totals row {index} has strict total despite missing NAV/FX/reconciliation"
                )
            if not reasons:
                structural_failures.append(
                    f"daily_portfolio_totals row {index} incomplete without missing_reasons"
                )
        elif active and declared_complete and not _nonblank(strict_value):
            structural_failures.append(
                f"daily_portfolio_totals row {index} complete but strict total is blank"
            )

    if structural_failures:
        status = "FAIL"
    elif counts.get("unavailable", 0):
        status = "FAIL"
    elif counts.get("anchored_gap", 0) or missing_coverage:
        status = "PARTIAL"
        if missing_coverage:
            acceptance_reasons.extend(sorted(coverage_reasons))
    else:
        status = "PASS"
    return {
        "ok": not structural_failures,
        "status": status,
        "reconciliation_counts": dict(sorted(counts.items())),
        "daily_holding_rows": len(holdings),
        "daily_total_rows": len(totals),
        "failures": structural_failures,
        "reasons": sorted(set(acceptance_reasons)),
    }


def verify_run(run_dir: Path, secret_needles: Iterable[bytes] = ()) -> dict[str, Any]:
    run_dir = run_dir.resolve()
    _run, datasets, _state_failures = _read_state(run_dir)
    checks = {
        "raw": verify_raw(run_dir),
        "sqlite": verify_sqlite(run_dir),
        "requests": verify_requests(run_dir),
        "coverage": verify_coverage(run_dir),
        "csv": verify_csvs(run_dir, datasets),
        "xlsx": verify_xlsx(run_dir, datasets),
        "parity": verify_parity(run_dir),
        "valuation": verify_valuation(datasets),
    }
    leaked = secret_scan(run_dir, secret_needles)
    checks["privacy"] = {"ok": not leaked, "matches": leaked}
    structural_ok = all(check.get("ok") is True for check in checks.values())
    acceptance_status = "FAIL"
    if structural_ok:
        acceptance_status = (
            "PARTIAL"
            if "PARTIAL"
            in {checks["coverage"].get("status"), checks["valuation"].get("status")}
            else checks["valuation"]["status"]
        )
    result = {
        "run_id": run_dir.name,
        "ok": structural_ok,
        "status": acceptance_status,
        "checks": checks,
    }
    target = run_dir / "reports" / "verification.json"
    atomic_write(target, (canonical_json(result) + "\n").encode("utf-8"))
    return result
